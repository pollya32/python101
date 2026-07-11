from flask import Flask, render_template, request, jsonify, send_file, Response, session, redirect, url_for
import sqlite3
import os
import csv
import io
import random
import re
import secrets
import socket
import shutil
import traceback
import html as html_lib
from urllib.parse import quote
from datetime import date, datetime, timedelta
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from werkzeug.security import generate_password_hash, check_password_hash
import threading
import time
import requests
from html.parser import HTMLParser

app = Flask(__name__)
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=30)
DB_PATH = os.path.join(os.path.dirname(__file__), "equipment.db")
BACKUP_DIR = os.path.join(os.path.dirname(__file__), "backups")

EQUIPMENT_COUNT = 20
EQUIPMENT_PREFIX = "TEAG"
MASTER_EQUIPMENT_ID = 1  # TEAG01호기: 이 설비에 추가한 부품은 동일한 이름의 유닛을 가진 나머지 설비에도 자동 복제된다
DEFAULT_PASSWORD = "0000"
# ══ 사내 메일 API 설정 ═══════════════════════════════════════════════
# (보안) 실제 값으로 교체한 뒤에는 이 파일을 외부에 공유/업로드하지 마세요.
os.environ["no_proxy"] = "openapi.samsung.net"
MAIL_SENDER_ID = "lbr-32.lee"                    # 발신자 녹스 ID (@samsung.com 앞부분)
MAIL_AUTHORIZATION = "XXXXXXXXXXXXXXX"           # ← 실제 Authorization 값으로 교체
MAIL_SYSTEM_ID = "XXXXXXXXXXXXXX"                # ← 실제 System-ID 값으로 교체
MAIL_SUBJECT = "[부품관리] TES 설비 부품 현황"
# ════════════════════════════════════════════════════════════════════
MAX_DRAWING_DATA_LEN = 8 * 1024 * 1024  # 도면 이미지(base64 data URL) 최대 길이, 원본 파일 약 5MB에 해당


def dashboard_grid_pos(index, cols=5):
    """대시보드 캔버스 안에서 index번째 설비의 기본 격자 위치(% 좌표)를 계산한다."""
    row = index // cols
    col = index % cols
    x = 10 + col * (80 / max(cols - 1, 1))
    y = min(15 + row * 22, 92)
    return round(x, 2), round(y, 2)

# 사진 속 설비(로드포트 4개 + HMI 제어패널 + 공정모듈 + 배기/시그널타워) 구조를 본뜬 기본 유닛
# pos_x, pos_y는 설비 캔버스 안에서 유닛 중심의 위치(% 좌표)
DEFAULT_UNITS = [
    ("배기 유닛 / 시그널 타워", "🚨", "#6b7280", 50, 18),
    ("로드포트 1", "📦", "#2563eb", 10, 60),
    ("로드포트 2", "📦", "#2563eb", 26, 60),
    ("로드포트 3", "📦", "#2563eb", 42, 60),
    ("로드포트 4", "📦", "#2563eb", 58, 60),
    ("HMI 제어 패널", "🖥️", "#0f766e", 76, 60),
    ("공정 모듈 (파워유닛)", "🔧", "#7c3aed", 92, 60),
]


def get_db():
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    return conn


_RICH_DANGEROUS_BLOCK_RE = re.compile(
    r"<(script|style|iframe|object|embed|link|meta|form)\b[^>]*>.*?</\1\s*>"
    r"|<(script|style|iframe|object|embed|link|meta|form)\b[^>]*/?>",
    re.IGNORECASE | re.DOTALL,
)
_RICH_EVENT_ATTR_RE = re.compile(r"""\son\w+\s*=\s*("[^"]*"|'[^']*'|[^\s>]+)""", re.IGNORECASE)
_RICH_JS_URI_RE = re.compile(r"""(href|src)\s*=\s*("|')\s*(javascript|data):[^"']*\2""", re.IGNORECASE)


def sanitize_rich_html(raw):
    """메모/노트에 저장되는 HTML에 대한 서버측 2차 방어선(주 방어는 클라이언트의 화이트리스트
    기반 sanitizeRichHtml). script/style/iframe 등 위험 태그와 on*= 이벤트 속성, javascript:/data:
    URI를 제거한다."""
    if not raw:
        return raw
    cleaned = _RICH_DANGEROUS_BLOCK_RE.sub("", raw)
    cleaned = _RICH_EVENT_ATTR_RE.sub("", cleaned)
    cleaned = _RICH_JS_URI_RE.sub(lambda m: f'{m.group(1)}="#"', cleaned)
    return cleaned


def legacy_text_to_rich_html(text):
    """리치 메모 기능 도입 이전에 평문으로 저장되어 있던 메모/노트를, 기존과 동일하게 보이는
    안전한 HTML로 1회 변환한다 (이스케이프 + 줄바꿈을 <br>로 + URL 자동 링크화)."""
    if not text:
        return text
    escaped = html_lib.escape(text, quote=False)
    escaped = re.sub(
        r"(https?://[^\s<]+)",
        r'<a href="\1" target="_blank" rel="noopener noreferrer">\1</a>',
        escaped,
    )
    return escaped.replace("\n", "<br>")


def get_lan_ip():
    """같은 네트워크의 다른 사람이 접속할 수 있는 이 컴퓨터의 IP 주소를 알아낸다."""
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))
        return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"
    finally:
        s.close()


def get_config(conn, key):
    row = conn.execute("SELECT value FROM app_config WHERE key = ?", (key,)).fetchone()
    return row["value"] if row else None


def set_config(conn, key, value):
    conn.execute(
        "INSERT INTO app_config (key, value) VALUES (?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
        (key, value),
    )


def log_activity(conn, action, target_type, target_id, target_name, detail=""):
    """설비/유닛/부품 등에 대한 주요 변경 작업을 활동 이력에 기록한다.
    현재 세션에 저장된 사용자 이름을 행위자로 남긴다."""
    actor = session.get("user_name") or "익명"
    conn.execute(
        "INSERT INTO activity_log (actor_name, action, target_type, target_id, target_name, detail) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (actor, action, target_type, target_id, target_name, detail),
    )


def init_db():
    conn = get_db()
    c = conn.cursor()

    c.execute("CREATE TABLE IF NOT EXISTS app_config (key TEXT PRIMARY KEY, value TEXT)")
    if not c.execute("SELECT 1 FROM app_config WHERE key = 'password_hash'").fetchone():
        c.execute(
            "INSERT INTO app_config (key, value) VALUES ('password_hash', ?)",
            (generate_password_hash(DEFAULT_PASSWORD),),
        )
    if not c.execute("SELECT 1 FROM app_config WHERE key = 'secret_key'").fetchone():
        c.execute(
            "INSERT INTO app_config (key, value) VALUES ('secret_key', ?)",
            (secrets.token_hex(32),),
        )
    conn.commit()

    c.execute("""
        CREATE TABLE IF NOT EXISTS activity_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            actor_name TEXT,
            action TEXT NOT NULL,
            target_type TEXT,
            target_id INTEGER,
            target_name TEXT,
            detail TEXT,
            created_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    conn.commit()

    c.execute("""
        CREATE TABLE IF NOT EXISTS equipments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            icon TEXT DEFAULT '🏭',
            pos_x REAL DEFAULT 50,
            pos_y REAL DEFAULT 50,
            location TEXT,
            setup_date TEXT,
            deleted_at TEXT,
            created_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    existing_eq_cols = {r["name"] for r in c.execute("PRAGMA table_info(equipments)").fetchall()}
    if "icon" not in existing_eq_cols:
        c.execute("ALTER TABLE equipments ADD COLUMN icon TEXT DEFAULT '🏭'")
        c.execute("UPDATE equipments SET icon = '🏭' WHERE icon IS NULL")
    if "pos_x" not in existing_eq_cols:
        c.execute("ALTER TABLE equipments ADD COLUMN pos_x REAL")
        c.execute("ALTER TABLE equipments ADD COLUMN pos_y REAL")
    if "location" not in existing_eq_cols:
        c.execute("ALTER TABLE equipments ADD COLUMN location TEXT")
    if "setup_date" not in existing_eq_cols:
        c.execute("ALTER TABLE equipments ADD COLUMN setup_date TEXT")
    if "deleted_at" not in existing_eq_cols:
        c.execute("ALTER TABLE equipments ADD COLUMN deleted_at TEXT")
    eq_count = c.execute("SELECT COUNT(*) AS n FROM equipments").fetchone()["n"]
    if eq_count == 0:
        for i in range(1, EQUIPMENT_COUNT + 1):
            x, y = dashboard_grid_pos(i - 1)
            c.execute(
                "INSERT INTO equipments (id, name, pos_x, pos_y) VALUES (?, ?, ?, ?)",
                (i, f"{EQUIPMENT_PREFIX}{i:02d}호기", x, y),
            )
    unplaced_eq = c.execute(
        "SELECT id FROM equipments WHERE pos_x IS NULL OR pos_y IS NULL ORDER BY id"
    ).fetchall()
    for i, row in enumerate(unplaced_eq):
        x, y = dashboard_grid_pos(i)
        c.execute("UPDATE equipments SET pos_x = ?, pos_y = ? WHERE id = ?", (x, y, row["id"]))

    c.execute("""
        CREATE TABLE IF NOT EXISTS units (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            equipment_id INTEGER NOT NULL DEFAULT 1,
            name TEXT NOT NULL,
            icon TEXT DEFAULT '⚙️',
            color TEXT DEFAULT '#1a3a5c',
            pos_x REAL DEFAULT 50,
            pos_y REAL DEFAULT 50,
            width REAL DEFAULT 140,
            height REAL DEFAULT 110,
            drawing_data TEXT,
            deleted_at TEXT,
            created_at TEXT DEFAULT (datetime('now','localtime')),
            FOREIGN KEY (equipment_id) REFERENCES equipments(id) ON DELETE CASCADE
        )
    """)
    existing_cols = {r["name"] for r in c.execute("PRAGMA table_info(units)").fetchall()}
    if "equipment_id" not in existing_cols:
        c.execute("ALTER TABLE units ADD COLUMN equipment_id INTEGER")
        c.execute("UPDATE units SET equipment_id = 1 WHERE equipment_id IS NULL")
    if "pos_x" not in existing_cols:
        c.execute("ALTER TABLE units ADD COLUMN pos_x REAL")
        c.execute("ALTER TABLE units ADD COLUMN pos_y REAL")
    if "width" not in existing_cols:
        c.execute("ALTER TABLE units ADD COLUMN width REAL")
        c.execute("ALTER TABLE units ADD COLUMN height REAL")
    if "deleted_at" not in existing_cols:
        c.execute("ALTER TABLE units ADD COLUMN deleted_at TEXT")
    if "drawing_data" not in existing_cols:
        c.execute("ALTER TABLE units ADD COLUMN drawing_data TEXT")
    unplaced = c.execute(
        "SELECT id FROM units WHERE pos_x IS NULL OR pos_y IS NULL ORDER BY id"
    ).fetchall()
    for i, row in enumerate(unplaced):
        x = 10 + (i * 84 / max(len(unplaced) - 1, 1)) if len(unplaced) > 1 else 50
        c.execute("UPDATE units SET pos_x = ?, pos_y = ? WHERE id = ?", (x, 50, row["id"]))
    c.execute("UPDATE units SET width = 140 WHERE width IS NULL")
    c.execute("UPDATE units SET height = 110 WHERE height IS NULL")

    c.execute("""
        CREATE TABLE IF NOT EXISTS parts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            unit_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            spec TEXT,
            cycle_days INTEGER,
            cycle_unit TEXT DEFAULT 'N/A',
            cost REAL DEFAULT 0,
            last_replaced_date TEXT,
            note TEXT,
            memo TEXT,
            drawing_data TEXT,
            icon TEXT DEFAULT '🔩',
            pos_x REAL DEFAULT 50,
            pos_y REAL DEFAULT 50,
            width REAL DEFAULT 130,
            height REAL DEFAULT 110,
            stock_qty INTEGER DEFAULT 0,
            supplier TEXT,
            supplier_contact TEXT,
            lead_time_days INTEGER,
            deleted_at TEXT,
            created_at TEXT DEFAULT (datetime('now','localtime')),
            FOREIGN KEY (unit_id) REFERENCES units(id) ON DELETE CASCADE
        )
    """)
    # 기존 DB는 cycle_days가 NOT NULL(기본값 90)이었다. 교체 주기를 "N/A"(주기 없음)로
    # 남겨둘 수 있으려면 NULL을 허용해야 하는데, SQLite는 컬럼의 NOT NULL 제약을
    # 직접 제거할 수 없으므로 테이블을 재생성해서 옮겨준다.
    # (parts를 다른 이름으로 RENAME했다가 다시 만드는 방식은, replacement_history의
    #  FOREIGN KEY ... ON DELETE CASCADE가 RENAME된 임시 이름을 따라가 버려서 임시
    #  테이블을 DROP하는 순간 거기 딸린 교체 이력이 CASCADE로 통째로 삭제되거나,
    #  FK 정의가 존재하지 않는 임시 테이블 이름을 계속 가리키게 되는 문제가 있었다.
    #  대신 새 테이블을 다른 이름으로 만들어 데이터를 옮긴 뒤 기존 parts를 지우고
    #  새 테이블을 parts로 RENAME해서, replacement_history의 FK 정의("parts" 참조)가
    #  한 번도 다른 이름을 가리키지 않도록 한다. SQLite 공식 가이드대로 이 구간만
    #  foreign_keys를 잠시 꺼서 진행한다.)
    part_col_info = c.execute("PRAGMA table_info(parts)").fetchall()
    if any(r["name"] == "cycle_days" and r["notnull"] for r in part_col_info):
        old_cols = ", ".join(r["name"] for r in part_col_info)
        conn.commit()  # foreign_keys pragma는 열려있는 트랜잭션이 없어야 적용된다
        c.execute("PRAGMA foreign_keys = OFF")
        c.execute("""
            CREATE TABLE parts_na_migrated (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                unit_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                spec TEXT,
                cycle_days INTEGER,
                cycle_unit TEXT DEFAULT 'N/A',
                cost REAL DEFAULT 0,
                last_replaced_date TEXT,
                note TEXT,
                memo TEXT,
                drawing_data TEXT,
                icon TEXT DEFAULT '🔩',
                pos_x REAL DEFAULT 50,
                pos_y REAL DEFAULT 50,
                width REAL DEFAULT 130,
                height REAL DEFAULT 110,
                stock_qty INTEGER DEFAULT 0,
                supplier TEXT,
                supplier_contact TEXT,
                lead_time_days INTEGER,
                deleted_at TEXT,
                created_at TEXT DEFAULT (datetime('now','localtime')),
                FOREIGN KEY (unit_id) REFERENCES units(id) ON DELETE CASCADE
            )
        """)
        c.execute(f"INSERT INTO parts_na_migrated ({old_cols}) SELECT {old_cols} FROM parts")
        c.execute("DROP TABLE parts")
        c.execute("ALTER TABLE parts_na_migrated RENAME TO parts")
        conn.commit()
        c.execute("PRAGMA foreign_keys = ON")

    existing_part_cols = {r["name"] for r in c.execute("PRAGMA table_info(parts)").fetchall()}
    if "icon" not in existing_part_cols:
        c.execute("ALTER TABLE parts ADD COLUMN icon TEXT DEFAULT '🔩'")
    if "pos_x" not in existing_part_cols:
        c.execute("ALTER TABLE parts ADD COLUMN pos_x REAL")
        c.execute("ALTER TABLE parts ADD COLUMN pos_y REAL")
        c.execute("ALTER TABLE parts ADD COLUMN width REAL")
        c.execute("ALTER TABLE parts ADD COLUMN height REAL")
    if "cost" not in existing_part_cols:
        c.execute("ALTER TABLE parts ADD COLUMN cost REAL DEFAULT 0")
        c.execute("UPDATE parts SET cost = 0 WHERE cost IS NULL")
    if "cycle_unit" not in existing_part_cols:
        c.execute("ALTER TABLE parts ADD COLUMN cycle_unit TEXT DEFAULT '일'")
        c.execute("UPDATE parts SET cycle_unit = '일' WHERE cycle_unit IS NULL")
    if "memo" not in existing_part_cols:
        c.execute("ALTER TABLE parts ADD COLUMN memo TEXT")
    if "drawing_data" not in existing_part_cols:
        c.execute("ALTER TABLE parts ADD COLUMN drawing_data TEXT")
    if "stock_qty" not in existing_part_cols:
        c.execute("ALTER TABLE parts ADD COLUMN stock_qty INTEGER DEFAULT 0")
        c.execute("ALTER TABLE parts ADD COLUMN supplier TEXT")
        c.execute("ALTER TABLE parts ADD COLUMN supplier_contact TEXT")
        c.execute("ALTER TABLE parts ADD COLUMN lead_time_days INTEGER")
        c.execute("UPDATE parts SET stock_qty = 0 WHERE stock_qty IS NULL")
    if "deleted_at" not in existing_part_cols:
        c.execute("ALTER TABLE parts ADD COLUMN deleted_at TEXT")
    for unit_row in c.execute("SELECT DISTINCT unit_id FROM parts").fetchall():
        unplaced_parts = c.execute(
            "SELECT id FROM parts WHERE unit_id = ? AND (pos_x IS NULL OR pos_y IS NULL) ORDER BY id",
            (unit_row["unit_id"],),
        ).fetchall()
        for i, prow in enumerate(unplaced_parts):
            x = 15 + (i * 70 / max(len(unplaced_parts) - 1, 1)) if len(unplaced_parts) > 1 else 50
            c.execute("UPDATE parts SET pos_x = ?, pos_y = ? WHERE id = ?", (x, 50, prow["id"]))
    c.execute("UPDATE parts SET width = 130 WHERE width IS NULL")
    c.execute("UPDATE parts SET height = 110 WHERE height IS NULL")
    c.execute("UPDATE parts SET icon = '🔩' WHERE icon IS NULL")

    c.execute("""
        CREATE TABLE IF NOT EXISTS replacement_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            part_id INTEGER NOT NULL,
            replaced_date TEXT NOT NULL,
            cost REAL DEFAULT 0,
            note TEXT,
            created_at TEXT DEFAULT (datetime('now','localtime')),
            FOREIGN KEY (part_id) REFERENCES parts(id) ON DELETE CASCADE
        )
    """)
    existing_history_cols = {r["name"] for r in c.execute("PRAGMA table_info(replacement_history)").fetchall()}
    if "cost" not in existing_history_cols:
        c.execute("ALTER TABLE replacement_history ADD COLUMN cost REAL DEFAULT 0")
        c.execute("UPDATE replacement_history SET cost = 0 WHERE cost IS NULL")

    c.execute("""
        CREATE TABLE IF NOT EXISTS equipment_notes (
            equipment_id INTEGER PRIMARY KEY,
            content TEXT DEFAULT '',
            updated_at TEXT DEFAULT (datetime('now','localtime')),
            FOREIGN KEY (equipment_id) REFERENCES equipments(id) ON DELETE CASCADE
        )
    """)
    old_notes_cols = {r["name"] for r in c.execute("PRAGMA table_info(equipment_notes)").fetchall()}
    if "equipment_id" not in old_notes_cols:
        c.execute("ALTER TABLE equipment_notes RENAME TO equipment_notes_old")
        c.execute("""
            CREATE TABLE equipment_notes (
                equipment_id INTEGER PRIMARY KEY,
                content TEXT DEFAULT '',
                updated_at TEXT DEFAULT (datetime('now','localtime')),
                FOREIGN KEY (equipment_id) REFERENCES equipments(id) ON DELETE CASCADE
            )
        """)
        old_row = c.execute("SELECT content, updated_at FROM equipment_notes_old WHERE id = 1").fetchone()
        if old_row:
            c.execute(
                "INSERT INTO equipment_notes (equipment_id, content, updated_at) VALUES (1, ?, ?)",
                (old_row["content"], old_row["updated_at"]),
            )
        c.execute("DROP TABLE equipment_notes_old")
    c.execute(
        "INSERT OR IGNORE INTO equipment_notes (equipment_id, content) "
        "SELECT id, '' FROM equipments"
    )

    c.execute("""
        CREATE TABLE IF NOT EXISTS unit_notes (
            unit_id INTEGER PRIMARY KEY,
            content TEXT DEFAULT '',
            updated_at TEXT DEFAULT (datetime('now','localtime')),
            FOREIGN KEY (unit_id) REFERENCES units(id) ON DELETE CASCADE
        )
    """)

    # 횡전개 현황판: 엑셀에서 붙여넣은 호기/Chamber별 진행 날짜 표(data_html)를 항목별로 저장
    c.execute("""
        CREATE TABLE IF NOT EXISTS rollout_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            data_html TEXT DEFAULT '',
            created_at TEXT DEFAULT (datetime('now','localtime')),
            updated_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)

    # 메일 수신자 목록 (녹스 ID만 저장, @samsung.com 은 발송 시 자동으로 붙임)
    c.execute("""
        CREATE TABLE IF NOT EXISTS mail_recipients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email_id TEXT NOT NULL UNIQUE,
            created_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS bulk_part_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            raw_unit_text TEXT,
            unit_id INTEGER,
            part_name TEXT NOT NULL,
            q_code TEXT,
            note TEXT,
            cost REAL DEFAULT 0,
            cycle_days INTEGER,
            cycle_unit TEXT DEFAULT 'N/A',
            status TEXT DEFAULT 'pending',
            created_at TEXT DEFAULT (datetime('now','localtime')),
            FOREIGN KEY (unit_id) REFERENCES units(id) ON DELETE SET NULL
        )
    """)
    existing_bulk_cols = {r["name"] for r in c.execute("PRAGMA table_info(bulk_part_entries)").fetchall()}
    if "cost" not in existing_bulk_cols:
        c.execute("ALTER TABLE bulk_part_entries ADD COLUMN cost REAL DEFAULT 0")
        c.execute("UPDATE bulk_part_entries SET cost = 0 WHERE cost IS NULL")
    if "cycle_days" not in existing_bulk_cols:
        c.execute("ALTER TABLE bulk_part_entries ADD COLUMN cycle_days INTEGER")
        c.execute("ALTER TABLE bulk_part_entries ADD COLUMN cycle_unit TEXT DEFAULT 'N/A'")
        c.execute("UPDATE bulk_part_entries SET cycle_unit = 'N/A' WHERE cycle_days IS NULL")

    c.execute("""
        CREATE TABLE IF NOT EXISTS bulk_part_entry_units (
            entry_id INTEGER NOT NULL,
            unit_id INTEGER NOT NULL,
            PRIMARY KEY (entry_id, unit_id),
            FOREIGN KEY (entry_id) REFERENCES bulk_part_entries(id) ON DELETE CASCADE,
            FOREIGN KEY (unit_id) REFERENCES units(id) ON DELETE CASCADE
        )
    """)
    if "unit_id" in existing_bulk_cols:
        for legacy in c.execute(
            "SELECT id, unit_id FROM bulk_part_entries WHERE unit_id IS NOT NULL"
        ).fetchall():
            c.execute(
                "INSERT OR IGNORE INTO bulk_part_entry_units (entry_id, unit_id) VALUES (?, ?)",
                (legacy["id"], legacy["unit_id"]),
            )

    # 모든 설비에 공통으로 반영되는 기본 유닛 구성 템플릿
    c.execute("""
        CREATE TABLE IF NOT EXISTS unit_templates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            icon TEXT DEFAULT '⚙️',
            color TEXT DEFAULT '#1a3a5c',
            pos_x REAL DEFAULT 50,
            pos_y REAL DEFAULT 50,
            width REAL DEFAULT 140,
            height REAL DEFAULT 110,
            created_at TEXT DEFAULT (datetime('now','localtime'))
        )
    """)
    template_count = c.execute("SELECT COUNT(*) AS n FROM unit_templates").fetchone()["n"]
    if template_count == 0:
        for name, icon, color, pos_x, pos_y in DEFAULT_UNITS:
            c.execute(
                "INSERT INTO unit_templates (name, icon, color, pos_x, pos_y) VALUES (?, ?, ?, ?, ?)",
                (name, icon, color, pos_x, pos_y),
            )

    # 각 설비에 유닛이 하나도 없으면 기본 유닛 구성 템플릿을 자동으로 반영
    templates = c.execute("SELECT * FROM unit_templates ORDER BY id").fetchall()
    for eq in c.execute("SELECT id FROM equipments").fetchall():
        unit_count = c.execute(
            "SELECT COUNT(*) AS n FROM units WHERE equipment_id = ?", (eq["id"],)
        ).fetchone()["n"]
        if unit_count == 0:
            for t in templates:
                c.execute(
                    "INSERT INTO units (equipment_id, name, icon, color, pos_x, pos_y, width, height) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (eq["id"], t["name"], t["icon"], t["color"], t["pos_x"], t["pos_y"], t["width"], t["height"]),
                )

    # 메모/노트에 리치 텍스트(엑셀 표 붙여넣기 등) 편집 기능이 도입되기 전에 평문으로 저장된
    # 기존 값을, 화면에 보이던 모습 그대로 안전한 HTML로 1회 변환한다 (재실행돼도 한 번만 수행).
    if not c.execute("SELECT 1 FROM app_config WHERE key = 'memo_rich_migrated_v1'").fetchone():
        for row in c.execute("SELECT id, memo FROM parts WHERE memo IS NOT NULL AND memo != ''").fetchall():
            c.execute(
                "UPDATE parts SET memo = ? WHERE id = ?",
                (legacy_text_to_rich_html(row["memo"]), row["id"]),
            )
        for row in c.execute(
            "SELECT equipment_id, content FROM equipment_notes WHERE content IS NOT NULL AND content != ''"
        ).fetchall():
            c.execute(
                "UPDATE equipment_notes SET content = ? WHERE equipment_id = ?",
                (legacy_text_to_rich_html(row["content"]), row["equipment_id"]),
            )
        for row in c.execute(
            "SELECT unit_id, content FROM unit_notes WHERE content IS NOT NULL AND content != ''"
        ).fetchall():
            c.execute(
                "UPDATE unit_notes SET content = ? WHERE unit_id = ?",
                (legacy_text_to_rich_html(row["content"]), row["unit_id"]),
            )
        c.execute("INSERT INTO app_config (key, value) VALUES ('memo_rich_migrated_v1', '1')")

    # 설비/유닛/부품이 많아져도 목록 조회가 느려지지 않도록, 자주 필터링되는 외래키 컬럼에
    # 인덱스를 추가한다 (이미 있으면 아무 일도 하지 않음).
    c.execute("CREATE INDEX IF NOT EXISTS idx_units_equipment_id ON units(equipment_id)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_parts_unit_id ON parts(unit_id)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_replacement_history_part_id ON replacement_history(part_id)")

    conn.commit()
    conn.close()


def part_status(cycle_days, last_replaced_date):
    """부품 교체 주기 대비 경과 상태 계산"""
    if cycle_days is None:
        return {"status": "unknown", "label": "N/A", "days_left": None, "next_due": None}
    if not last_replaced_date:
        return {"status": "unknown", "label": "미기록", "days_left": None, "next_due": None}
    last = datetime.strptime(last_replaced_date, "%Y-%m-%d").date()
    next_due = last + timedelta(days=cycle_days)
    days_left = (next_due - date.today()).days
    if days_left < 0:
        status, label = "overdue", "교체 필요"
    elif cycle_days > 0 and days_left <= max(cycle_days * 0.2, 3):
        status, label = "soon", "교체 임박"
    else:
        status, label = "ok", "정상"
    return {
        "status": status,
        "label": label,
        "days_left": days_left,
        "next_due": next_due.isoformat(),
    }


def serialize_part(row):
    info = part_status(row["cycle_days"], row["last_replaced_date"])
    d = dict(row)
    d.update(info)
    return d


def resolve_cycle(data, current_days=None, current_unit=None):
    """요청 바디의 cycle_days/cycle_unit을 정규화한다.
    두 필드가 모두 없으면(다른 필드만 부분 수정하는 요청) 기존 값을 그대로 유지하고,
    cycle_unit이 'N/A'면 주기 없음(cycle_days=None)으로 처리한다."""
    if "cycle_unit" not in data and "cycle_days" not in data:
        return current_days, current_unit
    cycle_unit = (data.get("cycle_unit") or "").strip()
    if cycle_unit == "N/A":
        return None, "N/A"
    cycle_days_raw = data.get("cycle_days")
    cycle_days = int(cycle_days_raw) if cycle_days_raw not in (None, "") else current_days
    return cycle_days, (cycle_unit or current_unit or "일")


def insert_part(conn, unit_id, name, spec="", cycle_days=None, cycle_unit="N/A", cost=0,
                 last_replaced_date=None, note="", memo="", drawing_data=None, icon="🔩",
                 pos_x=None, pos_y=None, width=130, height=110,
                 stock_qty=0, supplier="", supplier_contact="", lead_time_days=None):
    if pos_x is None or pos_y is None:
        pos_x, pos_y = random.uniform(15, 85), random.uniform(20, 80)
    cur = conn.execute(
        """INSERT INTO parts (unit_id, name, spec, cycle_days, cycle_unit, cost, last_replaced_date, note, memo,
           drawing_data, icon, pos_x, pos_y, width, height, stock_qty, supplier, supplier_contact, lead_time_days)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (unit_id, name, spec, cycle_days, cycle_unit, cost, last_replaced_date, note, memo, drawing_data, icon,
         pos_x, pos_y, width, height, stock_qty, supplier, supplier_contact, lead_time_days),
    )
    part_id = cur.lastrowid
    if last_replaced_date:
        conn.execute(
            "INSERT INTO replacement_history (part_id, replaced_date, note) VALUES (?, ?, ?)",
            (part_id, last_replaced_date, "최초 등록"),
        )
    return part_id


def apply_unit_parts_to_other_equipment(conn, unit_id):
    """기준 설비(TEAG01호기)의 특정 유닛에 등록된 부품 구성 전체를, 동일한 이름의 유닛을 가진
    나머지 설비에 일괄 동기화한다. 이름이 같은 부품은 규격/교체주기/비고/메모/도면/재고수량/구매처/
    리드타임/아이콘/위치/크기가 TEAG01호기 기준으로 갱신되고(재고는 전 설비가 동일하게 관리됨),
    새 부품은 추가되며, 여기 없는 이름의 부품은 삭제된다(교체 이력도 함께 삭제)."""
    master_unit = conn.execute("SELECT * FROM units WHERE id = ?", (unit_id,)).fetchone()
    master_parts = conn.execute(
        "SELECT * FROM parts WHERE unit_id = ? AND deleted_at IS NULL ORDER BY id", (unit_id,)
    ).fetchall()
    master_names = {p["name"] for p in master_parts}

    target_units = conn.execute(
        "SELECT id FROM units WHERE name = ? AND equipment_id != ? AND deleted_at IS NULL",
        (master_unit["name"], MASTER_EQUIPMENT_ID),
    ).fetchall()

    for t in target_units:
        existing = {
            p["name"]: p
            for p in conn.execute("SELECT * FROM parts WHERE unit_id = ? AND deleted_at IS NULL", (t["id"],)).fetchall()
        }
        for mp in master_parts:
            if mp["name"] in existing:
                ep = existing[mp["name"]]
                conn.execute(
                    """UPDATE parts SET spec = ?, cycle_days = ?, cycle_unit = ?, cost = ?, note = ?, memo = ?,
                       drawing_data = ?, icon = ?, pos_x = ?, pos_y = ?, width = ?, height = ?,
                       stock_qty = ?, supplier = ?, supplier_contact = ?, lead_time_days = ? WHERE id = ?""",
                    (
                        mp["spec"], mp["cycle_days"], mp["cycle_unit"], mp["cost"], mp["note"], mp["memo"],
                        mp["drawing_data"], mp["icon"], mp["pos_x"], mp["pos_y"], mp["width"], mp["height"],
                        mp["stock_qty"], mp["supplier"], mp["supplier_contact"], mp["lead_time_days"], ep["id"],
                    ),
                )
            else:
                conn.execute(
                    """INSERT INTO parts (unit_id, name, spec, cycle_days, cycle_unit, cost, note, memo, drawing_data,
                       icon, pos_x, pos_y, width, height, stock_qty, supplier, supplier_contact, lead_time_days)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        t["id"], mp["name"], mp["spec"], mp["cycle_days"], mp["cycle_unit"], mp["cost"],
                        mp["note"], mp["memo"], mp["drawing_data"], mp["icon"], mp["pos_x"], mp["pos_y"], mp["width"], mp["height"],
                        mp["stock_qty"], mp["supplier"], mp["supplier_contact"], mp["lead_time_days"],
                    ),
                )
        for name, ep in existing.items():
            if name not in master_names:
                conn.execute("DELETE FROM parts WHERE id = ?", (ep["id"],))

    return len(target_units), len(master_parts)


def units_with_status_bulk(conn, units):
    """유닛 여러 개의 상태/부품수를 부품 테이블 조회 1번으로 한꺼번에 계산한다
    (유닛마다 따로 쿼리를 날리는 N+1 패턴을 피하기 위함)."""
    if not units:
        return []
    unit_ids = [u["id"] for u in units]
    placeholders = ",".join("?" for _ in unit_ids)
    all_parts = conn.execute(
        f"SELECT * FROM parts WHERE unit_id IN ({placeholders}) AND deleted_at IS NULL", unit_ids
    ).fetchall()
    parts_by_unit = {}
    for p in all_parts:
        parts_by_unit.setdefault(p["unit_id"], []).append(p)

    result = []
    for u in units:
        parts = parts_by_unit.get(u["id"], [])
        statuses = [part_status(p["cycle_days"], p["last_replaced_date"])["status"] for p in parts]
        if "overdue" in statuses:
            overall = "overdue"
        elif "soon" in statuses:
            overall = "soon"
        elif "unknown" in statuses:
            overall = "unknown" if not any(s == "ok" for s in statuses) else "ok"
        elif statuses:
            overall = "ok"
        else:
            overall = "empty"
        d = dict(u)
        d["part_count"] = len(parts)
        d["overall_status"] = overall
        d["overdue_count"] = statuses.count("overdue")
        d["soon_count"] = statuses.count("soon")
        result.append(d)
    return result


def unit_with_status(conn, u):
    return units_with_status_bulk(conn, [u])[0]


STATUS_PRIORITY = ["overdue", "soon", "unknown", "ok", "empty"]


def calc_setup_runtime(setup_date):
    """SETUP 일자로부터 오늘까지 경과한 기간을 "N년 M개월" 형식으로 계산"""
    if not setup_date:
        return None
    setup = datetime.strptime(setup_date, "%Y-%m-%d").date()
    today = date.today()
    if setup > today:
        return None
    years = today.year - setup.year
    months = today.month - setup.month
    if today.day < setup.day:
        months -= 1
    if months < 0:
        years -= 1
        months += 12
    return f"{years}년 {months}개월"


def equipments_with_status_bulk(conn, equipments):
    """설비 여러 개의 상태/유닛수를 유닛+부품 조회 2번으로 한꺼번에 계산한다
    (설비마다, 유닛마다 따로 쿼리를 날리는 N+1 패턴을 피하기 위함)."""
    if not equipments:
        return []
    eq_ids = [e["id"] for e in equipments]
    placeholders = ",".join("?" for _ in eq_ids)
    all_units = conn.execute(
        f"SELECT * FROM units WHERE equipment_id IN ({placeholders}) AND deleted_at IS NULL", eq_ids
    ).fetchall()
    unit_statuses_all = units_with_status_bulk(conn, all_units)
    units_by_equipment = {}
    for us in unit_statuses_all:
        units_by_equipment.setdefault(us["equipment_id"], []).append(us)

    result = []
    for e in equipments:
        unit_statuses = units_by_equipment.get(e["id"], [])
        statuses = [us["overall_status"] for us in unit_statuses]
        overall = next((s for s in STATUS_PRIORITY if s in statuses), "empty")
        d = dict(e)
        d["unit_count"] = len(unit_statuses)
        d["overall_status"] = overall
        d["overdue_count"] = sum(us["overdue_count"] for us in unit_statuses)
        d["soon_count"] = sum(us["soon_count"] for us in unit_statuses)
        d["runtime_display"] = calc_setup_runtime(e["setup_date"])
        result.append(d)
    return result


def equipment_with_status(conn, e):
    return equipments_with_status_bulk(conn, [e])[0]


def seed_default_units_for_equipment(conn, equipment_id):
    templates = conn.execute("SELECT * FROM unit_templates ORDER BY id").fetchall()
    for t in templates:
        conn.execute(
            "INSERT INTO units (equipment_id, name, icon, color, pos_x, pos_y, width, height) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (equipment_id, t["name"], t["icon"], t["color"], t["pos_x"], t["pos_y"], t["width"], t["height"]),
        )


class _TableGridParser(HTMLParser):
    """rollout_items.data_html 의 첫 번째 표를 병합 셀(rowspan/colspan)까지
    반영한 2차원 격자로 펼친다 (화면의 게이지 집계 로직과 동일한 규칙)."""
    def __init__(self):
        super().__init__()
        self.grid = []
        self.row = -1
        self.col = 0
        self.in_cell = False
        self.cell_text = ""
        self.cell_span = (1, 1)
        self.cell_pos = (0, 0)
        self.table_depth = 0
        self.done_first_table = False

    def handle_starttag(self, tag, attrs):
        if self.done_first_table:
            return
        if tag == "table":
            self.table_depth += 1
            return
        if self.table_depth == 0:
            return
        if tag == "tr":
            self.row += 1
            self.col = 0
            while len(self.grid) <= self.row:
                self.grid.append({})
        elif tag in ("td", "th"):
            attrs = dict(attrs)
            colspan = int(attrs.get("colspan") or 1)
            rowspan = int(attrs.get("rowspan") or 1)
            while self.col in self.grid[self.row]:
                self.col += 1
            self.in_cell = True
            self.cell_text = ""
            self.cell_span = (rowspan, colspan)
            self.cell_pos = (self.row, self.col)

    def handle_endtag(self, tag):
        if self.done_first_table:
            return
        if tag == "table" and self.table_depth:
            self.table_depth -= 1
            if self.table_depth == 0:
                self.done_first_table = True
        elif tag in ("td", "th") and self.in_cell:
            r0, c0 = self.cell_pos
            rowspan, colspan = self.cell_span
            text = self.cell_text.strip()
            for dr in range(rowspan):
                while len(self.grid) <= r0 + dr:
                    self.grid.append({})
                for dc in range(colspan):
                    self.grid[r0 + dr][c0 + dc] = text
            self.col = c0 + colspan
            self.in_cell = False

    def handle_data(self, data):
        if self.in_cell:
            self.cell_text += data


def rollout_progress_from_html(data_html):
    """CH 이름(2열)이 있는 행만 대상으로 진행 날짜(3열) 기입 여부를 센다."""
    parser = _TableGridParser()
    parser.feed(data_html or "")
    done = pending = 0
    for r in range(1, len(parser.grid)):
        row = parser.grid[r]
        if not (row.get(1) or "").strip():
            continue
        if (row.get(2) or "").strip():
            done += 1
        else:
            pending += 1
    return done, pending


def build_mail_report_html():
    """메일 본문 HTML을 만든다.
    1번째: 사이트 접속 URL
    2번째: 횡전개 현황 (항목별 완료/미진행/전체/진행률)
    3번째: 이벤트 알림 - 재고 1개 이하 발생 리스트, 전날(발송 기준) 교체 기록 리스트"""
    conn = get_db()
    items = conn.execute("SELECT * FROM rollout_items ORDER BY id").fetchall()
    low_stock = [p for p in get_inventory_rows(conn) if (p["stock_qty"] or 0) <= 1]
    yesterday = (date.today() - timedelta(days=1)).isoformat()
    yesterday_history = conn.execute("""
        SELECT h.cost, p.name AS part_name, p.spec,
               u.name AS unit_name, e.name AS equipment_name
        FROM replacement_history h
        JOIN parts p ON h.part_id = p.id
        JOIN units u ON p.unit_id = u.id
        JOIN equipments e ON u.equipment_id = e.id
        WHERE h.replaced_date = ?
        ORDER BY e.id, u.id, p.id
    """, (yesterday,)).fetchall()
    conn.close()

    td = "border:1px solid #ccc;padding:6px 12px"
    site_url = f"http://{get_lan_ip()}:5000"

    rollout_rows = ""
    for it in items:
        done, pending = rollout_progress_from_html(it["data_html"])
        total = done + pending
        pct = round(done / total * 100) if total else 0
        rollout_rows += (
            f"<tr><td style='{td}'>{html_lib.escape(it['title'])}</td>"
            f"<td style='{td};text-align:center'>{done}</td>"
            f"<td style='{td};text-align:center'>{pending}</td>"
            f"<td style='{td};text-align:center'>{total}</td>"
            f"<td style='{td};text-align:center;font-weight:bold'>{pct}%</td></tr>"
        )
    if not rollout_rows:
        rollout_rows = f"<tr><td colspan='5' style='{td}'>등록된 횡전개 항목이 없습니다.</td></tr>"

    low_stock_rows = ""
    for p in low_stock:
        low_stock_rows += (
            f"<tr><td style='{td}'>{html_lib.escape(p['name'])}</td>"
            f"<td style='{td}'>{html_lib.escape(p['spec'] or '')}</td>"
            f"<td style='{td}'>{html_lib.escape(p['unit_name'])}</td>"
            f"<td style='{td};text-align:center;color:#c0392b;font-weight:bold'>{p['stock_qty'] or 0}</td></tr>"
        )
    if not low_stock_rows:
        low_stock_rows = f"<tr><td colspan='4' style='{td}'>재고 1개 이하 부품이 없습니다.</td></tr>"

    history_rows = ""
    for h in yesterday_history:
        history_rows += (
            f"<tr><td style='{td}'>{html_lib.escape(h['equipment_name'])}</td>"
            f"<td style='{td}'>{html_lib.escape(h['unit_name'])}</td>"
            f"<td style='{td}'>{html_lib.escape(h['part_name'])}</td>"
            f"<td style='{td}'>{html_lib.escape(h['spec'] or '')}</td>"
            f"<td style='{td};text-align:right'>{h['cost'] or 0:,.0f}원</td></tr>"
        )
    if not history_rows:
        history_rows = f"<tr><td colspan='5' style='{td}'>{yesterday} 교체 기록이 없습니다.</td></tr>"

    return (
        f"<p style='font-size:14px'>사이트 접속: <a href='{site_url}'>{site_url}</a></p>"
        f"<h3>횡전개 현황</h3>"
        f"<table style='border-collapse:collapse;font-size:14px;margin-bottom:20px'>"
        f"<tr style='background:#f3f4f6'>"
        f"<th style='{td}'>횡전개 항목</th><th style='{td}'>완료</th>"
        f"<th style='{td}'>미진행</th><th style='{td}'>전체</th><th style='{td}'>진행률</th></tr>"
        f"{rollout_rows}</table>"
        f"<h3>이벤트 알림</h3>"
        f"<p style='font-size:14px;margin-bottom:4px'><b>재고 1개 이하 발생</b></p>"
        f"<table style='border-collapse:collapse;font-size:14px;margin-bottom:16px'>"
        f"<tr style='background:#f3f4f6'>"
        f"<th style='{td}'>부품명</th><th style='{td}'>규격</th><th style='{td}'>소속 유닛</th><th style='{td}'>재고</th></tr>"
        f"{low_stock_rows}</table>"
        f"<p style='font-size:14px;margin-bottom:4px'><b>전날({yesterday}) 교체 기록</b></p>"
        f"<table style='border-collapse:collapse;font-size:14px'>"
        f"<tr style='background:#f3f4f6'>"
        f"<th style='{td}'>설비</th><th style='{td}'>유닛</th><th style='{td}'>부품</th>"
        f"<th style='{td}'>규격</th><th style='{td}'>금액</th></tr>"
        f"{history_rows}</table>"
    )


def send_status_mail():
    """사내 메일 API로 횡전개 현황 리포트를 발송한다. (성공 여부, 메시지) 반환."""
    conn = get_db()
    receivers = [
        r["email_id"]
        for r in conn.execute("SELECT email_id FROM mail_recipients ORDER BY id").fetchall()
    ]
    conn.close()
    if not receivers:
        return False, "수신자가 등록되어 있지 않습니다. 메일 설정에서 수신자를 먼저 추가하세요."
    if "XXXX" in MAIL_AUTHORIZATION or not MAIL_AUTHORIZATION:
        return False, "메일 API 키가 설정되지 않았습니다. app.py 상단의 MAIL_AUTHORIZATION / MAIL_SYSTEM_ID 를 확인하세요."
    api = "https://openapi.samsung.net/mail/api/v2.0/mails/send?userId=" + MAIL_SENDER_ID
    header = {"Authorization": MAIL_AUTHORIZATION, "System-ID": MAIL_SYSTEM_ID}
    body = {
        "subject": MAIL_SUBJECT,
        "contents": build_mail_report_html(),
        "contentType": "HTML",
        "docSecuType": "PERSONAL",
        "sender": {"emailAddress": f"{MAIL_SENDER_ID}@samsung.com"},
        "recipients": [
            {"emailAddress": f"{r}@samsung.com", "recipientType": "TO"} for r in receivers
        ],
    }
    try:
        res = requests.post(api, headers=header, json=body, timeout=15)
        if res.status_code // 100 == 2:
            return True, f"수신자 {len(receivers)}명에게 발송 완료"
        return False, f"메일 API 오류 (HTTP {res.status_code}): {res.text[:200]}"
    except Exception as e:
        return False, f"메일 발송 실패: {e}"


def get_alert_parts():
    """모든 설비를 통틀어 교체 필요/임박 상태인 부품 목록 (경과가 급한 순)"""
    conn = get_db()
    rows = conn.execute("""
        SELECT p.*, u.id AS unit_id, u.name AS unit_name,
               e.id AS equipment_id, e.name AS equipment_name, e.icon AS equipment_icon
        FROM parts p
        JOIN units u ON p.unit_id = u.id
        JOIN equipments e ON u.equipment_id = e.id
        WHERE p.deleted_at IS NULL AND u.deleted_at IS NULL AND e.deleted_at IS NULL
    """).fetchall()
    conn.close()
    result = []
    for r in rows:
        info = part_status(r["cycle_days"], r["last_replaced_date"])
        if info["status"] in ("overdue", "soon"):
            d = dict(r)
            d.update(info)
            result.append(d)
    result.sort(key=lambda x: x["days_left"])
    return result


def search_parts(query, status=None, equipment_id=None):
    """부품명/규격으로 모든 설비를 통틀어 검색 (상태/설비로 추가 필터링 가능)"""
    conn = get_db()
    like = f"%{query}%"
    sql = """
        SELECT p.*, u.id AS unit_id, u.name AS unit_name,
               e.id AS equipment_id, e.name AS equipment_name, e.icon AS equipment_icon
        FROM parts p
        JOIN units u ON p.unit_id = u.id
        JOIN equipments e ON u.equipment_id = e.id
        WHERE (p.name LIKE ? OR p.spec LIKE ?)
          AND p.deleted_at IS NULL AND u.deleted_at IS NULL AND e.deleted_at IS NULL
    """
    params = [like, like]
    if equipment_id:
        sql += " AND e.id = ?"
        params.append(equipment_id)
    sql += " ORDER BY e.id, u.id, p.id"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    result = []
    for r in rows:
        info = part_status(r["cycle_days"], r["last_replaced_date"])
        if status and info["status"] != status:
            continue
        d = dict(r)
        d.update(info)
        result.append(d)
    return result


def get_part_spec_stats(conn, unit_names=None, start_date=None, end_date=None):
    """부품명+규격을 기준으로 시스템 전체(선택된 유닛 이름으로 범위 제한 가능)를 집계한다.
    금액은 실제 교체 이력(replacement_history)이 있으면 그 금액을 사용하고, 기간 필터가
    없는데 교체 이력이 아직 없는 부품은 등록된 금액(예상 비용)을 대신 사용한다(설치만 해두고
    아직 한 번도 교체하지 않은 부품이 금액순 집계에서 통째로 사라지는 것을 막기 위함). 이때
    부품은 전 설비에 동일하게 동기화되어 있으므로, 등록된 금액은 유닛별로 합산하지 않고
    기준 설비(TEAG01호기)에 등록된 값 하나만 사용한다.
    기간 필터가 있으면 해당 기간에 실제로 발생한 교체 기록의 금액만 집계한다.
    사용량은 항상 실제 교체 이력 기준이다.
    교체주기는 항상 현재 부품 구성 기준으로 계산하되, 주기가 없는(N/A) 부품은 제외한다."""
    period_filter = bool(start_date or end_date)
    query = """
        SELECT p.*, u.name AS unit_name, e.id AS equipment_id
        FROM parts p
        JOIN units u ON p.unit_id = u.id
        JOIN equipments e ON u.equipment_id = e.id
        WHERE p.deleted_at IS NULL AND u.deleted_at IS NULL AND e.deleted_at IS NULL
    """
    params = []
    if unit_names:
        placeholders = ",".join("?" for _ in unit_names)
        query += f" AND u.name IN ({placeholders})"
        params = list(unit_names)
    parts = conn.execute(query, params).fetchall()

    # 그룹마다 따로 교체 이력을 조회하는 대신, 관련된 모든 부품의 이력을 한 번에 가져와
    # part_id 기준으로 집계해둔다 (N+1 쿼리 방지).
    hist_by_part = {}
    part_ids = [p["id"] for p in parts]
    if part_ids:
        placeholders = ",".join("?" for _ in part_ids)
        hist_query = f"SELECT part_id, cost FROM replacement_history WHERE part_id IN ({placeholders})"
        hist_params = list(part_ids)
        if start_date:
            hist_query += " AND replaced_date >= ?"
            hist_params.append(start_date)
        if end_date:
            hist_query += " AND replaced_date <= ?"
            hist_params.append(end_date)
        for h in conn.execute(hist_query, hist_params).fetchall():
            entry = hist_by_part.setdefault(h["part_id"], {"n": 0, "total": 0})
            entry["n"] += 1
            entry["total"] += h["cost"] or 0

    groups = {}
    for p in parts:
        key = (p["name"], p["spec"] or "")
        g = groups.get(key)
        if g is None:
            g = {
                "name": p["name"],
                "spec": p["spec"] or "",
                "total_cost": 0,
                "instance_count": 0,
                "usage_count": 0,
                "min_cycle_days": None,
                "min_cycle_unit": "일",
            }
            groups[key] = g
        g["instance_count"] += 1
        if p["cycle_days"] is not None and (g["min_cycle_days"] is None or p["cycle_days"] < g["min_cycle_days"]):
            g["min_cycle_days"] = p["cycle_days"]
            g["min_cycle_unit"] = p["cycle_unit"]
        hist = hist_by_part.get(p["id"], {"n": 0, "total": 0})
        g["usage_count"] += hist["n"]
        if hist["n"] > 0:
            g["total_cost"] += hist["total"]
        elif not period_filter and p["equipment_id"] == MASTER_EQUIPMENT_ID:
            g["total_cost"] += p["cost"] or 0

    return list(groups.values())


def csv_response(filename, header, rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(rows)
    data = "﻿" + buf.getvalue()
    encoded_name = quote(filename)
    return Response(
        data,
        mimetype="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename=export.csv; filename*=UTF-8''{encoded_name}"
        },
    )


LOGIN_EXEMPT_PREFIXES = ("/static/", "/login")


@app.before_request
def require_login():
    if request.path.startswith(LOGIN_EXEMPT_PREFIXES):
        return None
    if session.get("authenticated") and session.get("user_name"):
        return None
    if request.path.startswith("/api/"):
        return jsonify({"error": "로그인이 필요합니다"}), 401
    return redirect(url_for("login_page", next=request.path))


@app.route("/login", methods=["GET", "POST"])
def login_page():
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        password = request.form.get("password", "")
        if not name:
            return redirect(url_for("login_page", error="2"))
        conn = get_db()
        password_hash = get_config(conn, "password_hash")
        conn.close()
        if password_hash and check_password_hash(password_hash, password):
            session["authenticated"] = True
            session["user_name"] = name
            session.permanent = True
            next_url = request.args.get("next") or url_for("dashboard")
            return redirect(next_url)
        return redirect(url_for("login_page", error="1"))
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login_page"))


@app.route("/api/auth/change-password", methods=["POST"])
def change_password():
    data = request.get_json()
    current_password = data.get("current_password") or ""
    new_password = data.get("new_password") or ""
    if len(new_password) < 4:
        return jsonify({"error": "비밀번호는 4자 이상으로 설정하세요"}), 400
    conn = get_db()
    password_hash = get_config(conn, "password_hash")
    if not password_hash or not check_password_hash(password_hash, current_password):
        conn.close()
        return jsonify({"error": "현재 비밀번호가 올바르지 않습니다"}), 400
    set_config(conn, "password_hash", generate_password_hash(new_password))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/")
def dashboard():
    return render_template("dashboard.html", user_name=session.get("user_name", ""))


@app.route("/alerts")
def alerts_page():
    return render_template("alerts.html")


@app.route("/search")
def search_page():
    return render_template("search.html")


@app.route("/stats")
def stats_page():
    return render_template("stats.html")


@app.route("/bulk-add-parts")
def bulk_add_parts_page():
    return render_template("bulk_add_parts.html")


@app.route("/inventory")
def inventory_page():
    return render_template("inventory.html")


def get_inventory_rows(conn):
    """재고는 TEAG01호기(기준 설비) 기준으로 전 설비가 동일하게 관리되므로,
    나머지 설비는 동일한 내용이 중복되어 나타나는 것을 막기 위해 TEAG01호기 것만 보여준다.
    규격이 같으면(부품 이름이 달라도) 나란히 묶여 보이도록 규격 기준으로 먼저 정렬하고,
    같은 규격 안에서는 소속 유닛 기준으로 정렬한다."""
    return conn.execute("""
        SELECT p.id, p.name, p.spec, p.stock_qty, p.supplier, p.supplier_contact, p.lead_time_days,
               u.id AS unit_id, u.name AS unit_name
        FROM parts p
        JOIN units u ON p.unit_id = u.id
        JOIN equipments e ON u.equipment_id = e.id
        WHERE p.deleted_at IS NULL AND u.deleted_at IS NULL AND e.deleted_at IS NULL
          AND e.id = ?
        ORDER BY p.spec ASC, u.name ASC, p.name ASC
    """, (MASTER_EQUIPMENT_ID,)).fetchall()


@app.route("/api/inventory")
def api_inventory():
    conn = get_db()
    rows = get_inventory_rows(conn)
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route("/api/inventory/export")
def export_inventory_csv():
    low_only = request.args.get("low_only") == "1"
    conn = get_db()
    rows = get_inventory_rows(conn)
    conn.close()
    header = ["규격", "부품이름", "소속 유닛", "재고 수량", "구매처", "연락처", "리드타임(일)"]
    data_rows = []
    for p in rows:
        if low_only and (p["stock_qty"] or 0) > 0:
            continue
        data_rows.append([
            p["spec"] or "", p["name"], p["unit_name"],
            p["stock_qty"] or 0, p["supplier"] or "", p["supplier_contact"] or "",
            p["lead_time_days"] if p["lead_time_days"] is not None else "",
        ])
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return csv_response(f"재고관리_{timestamp}.csv", header, data_rows)


@app.route("/api/mail/recipients")
def list_mail_recipients():
    conn = get_db()
    rows = conn.execute("SELECT * FROM mail_recipients ORDER BY id").fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route("/api/mail/recipients", methods=["POST"])
def add_mail_recipient():
    data = request.get_json()
    email_id = (data.get("email_id") or "").strip().replace("@samsung.com", "")
    if not email_id:
        return jsonify({"error": "수신자 녹스 ID를 입력하세요"}), 400
    conn = get_db()
    try:
        conn.execute("INSERT INTO mail_recipients (email_id) VALUES (?)", (email_id,))
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "이미 등록된 수신자입니다"}), 409
    conn.close()
    return jsonify({"ok": True}), 201


@app.route("/api/mail/recipients/<int:rid>", methods=["DELETE"])
def delete_mail_recipient(rid):
    conn = get_db()
    conn.execute("DELETE FROM mail_recipients WHERE id = ?", (rid,))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/mail/send", methods=["POST"])
def api_send_mail():
    ok, msg = send_status_mail()
    if ok:
        conn = get_db()
        log_activity(conn, "mail", "mail", None, MAIL_SUBJECT, msg)
        conn.commit()
        conn.close()
    return jsonify({"ok": ok, "message": msg}), (200 if ok else 500)


@app.route("/api/mail/schedule", methods=["GET", "PUT"])
def api_mail_schedule():
    conn = get_db()
    if request.method == "PUT":
        t = (request.get_json().get("time") or "").strip()  # "HH:MM" 또는 "" (해제)
        set_config(conn, "mail_schedule_time", t)
        conn.commit()
    t = get_config(conn, "mail_schedule_time") or ""
    conn.close()
    return jsonify({"time": t})


@app.route("/rollout")
def rollout_page():
    return render_template("rollout.html")


@app.route("/api/rollout")
def list_rollout_items():
    conn = get_db()
    rows = conn.execute("SELECT * FROM rollout_items ORDER BY id").fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route("/api/rollout", methods=["POST"])
def add_rollout_item():
    data = request.get_json()
    title = (data.get("title") or "").strip()
    if not title:
        return jsonify({"error": "제목을 입력하세요"}), 400
    data_html = sanitize_rich_html(data.get("data_html") or "")
    conn = get_db()
    cur = conn.execute("INSERT INTO rollout_items (title, data_html) VALUES (?, ?)", (title, data_html))
    new_id = cur.lastrowid
    log_activity(conn, "create", "rollout", new_id, title, "횡전개 항목 추가")
    conn.commit()
    row = conn.execute("SELECT * FROM rollout_items WHERE id = ?", (new_id,)).fetchone()
    conn.close()
    return jsonify(dict(row)), 201


@app.route("/api/rollout/<int:item_id>", methods=["PUT"])
def update_rollout_item(item_id):
    data = request.get_json()
    conn = get_db()
    row = conn.execute("SELECT * FROM rollout_items WHERE id = ?", (item_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "항목을 찾을 수 없습니다"}), 404
    title = (data.get("title") or row["title"]).strip()
    data_html = sanitize_rich_html(data.get("data_html", row["data_html"]))
    conn.execute(
        "UPDATE rollout_items SET title = ?, data_html = ?, updated_at = datetime('now','localtime') WHERE id = ?",
        (title, data_html, item_id),
    )
    log_activity(conn, "update", "rollout", item_id, title, "횡전개 항목 수정")
    conn.commit()
    updated = conn.execute("SELECT * FROM rollout_items WHERE id = ?", (item_id,)).fetchone()
    conn.close()
    return jsonify(dict(updated))


@app.route("/api/rollout/<int:item_id>", methods=["DELETE"])
def delete_rollout_item(item_id):
    conn = get_db()
    row = conn.execute("SELECT * FROM rollout_items WHERE id = ?", (item_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "항목을 찾을 수 없습니다"}), 404
    conn.execute("DELETE FROM rollout_items WHERE id = ?", (item_id,))
    log_activity(conn, "delete", "rollout", item_id, row["title"], "횡전개 항목 삭제")
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/equipment/<int:equipment_id>")
def equipment_page(equipment_id):
    conn = get_db()
    equipment = conn.execute(
        "SELECT id FROM equipments WHERE id = ? AND deleted_at IS NULL", (equipment_id,)
    ).fetchone()
    conn.close()
    if not equipment:
        return "설비를 찾을 수 없습니다", 404
    return render_template("equipment.html", equipment_id=equipment_id)


@app.route("/config")
def unit_template_config():
    return render_template("config.html")


@app.route("/unit/<int:unit_id>")
def unit_detail(unit_id):
    conn = get_db()
    unit = conn.execute(
        "SELECT id FROM units WHERE id = ? AND deleted_at IS NULL", (unit_id,)
    ).fetchone()
    conn.close()
    if not unit:
        return "유닛을 찾을 수 없습니다", 404
    return render_template("unit.html", unit_id=unit_id)


@app.route("/api/equipments")
def list_equipments():
    conn = get_db()
    equipments = conn.execute("SELECT * FROM equipments WHERE deleted_at IS NULL ORDER BY id").fetchall()
    result = equipments_with_status_bulk(conn, equipments)
    conn.close()
    return jsonify(result)


@app.route("/api/equipments", methods=["POST"])
def add_equipment():
    data = request.get_json()
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "설비 이름을 입력하세요"}), 400
    icon = (data.get("icon") or "🏭").strip()
    conn = get_db()
    try:
        count = conn.execute("SELECT COUNT(*) AS n FROM equipments").fetchone()["n"]
        pos_x, pos_y = dashboard_grid_pos(count)
        cur = conn.execute(
            "INSERT INTO equipments (name, icon, pos_x, pos_y) VALUES (?, ?, ?, ?)",
            (name, icon, pos_x, pos_y),
        )
        new_id = cur.lastrowid
        conn.execute(
            "INSERT INTO equipment_notes (equipment_id, content) VALUES (?, '')", (new_id,)
        )
        seed_default_units_for_equipment(conn, new_id)
        log_activity(conn, "create", "equipment", new_id, name)
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "이미 사용 중인 설비 이름입니다"}), 409
    equipment = conn.execute("SELECT * FROM equipments WHERE id = ?", (new_id,)).fetchone()
    d = equipment_with_status(conn, equipment)
    conn.close()
    return jsonify(d), 201


@app.route("/api/equipments/<int:equipment_id>")
def get_equipment(equipment_id):
    conn = get_db()
    equipment = conn.execute(
        "SELECT * FROM equipments WHERE id = ? AND deleted_at IS NULL", (equipment_id,)
    ).fetchone()
    if not equipment:
        conn.close()
        return jsonify({"error": "설비를 찾을 수 없습니다"}), 404
    d = equipment_with_status(conn, equipment)
    conn.close()
    return jsonify(d)


@app.route("/api/equipments/<int:equipment_id>", methods=["PUT"])
def update_equipment(equipment_id):
    data = request.get_json()
    conn = get_db()
    equipment = conn.execute("SELECT * FROM equipments WHERE id = ?", (equipment_id,)).fetchone()
    if not equipment:
        conn.close()
        return jsonify({"error": "설비를 찾을 수 없습니다"}), 404
    name = (data.get("name") or equipment["name"]).strip()
    icon = (data.get("icon") or equipment["icon"]).strip()
    pos_x = data.get("pos_x", equipment["pos_x"])
    pos_y = data.get("pos_y", equipment["pos_y"])
    location = data.get("location", equipment["location"])
    setup_date = data.get("setup_date", equipment["setup_date"]) or None
    meaningful_change = (
        name != equipment["name"] or icon != equipment["icon"]
        or location != equipment["location"] or setup_date != equipment["setup_date"]
    )
    try:
        conn.execute(
            "UPDATE equipments SET name = ?, icon = ?, pos_x = ?, pos_y = ?, location = ?, setup_date = ? WHERE id = ?",
            (name, icon, pos_x, pos_y, location, setup_date, equipment_id),
        )
        if meaningful_change:
            log_activity(conn, "update", "equipment", equipment_id, name, "설비 정보 수정")
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "이미 사용 중인 설비 이름입니다"}), 409
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/equipments/<int:equipment_id>", methods=["DELETE"])
def delete_equipment(equipment_id):
    conn = get_db()
    equipment = conn.execute("SELECT * FROM equipments WHERE id = ?", (equipment_id,)).fetchone()
    if not equipment:
        conn.close()
        return jsonify({"error": "설비를 찾을 수 없습니다"}), 404
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn.execute("UPDATE equipments SET deleted_at = ? WHERE id = ?", (now, equipment_id))
    unit_ids = [
        r["id"] for r in conn.execute(
            "SELECT id FROM units WHERE equipment_id = ? AND deleted_at IS NULL", (equipment_id,)
        ).fetchall()
    ]
    if unit_ids:
        placeholders = ",".join("?" for _ in unit_ids)
        conn.execute(f"UPDATE units SET deleted_at = ? WHERE id IN ({placeholders})", (now, *unit_ids))
        conn.execute(f"UPDATE parts SET deleted_at = ? WHERE unit_id IN ({placeholders})", (now, *unit_ids))
    log_activity(conn, "delete", "equipment", equipment_id, equipment["name"], "휴지통으로 이동")
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/equipments/<int:equipment_id>/units")
def list_units(equipment_id):
    conn = get_db()
    units = conn.execute(
        "SELECT * FROM units WHERE equipment_id = ? AND deleted_at IS NULL ORDER BY id", (equipment_id,)
    ).fetchall()
    result = units_with_status_bulk(conn, units)
    conn.close()
    return jsonify(result)


@app.route("/api/equipments/<int:equipment_id>/units", methods=["POST"])
def add_unit(equipment_id):
    data = request.get_json()
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "유닛 이름을 입력하세요"}), 400
    icon = (data.get("icon") or "⚙️").strip()
    color = (data.get("color") or "#1a3a5c").strip()
    pos_x = data.get("pos_x")
    pos_y = data.get("pos_y")
    if pos_x is None or pos_y is None:
        pos_x, pos_y = random.uniform(15, 85), random.uniform(20, 80)
    width = data.get("width") or 140
    height = data.get("height") or 110
    drawing_data = data.get("drawing_data") or None
    if drawing_data and len(drawing_data) > MAX_DRAWING_DATA_LEN:
        return jsonify({"error": "도면 이미지 용량이 너무 큽니다 (최대 5MB)"}), 400
    conn = get_db()
    cur = conn.execute(
        "INSERT INTO units (equipment_id, name, icon, color, pos_x, pos_y, width, height, drawing_data) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (equipment_id, name, icon, color, pos_x, pos_y, width, height, drawing_data),
    )
    conn.commit()
    new_id = cur.lastrowid
    unit = conn.execute("SELECT * FROM units WHERE id = ?", (new_id,)).fetchone()
    log_activity(conn, "create", "unit", new_id, name)
    conn.commit()
    conn.close()
    d = dict(unit)
    d["part_count"] = 0
    d["overall_status"] = "empty"
    return jsonify(d), 201


@app.route("/api/units/<int:unit_id>")
def get_unit(unit_id):
    conn = get_db()
    unit = conn.execute(
        "SELECT * FROM units WHERE id = ? AND deleted_at IS NULL", (unit_id,)
    ).fetchone()
    if not unit:
        conn.close()
        return jsonify({"error": "유닛을 찾을 수 없습니다"}), 404
    d = unit_with_status(conn, unit)
    conn.close()
    return jsonify(d)


@app.route("/api/units/<int:unit_id>", methods=["PUT"])
def update_unit(unit_id):
    data = request.get_json()
    conn = get_db()
    unit = conn.execute("SELECT * FROM units WHERE id = ?", (unit_id,)).fetchone()
    if not unit:
        conn.close()
        return jsonify({"error": "유닛을 찾을 수 없습니다"}), 404
    name = (data.get("name") or unit["name"]).strip()
    icon = (data.get("icon") or unit["icon"]).strip()
    color = (data.get("color") or unit["color"]).strip()
    pos_x = data.get("pos_x", unit["pos_x"])
    pos_y = data.get("pos_y", unit["pos_y"])
    width = data.get("width", unit["width"])
    height = data.get("height", unit["height"])
    drawing_data = data.get("drawing_data", unit["drawing_data"])
    if drawing_data and len(drawing_data) > MAX_DRAWING_DATA_LEN:
        conn.close()
        return jsonify({"error": "도면 이미지 용량이 너무 큽니다 (최대 5MB)"}), 400
    meaningful_change = name != unit["name"] or icon != unit["icon"] or color != unit["color"]
    conn.execute(
        "UPDATE units SET name = ?, icon = ?, color = ?, pos_x = ?, pos_y = ?, width = ?, height = ?, "
        "drawing_data = ? WHERE id = ?",
        (name, icon, color, pos_x, pos_y, width, height, drawing_data, unit_id),
    )
    if meaningful_change:
        log_activity(conn, "update", "unit", unit_id, name, "유닛 정보 수정")
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/units/<int:unit_id>", methods=["DELETE"])
def delete_unit(unit_id):
    conn = get_db()
    unit = conn.execute("SELECT * FROM units WHERE id = ?", (unit_id,)).fetchone()
    if not unit:
        conn.close()
        return jsonify({"error": "유닛을 찾을 수 없습니다"}), 404
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn.execute("UPDATE units SET deleted_at = ? WHERE id = ?", (now, unit_id))
    conn.execute("UPDATE parts SET deleted_at = ? WHERE unit_id = ? AND deleted_at IS NULL", (now, unit_id))
    log_activity(conn, "delete", "unit", unit_id, unit["name"], "휴지통으로 이동")
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/units/<int:unit_id>/parts")
def list_parts(unit_id):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM parts WHERE unit_id = ? AND deleted_at IS NULL ORDER BY id", (unit_id,)
    ).fetchall()
    conn.close()
    return jsonify([serialize_part(r) for r in rows])


@app.route("/api/units/<int:unit_id>/parts", methods=["POST"])
def add_part(unit_id):
    data = request.get_json()
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "부품 이름을 입력하세요"}), 400
    spec = (data.get("spec") or "").strip()
    cycle_days, cycle_unit = resolve_cycle(data, None, "N/A")
    cost = float(data.get("cost") or 0)
    last_replaced_date = data.get("last_replaced_date") or None
    note = (data.get("note") or "").strip()
    memo = sanitize_rich_html((data.get("memo") or "").strip())
    drawing_data = data.get("drawing_data") or None
    if drawing_data and len(drawing_data) > MAX_DRAWING_DATA_LEN:
        return jsonify({"error": "도면 이미지 용량이 너무 큽니다 (최대 5MB)"}), 400
    icon = (data.get("icon") or "🔩").strip()
    width = data.get("width") or 130
    height = data.get("height") or 110
    stock_qty = int(data.get("stock_qty") or 0)
    supplier = (data.get("supplier") or "").strip()
    supplier_contact = (data.get("supplier_contact") or "").strip()
    lead_time_days = data.get("lead_time_days")
    lead_time_days = int(lead_time_days) if lead_time_days not in (None, "") else None

    conn = get_db()
    part_id = insert_part(
        conn, unit_id, name, spec=spec, cycle_days=cycle_days, cycle_unit=cycle_unit, cost=cost,
        last_replaced_date=last_replaced_date, note=note, memo=memo, drawing_data=drawing_data, icon=icon,
        pos_x=data.get("pos_x"), pos_y=data.get("pos_y"), width=width, height=height,
        stock_qty=stock_qty, supplier=supplier, supplier_contact=supplier_contact, lead_time_days=lead_time_days,
    )
    log_activity(conn, "create", "part", part_id, name)
    conn.commit()
    row = conn.execute("SELECT * FROM parts WHERE id = ?", (part_id,)).fetchone()
    conn.close()
    return jsonify(serialize_part(row)), 201


@app.route("/api/units/<int:unit_id>/apply-parts", methods=["POST"])
def apply_unit_parts(unit_id):
    conn = get_db()
    unit = conn.execute("SELECT * FROM units WHERE id = ?", (unit_id,)).fetchone()
    if not unit:
        conn.close()
        return jsonify({"error": "유닛을 찾을 수 없습니다"}), 404
    if unit["equipment_id"] != MASTER_EQUIPMENT_ID:
        conn.close()
        return jsonify({"error": "기준 설비(TEAG01호기)의 유닛에서만 사용할 수 있습니다"}), 400
    equipment_count, part_count = apply_unit_parts_to_other_equipment(conn, unit_id)
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "equipment_count": equipment_count, "part_count": part_count})


def parse_cycle_text(text):
    """"90", "90일", "2년" 등의 텍스트를 (cycle_days, cycle_unit) 튜플로 변환.
    비어있거나 해석할 수 없으면 기본값(N/A, 주기 없음)을 반환한다."""
    text = (text or "").strip()
    if not text:
        return None, "N/A"
    if text.endswith("년"):
        try:
            years = float(text[:-1].strip())
            return round(years * 365), "년"
        except ValueError:
            return None, "N/A"
    text = text[:-1].strip() if text.endswith("일") else text
    try:
        return round(float(text)), "일"
    except ValueError:
        return None, "N/A"


def parse_bulk_paste_text(text):
    """붙여넣은 텍스트(탭 또는 쉼표 구분)를 (유닛이름, 부품이름, Q-CODE, 부가설명, 금액, 교체주기일수, 교체주기단위)
    튜플 목록으로 변환. 교체주기는 "90", "90일", "2년"과 같이 입력할 수 있다."""
    rows = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        cols = line.split("\t")
        if len(cols) < 2:
            cols = line.split(",")
        cols = [c.strip() for c in cols]
        while len(cols) < 6:
            cols.append("")
        unit_text, part_name, q_code, note, cost_text, cycle_text = cols[0], cols[1], cols[2], cols[3], cols[4], cols[5]
        if not part_name:
            continue
        try:
            cost = float(cost_text) if cost_text else 0
        except ValueError:
            cost = 0
        cycle_days, cycle_unit = parse_cycle_text(cycle_text)
        rows.append((unit_text, part_name, q_code, note, cost, cycle_days, cycle_unit))
    return rows


def serialize_bulk_entry(row):
    d = dict(row)
    return d


def get_template_mapped_units(conn):
    """기본 유닛 구성(unit_templates)의 이름을 기준으로, 기준 설비(TEAG01호기)의 실제
    유닛 id에 매핑한 목록을 반환한다. (부품은 실제 유닛에만 등록할 수 있으므로, 선택 항목의
    이름 기준은 기본 유닛 구성을 따르되 등록 대상은 TEAG01호기의 해당 유닛으로 연결한다.)"""
    return conn.execute("""
        SELECT u.id AS id, t.name AS name
        FROM unit_templates t
        JOIN units u ON u.equipment_id = ? AND u.name = t.name
        WHERE u.deleted_at IS NULL
        ORDER BY t.id
    """, (MASTER_EQUIPMENT_ID,)).fetchall()


@app.route("/api/bulk-parts")
def list_bulk_parts():
    conn = get_db()
    rows = conn.execute("SELECT * FROM bulk_part_entries ORDER BY id DESC").fetchall()
    entries = []
    for r in rows:
        units = conn.execute("""
            SELECT u.id, u.name, u.icon
            FROM bulk_part_entry_units beu
            JOIN units u ON beu.unit_id = u.id
            WHERE beu.entry_id = ?
            ORDER BY u.id
        """, (r["id"],)).fetchall()
        d = serialize_bulk_entry(r)
        d["units"] = [dict(u) for u in units]
        entries.append(d)
    master_units = get_template_mapped_units(conn)
    conn.close()
    return jsonify({
        "entries": entries,
        "master_units": [dict(u) for u in master_units],
    })


@app.route("/api/bulk-parts/paste", methods=["POST"])
def paste_bulk_parts():
    data = request.get_json()
    text = data.get("text") or ""
    parsed = parse_bulk_paste_text(text)
    if not parsed:
        return jsonify({"error": "붙여넣은 내용에서 부품 정보를 찾을 수 없습니다"}), 400

    conn = get_db()
    created_ids = []
    for unit_text, part_name, q_code, note, cost, cycle_days, cycle_unit in parsed:
        cur = conn.execute(
            "INSERT INTO bulk_part_entries (raw_unit_text, part_name, q_code, note, cost, cycle_days, cycle_unit) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (unit_text, part_name, q_code, note, cost, cycle_days, cycle_unit),
        )
        created_ids.append(cur.lastrowid)
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "created_count": len(created_ids)}), 201


@app.route("/api/bulk-parts/<int:entry_id>", methods=["PUT"])
def update_bulk_part(entry_id):
    data = request.get_json()
    conn = get_db()
    entry = conn.execute("SELECT * FROM bulk_part_entries WHERE id = ?", (entry_id,)).fetchone()
    if not entry:
        conn.close()
        return jsonify({"error": "항목을 찾을 수 없습니다"}), 404
    part_name = (data.get("part_name") or entry["part_name"]).strip()
    q_code = data.get("q_code", entry["q_code"])
    note = data.get("note", entry["note"])
    cost = data.get("cost", entry["cost"])
    cycle_days, cycle_unit = resolve_cycle(data, entry["cycle_days"], entry["cycle_unit"])
    conn.execute(
        "UPDATE bulk_part_entries SET part_name = ?, q_code = ?, note = ?, cost = ?, cycle_days = ?, cycle_unit = ? WHERE id = ?",
        (part_name, q_code, note, cost, cycle_days, cycle_unit, entry_id),
    )
    if "unit_ids" in data:
        conn.execute("DELETE FROM bulk_part_entry_units WHERE entry_id = ?", (entry_id,))
        for uid in data["unit_ids"]:
            conn.execute(
                "INSERT OR IGNORE INTO bulk_part_entry_units (entry_id, unit_id) VALUES (?, ?)",
                (entry_id, uid),
            )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/bulk-parts/<int:entry_id>/register", methods=["POST"])
def register_bulk_part(entry_id):
    conn = get_db()
    entry = conn.execute("SELECT * FROM bulk_part_entries WHERE id = ?", (entry_id,)).fetchone()
    if not entry:
        conn.close()
        return jsonify({"error": "항목을 찾을 수 없습니다"}), 404
    unit_ids = [
        r["unit_id"] for r in conn.execute(
            "SELECT unit_id FROM bulk_part_entry_units WHERE entry_id = ?", (entry_id,)
        ).fetchall()
    ]
    if not unit_ids:
        conn.close()
        return jsonify({"error": "유닛을 먼저 선택하세요"}), 400

    part_ids = []
    for uid in unit_ids:
        unit = conn.execute("SELECT * FROM units WHERE id = ? AND deleted_at IS NULL", (uid,)).fetchone()
        if not unit:
            continue
        part_id = insert_part(
            conn, uid, entry["part_name"], spec=entry["q_code"] or "",
            cost=entry["cost"] or 0, note=entry["note"] or "",
            cycle_days=entry["cycle_days"], cycle_unit=entry["cycle_unit"] or "N/A",
        )
        part_ids.append(part_id)
        log_activity(conn, "create", "part", part_id, entry["part_name"], "부품 일괄 등록")
    conn.execute("UPDATE bulk_part_entries SET status = 'registered' WHERE id = ?", (entry_id,))
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "part_ids": part_ids})


@app.route("/api/bulk-parts/<int:entry_id>", methods=["DELETE"])
def delete_bulk_part(entry_id):
    conn = get_db()
    conn.execute("DELETE FROM bulk_part_entries WHERE id = ?", (entry_id,))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/parts/<int:part_id>", methods=["PUT"])
def update_part(part_id):
    data = request.get_json()
    conn = get_db()
    part = conn.execute("SELECT * FROM parts WHERE id = ?", (part_id,)).fetchone()
    if not part:
        conn.close()
        return jsonify({"error": "부품을 찾을 수 없습니다"}), 404
    name = (data.get("name") or part["name"]).strip()
    spec = data.get("spec", part["spec"])
    cycle_days, cycle_unit = resolve_cycle(data, part["cycle_days"], part["cycle_unit"])
    cost = data.get("cost", part["cost"])
    note = data.get("note", part["note"])
    memo = sanitize_rich_html(data.get("memo", part["memo"]))
    drawing_data = data.get("drawing_data", part["drawing_data"])
    if drawing_data and len(drawing_data) > MAX_DRAWING_DATA_LEN:
        conn.close()
        return jsonify({"error": "도면 이미지 용량이 너무 큽니다 (최대 5MB)"}), 400
    icon = (data.get("icon") or part["icon"]).strip()
    pos_x = data.get("pos_x", part["pos_x"])
    pos_y = data.get("pos_y", part["pos_y"])
    width = data.get("width", part["width"])
    height = data.get("height", part["height"])
    stock_qty = int(data.get("stock_qty", part["stock_qty"]) or 0)
    supplier = data.get("supplier", part["supplier"])
    supplier_contact = data.get("supplier_contact", part["supplier_contact"])
    lead_time_days = data.get("lead_time_days", part["lead_time_days"])
    lead_time_days = int(lead_time_days) if lead_time_days not in (None, "") else None
    meaningful_change = (
        name != part["name"] or spec != part["spec"] or cycle_days != part["cycle_days"]
        or cycle_unit != part["cycle_unit"] or cost != part["cost"] or note != part["note"]
        or stock_qty != (part["stock_qty"] or 0) or supplier != part["supplier"]
    )
    conn.execute(
        """UPDATE parts SET name = ?, spec = ?, cycle_days = ?, cycle_unit = ?, cost = ?, note = ?, memo = ?,
           drawing_data = ?, icon = ?, pos_x = ?, pos_y = ?, width = ?, height = ?,
           stock_qty = ?, supplier = ?, supplier_contact = ?, lead_time_days = ? WHERE id = ?""",
        (name, spec, cycle_days, cycle_unit, cost, note, memo, drawing_data, icon, pos_x, pos_y, width, height,
         stock_qty, supplier, supplier_contact, lead_time_days, part_id),
    )
    if meaningful_change:
        log_activity(conn, "update", "part", part_id, name, "부품 정보 수정")
    conn.commit()
    row = conn.execute("SELECT * FROM parts WHERE id = ?", (part_id,)).fetchone()
    conn.close()
    return jsonify(serialize_part(row))


@app.route("/api/parts/<int:part_id>", methods=["DELETE"])
def delete_part(part_id):
    conn = get_db()
    part = conn.execute("SELECT * FROM parts WHERE id = ?", (part_id,)).fetchone()
    if not part:
        conn.close()
        return jsonify({"error": "부품을 찾을 수 없습니다"}), 404
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn.execute("UPDATE parts SET deleted_at = ? WHERE id = ?", (now, part_id))
    log_activity(conn, "delete", "part", part_id, part["name"], "휴지통으로 이동")
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/parts/<int:part_id>/replace", methods=["POST"])
def replace_part(part_id):
    data = request.get_json()
    replaced_date = data.get("replaced_date") or str(date.today())
    cost = float(data.get("cost") or 0)
    note = (data.get("note") or "").strip()

    conn = get_db()
    part = conn.execute("SELECT * FROM parts WHERE id = ?", (part_id,)).fetchone()
    if not part:
        conn.close()
        return jsonify({"error": "부품을 찾을 수 없습니다"}), 404
    conn.execute(
        "INSERT INTO replacement_history (part_id, replaced_date, cost, note) VALUES (?, ?, ?, ?)",
        (part_id, replaced_date, cost, note),
    )
    conn.execute(
        "UPDATE parts SET last_replaced_date = ? WHERE id = ? AND (last_replaced_date IS NULL OR ? >= last_replaced_date)",
        (replaced_date, part_id, replaced_date),
    )
    log_activity(conn, "replace", "part", part_id, part["name"], f"교체 기록 추가 ({replaced_date})")
    conn.commit()
    row = conn.execute("SELECT * FROM parts WHERE id = ?", (part_id,)).fetchone()
    conn.close()
    return jsonify(serialize_part(row)), 201


@app.route("/api/parts/<int:part_id>/history")
def part_history(part_id):
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM replacement_history WHERE part_id = ? ORDER BY replaced_date DESC, id DESC",
        (part_id,),
    ).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route("/api/history/<int:history_id>", methods=["DELETE"])
def delete_history(history_id):
    conn = get_db()
    row = conn.execute("SELECT * FROM replacement_history WHERE id = ?", (history_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "이력을 찾을 수 없습니다"}), 404
    part_id = row["part_id"]
    conn.execute("DELETE FROM replacement_history WHERE id = ?", (history_id,))
    # 삭제 후 남아있는 이력 중 가장 최근 날짜를 기준으로 교체주기 계산용 last_replaced_date를 다시 계산한다
    # (남은 이력이 없으면 "미기록" 상태로 되돌아간다).
    remaining = conn.execute(
        "SELECT MAX(replaced_date) AS last_date FROM replacement_history WHERE part_id = ?", (part_id,)
    ).fetchone()
    conn.execute(
        "UPDATE parts SET last_replaced_date = ? WHERE id = ?",
        (remaining["last_date"], part_id),
    )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/equipments/<int:equipment_id>/notes")
def get_notes(equipment_id):
    conn = get_db()
    row = conn.execute(
        "SELECT content, updated_at FROM equipment_notes WHERE equipment_id = ?", (equipment_id,)
    ).fetchone()
    conn.close()
    return jsonify(dict(row) if row else {"content": "", "updated_at": None})


@app.route("/api/equipments/<int:equipment_id>/notes", methods=["PUT"])
def update_notes(equipment_id):
    data = request.get_json()
    content = sanitize_rich_html(data.get("content") or "")
    conn = get_db()
    conn.execute(
        "INSERT INTO equipment_notes (equipment_id, content, updated_at) VALUES (?, ?, datetime('now','localtime')) "
        "ON CONFLICT(equipment_id) DO UPDATE SET content = excluded.content, updated_at = excluded.updated_at",
        (equipment_id, content),
    )
    conn.commit()
    row = conn.execute(
        "SELECT content, updated_at FROM equipment_notes WHERE equipment_id = ?", (equipment_id,)
    ).fetchone()
    conn.close()
    return jsonify(dict(row))


@app.route("/api/units/<int:unit_id>/notes")
def get_unit_notes(unit_id):
    conn = get_db()
    row = conn.execute(
        "SELECT content, updated_at FROM unit_notes WHERE unit_id = ?", (unit_id,)
    ).fetchone()
    conn.close()
    return jsonify(dict(row) if row else {"content": "", "updated_at": None})


@app.route("/api/units/<int:unit_id>/notes", methods=["PUT"])
def update_unit_notes(unit_id):
    data = request.get_json()
    content = sanitize_rich_html(data.get("content") or "")
    conn = get_db()
    conn.execute(
        "INSERT INTO unit_notes (unit_id, content, updated_at) VALUES (?, ?, datetime('now','localtime')) "
        "ON CONFLICT(unit_id) DO UPDATE SET content = excluded.content, updated_at = excluded.updated_at",
        (unit_id, content),
    )
    conn.commit()
    row = conn.execute(
        "SELECT content, updated_at FROM unit_notes WHERE unit_id = ?", (unit_id,)
    ).fetchone()
    conn.close()
    return jsonify(dict(row))


@app.route("/api/unit-templates")
def list_unit_templates():
    conn = get_db()
    rows = conn.execute("SELECT * FROM unit_templates ORDER BY id").fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


@app.route("/api/unit-templates", methods=["POST"])
def add_unit_template():
    data = request.get_json()
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "유닛 이름을 입력하세요"}), 400
    icon = (data.get("icon") or "⚙️").strip()
    color = (data.get("color") or "#1a3a5c").strip()
    pos_x = data.get("pos_x")
    pos_y = data.get("pos_y")
    if pos_x is None or pos_y is None:
        pos_x, pos_y = random.uniform(15, 85), random.uniform(20, 80)
    width = data.get("width") or 140
    height = data.get("height") or 110
    conn = get_db()
    cur = conn.execute(
        "INSERT INTO unit_templates (name, icon, color, pos_x, pos_y, width, height) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (name, icon, color, pos_x, pos_y, width, height),
    )
    conn.commit()
    row = conn.execute("SELECT * FROM unit_templates WHERE id = ?", (cur.lastrowid,)).fetchone()
    conn.close()
    return jsonify(dict(row)), 201


@app.route("/api/unit-templates/<int:template_id>", methods=["PUT"])
def update_unit_template(template_id):
    data = request.get_json()
    conn = get_db()
    t = conn.execute("SELECT * FROM unit_templates WHERE id = ?", (template_id,)).fetchone()
    if not t:
        conn.close()
        return jsonify({"error": "유닛을 찾을 수 없습니다"}), 404
    name = (data.get("name") or t["name"]).strip()
    icon = (data.get("icon") or t["icon"]).strip()
    color = (data.get("color") or t["color"]).strip()
    pos_x = data.get("pos_x", t["pos_x"])
    pos_y = data.get("pos_y", t["pos_y"])
    width = data.get("width", t["width"])
    height = data.get("height", t["height"])
    conn.execute(
        "UPDATE unit_templates SET name = ?, icon = ?, color = ?, pos_x = ?, pos_y = ?, width = ?, height = ? WHERE id = ?",
        (name, icon, color, pos_x, pos_y, width, height, template_id),
    )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/unit-templates/<int:template_id>", methods=["DELETE"])
def delete_unit_template(template_id):
    conn = get_db()
    conn.execute("DELETE FROM unit_templates WHERE id = ?", (template_id,))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/unit-templates/apply", methods=["POST"])
def apply_unit_templates():
    conn = get_db()
    templates = conn.execute("SELECT * FROM unit_templates ORDER BY id").fetchall()
    equipments = conn.execute("SELECT id FROM equipments WHERE deleted_at IS NULL").fetchall()
    template_names = {t["name"] for t in templates}

    for eq in equipments:
        existing = {
            u["name"]: u
            for u in conn.execute(
                "SELECT * FROM units WHERE equipment_id = ? AND deleted_at IS NULL", (eq["id"],)
            ).fetchall()
        }
        for t in templates:
            if t["name"] in existing:
                u = existing[t["name"]]
                conn.execute(
                    "UPDATE units SET icon = ?, color = ?, pos_x = ?, pos_y = ?, width = ?, height = ? WHERE id = ?",
                    (t["icon"], t["color"], t["pos_x"], t["pos_y"], t["width"], t["height"], u["id"]),
                )
            else:
                conn.execute(
                    "INSERT INTO units (equipment_id, name, icon, color, pos_x, pos_y, width, height) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (eq["id"], t["name"], t["icon"], t["color"], t["pos_x"], t["pos_y"], t["width"], t["height"]),
                )
        for name, u in existing.items():
            if name not in template_names:
                conn.execute("DELETE FROM units WHERE id = ?", (u["id"],))

    conn.commit()
    conn.close()
    return jsonify({"ok": True, "equipment_count": len(equipments), "unit_count": len(templates)})


@app.route("/api/alerts")
def api_alerts():
    return jsonify(get_alert_parts())


@app.route("/api/alerts/export")
def export_alerts_csv():
    parts = get_alert_parts()
    header = ["상태", "설비", "유닛", "부품명", "규격", "교체주기(일)", "최근교체일", "다음교체예정일", "남은/초과일수", "비고"]
    rows = []
    for p in parts:
        days_text = f"{abs(p['days_left'])}일 초과" if p["status"] == "overdue" else f"{p['days_left']}일 남음"
        rows.append([
            p["label"], p["equipment_name"], p["unit_name"], p["name"], p.get("spec") or "",
            p["cycle_days"], p.get("last_replaced_date") or "", p.get("next_due") or "",
            days_text, p.get("note") or "",
        ])
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return csv_response(f"교체현황_{timestamp}.csv", header, rows)


@app.route("/api/search")
def api_search():
    q = (request.args.get("q") or "").strip()
    status = request.args.get("status") or None
    equipment_id = request.args.get("equipment_id")
    equipment_id = int(equipment_id) if equipment_id else None
    return jsonify(search_parts(q, status=status, equipment_id=equipment_id))


def build_stats_payload(conn, unit_names, start_date=None, end_date=None):
    """부품 규격 기준(금액순/사용량 많은순/교체주기 짧은순)과 유닛 기준(부품수 많은순) 통계를 함께 만든다."""
    period_active = bool(start_date or end_date)
    spec_rows = get_part_spec_stats(conn, unit_names, start_date=start_date, end_date=end_date)
    # 기간 필터가 없으면 아직 교체 이력이 없는 부품도 등록된 금액으로 집계되므로(get_part_spec_stats
    # 참고) 전부 보여준다. 기간 필터가 있으면 그 기간에 실제 교체 기록이 있는 부품만 보여준다.
    by_cost = sorted(
        (r for r in spec_rows if period_active is False or r["usage_count"] > 0),
        key=lambda r: r["total_cost"], reverse=True
    )
    by_usage = sorted(spec_rows, key=lambda r: r["usage_count"], reverse=True)
    by_short_cycle = sorted(
        (r for r in spec_rows if r["min_cycle_days"] is not None), key=lambda r: r["min_cycle_days"]
    )

    unit_query = """
        SELECT u.id AS unit_id, u.name AS unit_name, u.icon AS unit_icon,
               e.id AS equipment_id, e.name AS equipment_name, e.icon AS equipment_icon,
               (SELECT COUNT(*) FROM parts p WHERE p.unit_id = u.id AND p.deleted_at IS NULL) AS part_count
        FROM units u
        JOIN equipments e ON u.equipment_id = e.id
        WHERE u.deleted_at IS NULL AND e.deleted_at IS NULL
    """
    params = []
    if unit_names:
        placeholders = ",".join("?" for _ in unit_names)
        unit_query += f" AND u.name IN ({placeholders})"
        params = unit_names
    unit_query += " ORDER BY u.id"
    unit_rows = [dict(r) for r in conn.execute(unit_query, params).fetchall()]
    by_part_count = sorted(unit_rows, key=lambda r: r["part_count"], reverse=True)

    all_unit_names = [
        r["name"] for r in conn.execute(
            "SELECT DISTINCT name FROM units WHERE deleted_at IS NULL ORDER BY name"
        ).fetchall()
    ]

    return {
        "by_cost": by_cost,
        "by_usage": by_usage,
        "by_short_cycle": by_short_cycle,
        "by_part_count": by_part_count,
        "unit_names": all_unit_names,
        "period_active": bool(start_date or end_date),
    }


@app.route("/api/stats")
def api_stats():
    unit_names = request.args.getlist("unit_name")
    start_date = request.args.get("start_date") or None
    end_date = request.args.get("end_date") or None
    conn = get_db()
    payload = build_stats_payload(conn, unit_names, start_date=start_date, end_date=end_date)
    conn.close()
    return jsonify(payload)


def format_cycle_for_export(days, unit):
    if days is None or unit == "N/A":
        return "N/A"
    if unit == "년":
        return f"{round(days / 365, 2)}년"
    return f"{days}일"


@app.route("/api/stats/export.csv")
def export_stats_csv():
    unit_names = request.args.getlist("unit_name")
    start_date = request.args.get("start_date") or None
    end_date = request.args.get("end_date") or None
    conn = get_db()
    payload = build_stats_payload(conn, unit_names, start_date=start_date, end_date=end_date)
    conn.close()

    buf = io.StringIO()
    writer = csv.writer(buf)

    writer.writerow(["[금액순 (부품 규격 기준)]"])
    writer.writerow(["순위", "부품명", "규격", "총 금액", "등록 수", "교체 횟수"])
    for i, r in enumerate(payload["by_cost"], 1):
        writer.writerow([i, r["name"], r["spec"], r["total_cost"], r["instance_count"], r["usage_count"]])
    writer.writerow([])

    writer.writerow(["[사용량 많은순 (부품 규격 기준)]"])
    writer.writerow(["순위", "부품명", "규격", "교체 횟수", "등록 수", "총 금액"])
    for i, r in enumerate(payload["by_usage"], 1):
        writer.writerow([i, r["name"], r["spec"], r["usage_count"], r["instance_count"], r["total_cost"]])
    writer.writerow([])

    writer.writerow(["[교체 주기 짧은순 (부품 규격 기준)]"])
    writer.writerow(["순위", "부품명", "규격", "교체 주기", "등록 수"])
    for i, r in enumerate(payload["by_short_cycle"], 1):
        writer.writerow([
            i, r["name"], r["spec"], format_cycle_for_export(r["min_cycle_days"], r["min_cycle_unit"]),
            r["instance_count"],
        ])
    writer.writerow([])

    writer.writerow(["[부품수 많은순 (유닛 기준)]"])
    writer.writerow(["순위", "설비", "유닛", "부품수"])
    for i, r in enumerate(payload["by_part_count"], 1):
        writer.writerow([i, r["equipment_name"], r["unit_name"], r["part_count"]])

    data = "﻿" + buf.getvalue()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"통계_{timestamp}.csv"
    encoded_name = quote(filename)
    return Response(
        data,
        mimetype="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename=stats.csv; filename*=UTF-8''{encoded_name}"
        },
    )


def build_stats_pptx(payload, unit_names):
    """조건별(금액순/사용량순/교체주기순/부품수순) TOP 5를 표+막대 그래프로 함께 보여주는
    보고서 형태의 PPT를 생성한다."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "부품/유닛 통계 리포트"
    filter_text = ", ".join(unit_names) if unit_names else "전체 유닛"
    title_slide.placeholders[1].text = (
        f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n필터: {filter_text}\n"
        f"조건별 TOP 5 표 및 막대 그래프"
    )

    blank_layout = prs.slide_layouts[6]

    def add_report_slide(title, headers, rows, chart_labels, chart_values, chart_value_title, number_format="0"):
        """표(왼쪽, 상위 5건)와 막대 그래프(오른쪽)를 함께 보여주는 보고서 슬라이드를 추가한다."""
        slide = prs.slides.add_slide(blank_layout)
        title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(12.5), Inches(0.7))
        tf = title_box.text_frame
        tf.text = title
        tf.paragraphs[0].font.size = Pt(28)
        tf.paragraphs[0].font.bold = True

        n_rows = len(rows) + 1
        n_cols = len(headers)
        if rows:
            table = slide.shapes.add_table(
                n_rows, n_cols, Inches(0.4), Inches(1.3), Inches(6.1), Inches(0.5 * n_rows)
            ).table
            for c, h in enumerate(headers):
                cell = table.cell(0, c)
                cell.text = str(h)
                cell.text_frame.paragraphs[0].font.bold = True
                cell.text_frame.paragraphs[0].font.size = Pt(13)
            for r_i, row in enumerate(rows, 1):
                for c_i, val in enumerate(row):
                    cell = table.cell(r_i, c_i)
                    cell.text = str(val)
                    cell.text_frame.paragraphs[0].font.size = Pt(12)
        else:
            empty_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.5), Inches(6.1), Inches(1.0))
            empty_box.text_frame.text = "표시할 데이터가 없습니다."

        if chart_labels:
            chart_data = CategoryChartData()
            chart_data.categories = chart_labels
            chart_data.add_series(chart_value_title, chart_values)
            chart_frame = slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.9), Inches(1.3), Inches(6.0), Inches(4.8), chart_data
            )
            chart = chart_frame.chart
            chart.has_legend = False
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(11)
            plot.data_labels.number_format = number_format
            plot.data_labels.number_format_is_linked = False
            chart.category_axis.tick_labels.font.size = Pt(10)
            chart.value_axis.tick_labels.font.size = Pt(10)

    def chart_label(r):
        return f'{r["name"]} ({r["spec"]})' if r.get("spec") else r["name"]

    by_cost5 = payload["by_cost"][:5]
    add_report_slide(
        "금액순 TOP 5 (부품 규격 기준)",
        ["순위", "부품명", "규격", "총 금액(원)", "등록 수", "교체 횟수"],
        [
            [i, r["name"], r["spec"], f'{r["total_cost"]:,.0f}', r["instance_count"], r["usage_count"]]
            for i, r in enumerate(by_cost5, 1)
        ],
        [chart_label(r) for r in by_cost5],
        [r["total_cost"] for r in by_cost5],
        "총 금액(원)",
        number_format="#,##0",
    )

    by_usage5 = payload["by_usage"][:5]
    add_report_slide(
        "사용량 많은순 TOP 5 (부품 규격 기준)",
        ["순위", "부품명", "규격", "교체 횟수", "등록 수", "총 금액(원)"],
        [
            [i, r["name"], r["spec"], r["usage_count"], r["instance_count"], f'{r["total_cost"]:,.0f}']
            for i, r in enumerate(by_usage5, 1)
        ],
        [chart_label(r) for r in by_usage5],
        [r["usage_count"] for r in by_usage5],
        "교체 횟수",
    )

    by_short5 = payload["by_short_cycle"][:5]
    add_report_slide(
        "교체 주기 짧은순 TOP 5 (부품 규격 기준)",
        ["순위", "부품명", "규격", "교체 주기", "등록 수"],
        [
            [i, r["name"], r["spec"], format_cycle_for_export(r["min_cycle_days"], r["min_cycle_unit"]), r["instance_count"]]
            for i, r in enumerate(by_short5, 1)
        ],
        [chart_label(r) for r in by_short5],
        [r["min_cycle_days"] for r in by_short5],
        "교체 주기(일)",
    )

    by_part5 = payload["by_part_count"][:5]
    add_report_slide(
        "부품수 많은순 TOP 5 (유닛 기준)",
        ["순위", "설비", "유닛", "부품수"],
        [
            [i, r["equipment_name"], r["unit_name"], r["part_count"]]
            for i, r in enumerate(by_part5, 1)
        ],
        [f'{r["equipment_name"]} {r["unit_name"]}' for r in by_part5],
        [r["part_count"] for r in by_part5],
        "부품수",
    )

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


@app.route("/api/stats/export.pptx")
def export_stats_pptx():
    unit_names = request.args.getlist("unit_name")
    start_date = request.args.get("start_date") or None
    end_date = request.args.get("end_date") or None
    try:
        conn = get_db()
        payload = build_stats_payload(conn, unit_names, start_date=start_date, end_date=end_date)
        conn.close()
        buf = build_stats_pptx(payload, unit_names)
    except Exception:
        # PPT 생성 중 예기치 못한 오류가 나도 서버가 조용히 죽거나 모호한 메시지 대신,
        # 콘솔에 원인을 남기고 클라이언트에는 명확한 오류를 반환한다.
        traceback.print_exc()
        return jsonify({"error": "PPT 생성 중 오류가 발생했습니다. 서버 콘솔의 오류 메시지를 확인해주세요."}), 500
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"통계_{timestamp}.pptx",
        mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


@app.route("/api/backup", methods=["POST"])
def download_backup():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"equipment_backup_{timestamp}.db"
    dest_path = os.path.join(BACKUP_DIR, filename)
    shutil.copyfile(DB_PATH, dest_path)
    conn = get_db()
    log_activity(conn, "backup", "backup", None, filename, f"백업 파일 저장: {dest_path}")
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "filename": filename, "path": dest_path})


@app.route("/trash")
def trash_page():
    return render_template("trash.html")


@app.route("/api/trash")
def api_trash():
    conn = get_db()
    equipments = conn.execute("""
        SELECT *, (SELECT COUNT(*) FROM units WHERE equipment_id = equipments.id) AS unit_count
        FROM equipments WHERE deleted_at IS NOT NULL ORDER BY deleted_at DESC
    """).fetchall()
    # 소속 설비/유닛이 함께 삭제된(연쇄 삭제된) 항목은 부모를 복원하면 같이 복원되므로
    # 여기서는 "단독으로" 삭제된 유닛/부품만 보여준다 (부모는 살아있는데 이것만 삭제된 경우).
    units = conn.execute("""
        SELECT u.*, e.name AS equipment_name
        FROM units u
        JOIN equipments e ON u.equipment_id = e.id
        WHERE u.deleted_at IS NOT NULL AND e.deleted_at IS NULL
        ORDER BY u.deleted_at DESC
    """).fetchall()
    parts = conn.execute("""
        SELECT p.*, u.name AS unit_name, e.name AS equipment_name
        FROM parts p
        JOIN units u ON p.unit_id = u.id
        JOIN equipments e ON u.equipment_id = e.id
        WHERE p.deleted_at IS NOT NULL AND u.deleted_at IS NULL
        ORDER BY p.deleted_at DESC
    """).fetchall()
    conn.close()
    return jsonify({
        "equipments": [dict(r) for r in equipments],
        "units": [dict(r) for r in units],
        "parts": [dict(r) for r in parts],
    })


@app.route("/api/trash/equipment/<int:equipment_id>/restore", methods=["POST"])
def restore_equipment(equipment_id):
    conn = get_db()
    equipment = conn.execute("SELECT * FROM equipments WHERE id = ?", (equipment_id,)).fetchone()
    if not equipment:
        conn.close()
        return jsonify({"error": "설비를 찾을 수 없습니다"}), 404
    conn.execute("UPDATE equipments SET deleted_at = NULL WHERE id = ?", (equipment_id,))
    unit_ids = [
        r["id"] for r in conn.execute(
            "SELECT id FROM units WHERE equipment_id = ?", (equipment_id,)
        ).fetchall()
    ]
    if unit_ids:
        placeholders = ",".join("?" for _ in unit_ids)
        conn.execute(f"UPDATE units SET deleted_at = NULL WHERE id IN ({placeholders})", unit_ids)
        conn.execute(f"UPDATE parts SET deleted_at = NULL WHERE unit_id IN ({placeholders})", unit_ids)
    log_activity(conn, "restore", "equipment", equipment_id, equipment["name"], "휴지통에서 복원")
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/trash/unit/<int:unit_id>/restore", methods=["POST"])
def restore_unit(unit_id):
    conn = get_db()
    unit = conn.execute("SELECT * FROM units WHERE id = ?", (unit_id,)).fetchone()
    if not unit:
        conn.close()
        return jsonify({"error": "유닛을 찾을 수 없습니다"}), 404
    conn.execute("UPDATE units SET deleted_at = NULL WHERE id = ?", (unit_id,))
    conn.execute("UPDATE parts SET deleted_at = NULL WHERE unit_id = ?", (unit_id,))
    log_activity(conn, "restore", "unit", unit_id, unit["name"], "휴지통에서 복원")
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/trash/part/<int:part_id>/restore", methods=["POST"])
def restore_part(part_id):
    conn = get_db()
    part = conn.execute("SELECT * FROM parts WHERE id = ?", (part_id,)).fetchone()
    if not part:
        conn.close()
        return jsonify({"error": "부품을 찾을 수 없습니다"}), 404
    conn.execute("UPDATE parts SET deleted_at = NULL WHERE id = ?", (part_id,))
    log_activity(conn, "restore", "part", part_id, part["name"], "휴지통에서 복원")
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/trash/equipment/<int:equipment_id>", methods=["DELETE"])
def permanent_delete_equipment(equipment_id):
    conn = get_db()
    equipment = conn.execute("SELECT * FROM equipments WHERE id = ?", (equipment_id,)).fetchone()
    if equipment:
        conn.execute("DELETE FROM equipments WHERE id = ?", (equipment_id,))
        log_activity(conn, "permanent_delete", "equipment", equipment_id, equipment["name"], "영구 삭제")
        conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/trash/unit/<int:unit_id>", methods=["DELETE"])
def permanent_delete_unit(unit_id):
    conn = get_db()
    unit = conn.execute("SELECT * FROM units WHERE id = ?", (unit_id,)).fetchone()
    if unit:
        conn.execute("DELETE FROM units WHERE id = ?", (unit_id,))
        log_activity(conn, "permanent_delete", "unit", unit_id, unit["name"], "영구 삭제")
        conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/trash/part/<int:part_id>", methods=["DELETE"])
def permanent_delete_part(part_id):
    conn = get_db()
    part = conn.execute("SELECT * FROM parts WHERE id = ?", (part_id,)).fetchone()
    if part:
        conn.execute("DELETE FROM parts WHERE id = ?", (part_id,))
        log_activity(conn, "permanent_delete", "part", part_id, part["name"], "영구 삭제")
        conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/activity-log")
def activity_log_page():
    return render_template("activity_log.html")


@app.route("/api/activity-log")
def api_activity_log():
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM activity_log ORDER BY id DESC LIMIT 300"
    ).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])


def mail_scheduler_loop():
    """30초마다 예약 시간을 확인해, 설정된 시각(HH:MM)이 되면 하루 1회 발송한다."""
    last_sent_date = None
    while True:
        try:
            conn = get_db()
            t = get_config(conn, "mail_schedule_time")
            conn.close()
            if t:
                now = datetime.now()
                if now.strftime("%H:%M") == t and last_sent_date != now.strftime("%Y-%m-%d"):
                    ok, msg = send_status_mail()
                    print(f"[예약 메일] {now:%Y-%m-%d %H:%M} → {msg}")
                    last_sent_date = now.strftime("%Y-%m-%d")
        except Exception as e:
            print("[예약 메일] 오류:", e)
        time.sleep(30)


if __name__ == "__main__":
    init_db()
    _conn = get_db()
    app.secret_key = get_config(_conn, "secret_key")
    _conn.close()
    lan_ip = get_lan_ip()
    print("설비 부품 교체 관리 시스템 시작!")
    print(f"  이 컴퓨터에서 접속: http://localhost:5000")
    print(f"  같은 네트워크의 다른 사람 접속: http://{lan_ip}:5000")
    print("  (다른 사람이 접속 안 되면 Windows 방화벽에서 Python 허용 여부를 확인하세요)")
    print(f"  최초 접속 비밀번호: {DEFAULT_PASSWORD} (로그인 후 반드시 변경해주세요)")
    threading.Thread(target=mail_scheduler_loop, daemon=True).start()
    app.run(host="0.0.0.0", port=5000, debug=False, threaded=True)
