"""
RPA 자동화 스튜디오 (Windows)

정상적인 접근 권한이 있는 프로그램과 문서에서 반복 작업을 자동화합니다.
보안 기능, DRM, 접근 권한 또는 복사 방지 기능을 우회하지 않습니다.
문서 내용과 입력 문자열은 실행 로그나 스크린샷 이름에 기록하지 않습니다.

주요 기능
---------
- 마우스/키보드 실제 동작 녹화(F8 종료), 단계 편집/복제/비활성화
- 클릭, 이동, 드래그, 스크롤, 모든 키/단축키(가운데 클릭 포함), 한글·유니코드 입력
- 화면 이미지 찾기/클릭, 화면을 드래그해 바로 이미지 등록, Windows 창 제목 대기/활성화
- 이미지 있음/없음에 따라 다음 단계를 건너뛰는 조건 분기
- CSV/XLSX/XLSM 행별 데이터 반복 및 {열이름}, {A}, {row} 템플릿
- 실행 일시정지, 단계별 재시도, 실패 스크린샷, 실행 로그, 선택 단계 즉시 테스트 실행
- 실행 완료 후 성공/건너뜀/재시도 횟수와 소요 시간 요약
- 로그·스크린샷 보관 기간 설정 및 자동 정리
- 프로그램 내부 예약 실행, Windows 작업 스케줄러 등록(창 최소화·완료 후 자동 종료 옵션)
- 예약 실행과 수동 실행이 겹치지 않도록 중복 실행 방지
- 기존 v1~v3 JSON 설정 자동 호환

필수 설치
---------
    py -m pip install pyautogui pynput pillow

선택 설치
---------
    py -m pip install openpyxl opencv-python
"""

from __future__ import annotations

import argparse
import copy
import csv
import ctypes
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import traceback
import uuid
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk
import tkinter as tk
from typing import Any, Callable, Iterable

try:
    import pyautogui
except Exception:
    pyautogui = None

try:
    from pynput import keyboard as pynput_keyboard
    from pynput import mouse as pynput_mouse
except Exception:
    pynput_keyboard = None
    pynput_mouse = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


APP_TITLE = "RPA 자동화 스튜디오"
PROFILE_VERSION = 4
MAX_ACTIONS = 20_000
STOP_RECORD_KEY = "f8"
WINDOWS = sys.platform == "win32"


class AutomationStopped(Exception):
    """사용자가 실행을 중지했을 때 작업 스레드를 빠져나오기 위한 예외."""


class SafeFormatDict(dict[str, Any]):
    def __missing__(self, key: str) -> str:
        return "{" + key + "}"


def new_action_id() -> str:
    return uuid.uuid4().hex[:12]


def excel_column_name(index: int) -> str:
    result = ""
    value = index
    while value > 0:
        value, remainder = divmod(value - 1, 26)
        result = chr(65 + remainder) + result
    return result


def render_template(template: str, context: dict[str, Any]) -> str:
    """{열이름} 형식의 템플릿을 채운다. 짝이 안 맞는 중괄호 등 형식 오류가 있으면
    입력한 문자열을 그대로 반환한다(템플릿을 의도하지 않은 일반 텍스트 입력 보호)."""
    normalized = {str(key): "" if value is None else value for key, value in context.items()}
    try:
        return template.format_map(SafeFormatDict(normalized))
    except (ValueError, IndexError):
        return template


def enable_windows_dpi_awareness() -> None:
    if not WINDOWS:
        return
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)
    except (AttributeError, OSError):
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except (AttributeError, OSError):
            pass


def default_log_directory() -> Path:
    documents = Path.home() / "Documents"
    base = documents if documents.exists() else Path.home()
    return base / "RPA_Automation_Logs"


def send_unicode_text_windows(text: str, interval: float = 0.01) -> None:
    """Windows SendInput으로 클립보드를 사용하지 않고 유니코드 문자열을 입력한다."""
    if not WINDOWS:
        raise RuntimeError("한글·유니코드 직접 입력은 Windows에서 지원됩니다.")

    from ctypes import wintypes

    ulong_ptr = wintypes.WPARAM

    class MOUSEINPUT(ctypes.Structure):
        _fields_ = [
            ("dx", wintypes.LONG),
            ("dy", wintypes.LONG),
            ("mouseData", wintypes.DWORD),
            ("dwFlags", wintypes.DWORD),
            ("time", wintypes.DWORD),
            ("dwExtraInfo", ulong_ptr),
        ]

    class KEYBDINPUT(ctypes.Structure):
        _fields_ = [
            ("wVk", wintypes.WORD),
            ("wScan", wintypes.WORD),
            ("dwFlags", wintypes.DWORD),
            ("time", wintypes.DWORD),
            ("dwExtraInfo", ulong_ptr),
        ]

    class HARDWAREINPUT(ctypes.Structure):
        _fields_ = [
            ("uMsg", wintypes.DWORD),
            ("wParamL", wintypes.WORD),
            ("wParamH", wintypes.WORD),
        ]

    class INPUTUNION(ctypes.Union):
        _fields_ = [("mi", MOUSEINPUT), ("ki", KEYBDINPUT), ("hi", HARDWAREINPUT)]

    class INPUT(ctypes.Structure):
        _anonymous_ = ("union",)
        _fields_ = [("type", wintypes.DWORD), ("union", INPUTUNION)]

    keyeventf_keyup = 0x0002
    keyeventf_unicode = 0x0004
    units = text.encode("utf-16-le")
    for offset in range(0, len(units), 2):
        scan = int.from_bytes(units[offset : offset + 2], "little")
        for flags in (keyeventf_unicode, keyeventf_unicode | keyeventf_keyup):
            item = INPUT(type=1)
            item.ki = KEYBDINPUT(0, scan, flags, 0, 0)
            sent = ctypes.windll.user32.SendInput(1, ctypes.byref(item), ctypes.sizeof(INPUT))
            if sent != 1:
                raise OSError("Windows SendInput 문자 입력에 실패했습니다.")
        if interval > 0:
            time.sleep(interval)


CF_UNICODETEXT = 13
GMEM_MOVEABLE = 0x0002


def win32_get_clipboard_text() -> str | None:
    if not WINDOWS:
        return None
    user32 = ctypes.windll.user32
    kernel32 = ctypes.windll.kernel32
    if not user32.OpenClipboard(None):
        return None
    try:
        handle = user32.GetClipboardData(CF_UNICODETEXT)
        if not handle:
            return None
        pointer = kernel32.GlobalLock(handle)
        if not pointer:
            return None
        try:
            return ctypes.wstring_at(pointer)
        finally:
            kernel32.GlobalUnlock(handle)
    finally:
        user32.CloseClipboard()


def win32_set_clipboard_text(text: str) -> None:
    if not WINDOWS:
        raise RuntimeError("클립보드 붙여넣기 입력은 Windows에서 지원됩니다.")
    user32 = ctypes.windll.user32
    kernel32 = ctypes.windll.kernel32
    data = text.encode("utf-16-le") + b"\x00\x00"
    handle = kernel32.GlobalAlloc(GMEM_MOVEABLE, len(data))
    if not handle:
        raise OSError("클립보드에 쓸 메모리를 할당하지 못했습니다.")
    pointer = kernel32.GlobalLock(handle)
    if not pointer:
        kernel32.GlobalFree(handle)
        raise OSError("클립보드 메모리를 잠그지 못했습니다.")
    ctypes.memmove(pointer, data, len(data))
    kernel32.GlobalUnlock(handle)
    if not user32.OpenClipboard(None):
        kernel32.GlobalFree(handle)
        raise OSError("클립보드를 열지 못했습니다.")
    try:
        user32.EmptyClipboard()
        if not user32.SetClipboardData(CF_UNICODETEXT, handle):
            kernel32.GlobalFree(handle)
            raise OSError("클립보드에 값을 설정하지 못했습니다.")
    finally:
        user32.CloseClipboard()


def list_visible_windows() -> list[tuple[int, str]]:
    if not WINDOWS:
        return []
    from ctypes import wintypes

    user32 = ctypes.windll.user32
    results: list[tuple[int, str]] = []
    callback_type = ctypes.WINFUNCTYPE(ctypes.c_bool, wintypes.HWND, wintypes.LPARAM)

    def callback(hwnd: int, _lparam: int) -> bool:
        if not user32.IsWindowVisible(hwnd):
            return True
        length = user32.GetWindowTextLengthW(hwnd)
        if length <= 0:
            return True
        buffer = ctypes.create_unicode_buffer(length + 1)
        user32.GetWindowTextW(hwnd, buffer, length + 1)
        title = buffer.value.strip()
        if title:
            results.append((int(hwnd), title))
        return True

    user32.EnumWindows(callback_type(callback), 0)
    return results


def find_window(title_pattern: str, regex: bool = False) -> tuple[int, str] | None:
    for hwnd, title in list_visible_windows():
        if regex:
            try:
                if re.search(title_pattern, title, re.IGNORECASE):
                    return hwnd, title
            except re.error as error:
                raise ValueError(f"창 제목 정규식 오류: {error}") from error
        elif title_pattern.casefold() in title.casefold():
            return hwnd, title
    return None


def activate_window(hwnd: int) -> None:
    if not WINDOWS:
        return
    user32 = ctypes.windll.user32
    user32.ShowWindow(hwnd, 9)
    user32.SetForegroundWindow(hwnd)


RUN_MUTEX_NAME = "Global\\RPA_Automation_Studio_Run_Mutex"
ERROR_ALREADY_EXISTS = 183


def acquire_run_mutex() -> Any:
    """자동화 실행 시작 시 호출한다. Windows 작업 스케줄러 실행과 수동 실행이 겹치는
    것을 막기 위한 시스템 전역 뮤텍스이다. 이미 다른 프로세스가 실행 중이면 None을 반환한다."""
    if not WINDOWS:
        return object()
    handle = ctypes.windll.kernel32.CreateMutexW(None, False, RUN_MUTEX_NAME)
    if not handle:
        return object()
    if ctypes.GetLastError() == ERROR_ALREADY_EXISTS:
        ctypes.windll.kernel32.CloseHandle(handle)
        return None
    return handle


def release_run_mutex(handle: Any) -> None:
    if WINDOWS and handle:
        try:
            ctypes.windll.kernel32.CloseHandle(handle)
        except OSError:
            pass


@dataclass
class Action:
    kind: str
    params: dict[str, Any]
    enabled: bool = True
    action_id: str = field(default_factory=new_action_id)

    @classmethod
    def from_dict(cls, item: dict[str, Any]) -> "Action":
        return cls(
            kind=str(item["kind"]),
            params=dict(item.get("params", {})),
            enabled=bool(item.get("enabled", True)),
            action_id=str(item.get("action_id") or new_action_id()),
        )

    def clone(self) -> "Action":
        return Action(self.kind, copy.deepcopy(self.params), self.enabled)

    def label(self) -> str:
        p = self.params
        labels = {
            "click": "한 번 클릭",
            "double_click": "두 번 클릭",
            "right_click": "오른쪽 클릭",
            "middle_click": "가운데 클릭",
        }
        if self.kind in labels:
            return f"{labels[self.kind]}: ({p.get('x')}, {p.get('y')})"
        if self.kind == "move":
            return (
                f"마우스 이동: ({p.get('x')}, {p.get('y')}) / "
                f"{float(p.get('duration', 0.2)):.2f}초"
            )
        if self.kind == "drag":
            return (
                f"드래그: ({p.get('x1')}, {p.get('y1')}) → "
                f"({p.get('x2')}, {p.get('y2')})"
            )
        if self.kind == "scroll":
            return f"스크롤: {p.get('clicks', 0)}"
        if self.kind == "wait":
            return f"대기: {float(p.get('seconds', 0)):.2f}초"
        if self.kind == "hotkey":
            return "단축키: " + "+".join(str(key).upper() for key in p.get("keys", []))
        if self.kind in {"key", "key_down", "key_up"}:
            name = {"key": "키 입력", "key_down": "키 누름", "key_up": "키 놓음"}[self.kind]
            return f"{name}: {str(p.get('key', '')).upper()}"
        if self.kind == "text":
            template = str(p.get("template", ""))
            preview = template.replace("\n", "↵")
            if len(preview) > 45:
                preview = preview[:42] + "..."
            mode = " [붙여넣기]" if p.get("use_clipboard") else ""
            return f"문자 입력{mode}: {preview}"
        if self.kind == "image_click":
            filename = Path(str(p.get("image", ""))).name
            return (
                f"이미지 찾기/클릭: {filename} "
                f"(정확도 {float(p.get('confidence', 0.85)):.2f}, "
                f"{float(p.get('timeout', 15)):.0f}초)"
            )
        if self.kind == "wait_window":
            return f"창 대기: {p.get('title', '')} ({float(p.get('timeout', 30)):.0f}초)"
        if self.kind == "image_condition_skip":
            filename = Path(str(p.get("image", ""))).name
            condition_text = "있으면" if p.get("condition") == "found" else "없으면"
            return f"조건 분기: {filename} {condition_text} 다음 {p.get('skip_count', 1)}단계 건너뜀"
        return f"{self.kind}: {p}"


def normalize_key_name(raw: str) -> str:
    value = raw.strip().lower().replace("key.", "")
    mapping = {
        "ctrl": "ctrl",
        "ctrl_l": "ctrlleft",
        "ctrl_r": "ctrlright",
        "shift_l": "shiftleft",
        "shift_r": "shiftright",
        "alt_l": "altleft",
        "alt_r": "altright",
        "cmd": "win",
        "cmd_l": "winleft",
        "cmd_r": "winright",
        "return": "enter",
        "escape": "esc",
        "page_up": "pgup",
        "page_down": "pgdn",
        "backspace": "backspace",
        "delete": "delete",
        "space": "space",
    }
    return mapping.get(value, value)


def pynput_key_name(key: Any) -> str | None:
    char = getattr(key, "char", None)
    vk = getattr(key, "vk", None)
    is_control_char = isinstance(char, str) and char != "" and ord(char) < 32
    if (char is None or is_control_char) and isinstance(vk, int):
        # Ctrl을 누른 채 문자 키를 누르면 Windows/pynput이 실제 글자 대신 제어 문자를
        # 돌려주거나(예: Ctrl+A -> chr(1)) char 자체가 비어 오는 경우가 있다. pyautogui는
        # 이런 값을 키 이름으로 인식하지 못해 조용히 아무 동작도 하지 않으므로,
        # 가상 키 코드로 원래 글자를 복원한다.
        if 65 <= vk <= 90:  # A-Z
            return chr(vk).lower()
        if 48 <= vk <= 57:  # 0-9
            return chr(vk)
    if isinstance(char, str) and char:
        if len(char) == 1 and ord(char) < 128:
            return char.lower()
        return char
    text = str(key)
    if text.startswith("Key."):
        return normalize_key_name(text)
    return None


def load_tabular_rows(
    path: Path,
    sheet_name: str,
    header_row: int,
    start_row: int,
    end_row: int | None,
) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"데이터 파일을 찾을 수 없습니다: {path}")
    suffix = path.suffix.lower()
    raw_rows: list[list[Any]] = []

    if suffix in {".csv", ".tsv"}:
        encoding_candidates = ("utf-8-sig", "cp949", "utf-8")
        last_error: Exception | None = None
        for encoding in encoding_candidates:
            try:
                with path.open("r", encoding=encoding, newline="") as stream:
                    if suffix == ".tsv":
                        reader = csv.reader(stream, delimiter="\t")
                    else:
                        sample = stream.read(4096)
                        stream.seek(0)
                        try:
                            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                        except csv.Error:
                            dialect = csv.excel
                        reader = csv.reader(stream, dialect)
                    raw_rows = [list(row) for row in reader]
                break
            except UnicodeDecodeError as error:
                last_error = error
        else:
            raise ValueError(f"CSV 인코딩을 확인하세요: {last_error}")
    elif suffix in {".xlsx", ".xlsm"}:
        if openpyxl is None:
            raise RuntimeError("Excel 파일 사용에는 openpyxl이 필요합니다: py -m pip install openpyxl")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name and sheet_name not in workbook.sheetnames:
                raise ValueError(f"시트를 찾을 수 없습니다: {sheet_name}")
            sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
            raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    else:
        raise ValueError("지원 데이터 형식은 CSV, TSV, XLSX, XLSM입니다.")

    if header_row < 1 or header_row > len(raw_rows):
        raise ValueError("헤더 행 번호가 데이터 범위를 벗어났습니다.")
    headers = [
        str(value).strip() if value not in (None, "") else excel_column_name(index)
        for index, value in enumerate(raw_rows[header_row - 1], start=1)
    ]
    first = max(header_row + 1, start_row)
    last = len(raw_rows) if not end_row else min(len(raw_rows), end_row)
    contexts: list[dict[str, Any]] = []
    for absolute_row in range(first, last + 1):
        values = raw_rows[absolute_row - 1]
        if not any(value not in (None, "") for value in values):
            continue
        context: dict[str, Any] = {"row": absolute_row}
        for index, header in enumerate(headers, start=1):
            value = values[index - 1] if index <= len(values) else ""
            context[header] = value
            context[excel_column_name(index)] = value
        contexts.append(context)
    return contexts


class MultilineDialog(tk.Toplevel):
    def __init__(self, parent: tk.Misc, title: str, prompt: str, initial: str = "") -> None:
        super().__init__(parent)
        self.result: str | None = None
        self.title(title)
        self.transient(parent)
        self.grab_set()
        self.geometry("620x360")
        self.minsize(460, 280)

        frame = ttk.Frame(self, padding=12)
        frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(frame, text=prompt, wraplength=580).pack(anchor=tk.W)
        self.text = tk.Text(frame, wrap=tk.WORD, font=("Malgun Gothic", 11), undo=True)
        self.text.pack(fill=tk.BOTH, expand=True, pady=8)
        self.text.insert("1.0", initial)
        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X)
        ttk.Button(buttons, text="취소", command=self.destroy).pack(side=tk.RIGHT, padx=4)
        ttk.Button(buttons, text="확인", command=self._accept).pack(side=tk.RIGHT)
        self.bind("<Control-Return>", lambda _event: self._accept())
        self.text.focus_set()
        self.wait_window()

    def _accept(self) -> None:
        self.result = self.text.get("1.0", "end-1c")
        self.destroy()


class AutomationApp:
    def __init__(
        self,
        root: tk.Tk,
        startup_profile: Path | None = None,
        autorun: bool = False,
    ) -> None:
        self.root = root
        self.actions: list[Action] = []
        self.stop_event = threading.Event()
        self.pause_event = threading.Event()
        self.running = False
        self.recording = False
        self.record_lock = threading.Lock()
        self.recorded_actions: list[Action] = []
        self.record_last_time = 0.0
        self.record_mouse_moves = False
        self.last_move_time = 0.0
        self.last_move_point: tuple[int, int] | None = None
        self.mouse_press: dict[str, tuple[int, int, float]] = {}
        self.keys_down: set[str] = set()
        self.mouse_listener: Any = None
        self.keyboard_listener: Any = None
        self.current_profile_path: Path | None = None
        self.last_schedule_stamp = ""
        self.log_file: Path | None = None
        self.log_directory: Path | None = None
        self.log_lock = threading.Lock()
        self.run_mutex_handle: Any = None
        self.is_autorun_session = bool(autorun)

        root.title(APP_TITLE)
        root.geometry("1160x790")
        root.minsize(980, 680)
        root.protocol("WM_DELETE_WINDOW", self.on_close)

        self.repeat_var = tk.IntVar(value=1)
        self.step_delay_var = tk.DoubleVar(value=0.25)
        self.retry_count_var = tk.IntVar(value=2)
        self.retry_delay_var = tk.DoubleVar(value=1.0)
        self.status_var = tk.StringVar(value="동작을 추가하거나 [실제 동작 녹화]를 시작하세요.")
        self.record_moves_var = tk.BooleanVar(value=False)
        self.data_enabled_var = tk.BooleanVar(value=False)
        self.data_path_var = tk.StringVar()
        self.sheet_name_var = tk.StringVar()
        self.header_row_var = tk.IntVar(value=1)
        self.start_row_var = tk.IntVar(value=2)
        self.end_row_var = tk.IntVar(value=0)
        self.schedule_enabled_var = tk.BooleanVar(value=False)
        self.schedule_time_var = tk.StringVar(value="09:00")
        self.schedule_daily_var = tk.BooleanVar(value=True)
        self.schedule_minimized_var = tk.BooleanVar(value=True)
        self.schedule_exit_after_var = tk.BooleanVar(value=False)
        self.log_dir_var = tk.StringVar(value=str(default_log_directory()))
        self.log_retention_var = tk.IntVar(value=30)

        self._build_ui()
        self._refresh_list()

        if pyautogui is not None:
            pyautogui.FAILSAFE = True
            pyautogui.PAUSE = 0.03
        else:
            root.after(300, self._show_dependency_error)

        if startup_profile:
            try:
                self._load_profile_path(startup_profile)
            except Exception as error:
                # except 블록을 벗어나면 'error' 바인딩이 사라지므로 지연 실행되는
                # lambda가 참조하기 전에 메시지를 문자열로 미리 저장해 둔다.
                error_text = str(error)
                root.after(300, lambda text=error_text: messagebox.showerror("설정 불러오기 실패", text))
            else:
                if autorun:
                    if self.schedule_minimized_var.get():
                        root.iconify()
                    root.after(800, lambda: self.start_automation(confirm=False))
        self.root.after(1000, self._scheduler_tick)

    def _build_menu(self) -> None:
        menubar = tk.Menu(self.root)
        file_menu = tk.Menu(menubar, tearoff=False)
        file_menu.add_command(label="새로 만들기", accelerator="Ctrl+N", command=self.new_profile)
        file_menu.add_command(label="열기...", accelerator="Ctrl+O", command=self.load_profile)
        file_menu.add_separator()
        file_menu.add_command(label="저장", accelerator="Ctrl+S", command=self.save_profile)
        file_menu.add_command(
            label="다른 이름으로 저장...", accelerator="Ctrl+Shift+S", command=self.save_profile_as
        )
        file_menu.add_separator()
        file_menu.add_command(label="로그 폴더 열기", command=self.open_log_directory)
        file_menu.add_separator()
        file_menu.add_command(label="종료", command=self.on_close)
        menubar.add_cascade(label="파일", menu=file_menu)
        self.root.config(menu=menubar)

        self.root.bind_all("<Control-n>", lambda _event: self.new_profile())
        self.root.bind_all("<Control-o>", lambda _event: self.load_profile())
        self.root.bind_all("<Control-s>", lambda _event: self.save_profile())
        for sequence in ("<Control-Shift-S>", "<Control-Shift-s>"):
            self.root.bind_all(sequence, lambda _event: self.save_profile_as())

    def _build_ui(self) -> None:
        self._build_menu()
        outer = ttk.Frame(self.root, padding=12)
        outer.pack(fill=tk.BOTH, expand=True)
        header = ttk.Frame(outer)
        header.pack(fill=tk.X)
        ttk.Label(header, text=APP_TITLE, font=("Malgun Gothic", 18, "bold")).pack(side=tk.LEFT)
        ttk.Label(
            header,
            text="권한이 있는 업무만 자동화 · F8 녹화 종료 · 마우스 좌상단 긴급정지",
            foreground="#8A3B12",
        ).pack(side=tk.RIGHT)

        self.notebook = ttk.Notebook(outer)
        self.notebook.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        self.action_tab = ttk.Frame(self.notebook, padding=10)
        self.settings_tab = ttk.Frame(self.notebook, padding=10)
        self.data_tab = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.action_tab, text="동작 편집")
        self.notebook.add(self.settings_tab, text="실행 설정")
        self.notebook.add(self.data_tab, text="데이터 · 예약")

        self._build_action_tab()
        self._build_settings_tab()
        self._build_data_tab()

        run_area = ttk.Frame(outer)
        run_area.pack(fill=tk.X, pady=(10, 0))
        ttk.Label(
            run_area,
            textvariable=self.status_var,
            foreground="#174A7E",
            wraplength=650,
        ).pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.run_button = ttk.Button(run_area, text="▶ 자동화 실행", command=self.start_automation)
        self.run_button.pack(side=tk.RIGHT, padx=4)
        self.pause_button = ttk.Button(
            run_area, text="Ⅱ 일시정지", command=self.toggle_pause, state=tk.DISABLED
        )
        self.pause_button.pack(side=tk.RIGHT, padx=4)
        self.stop_button = tk.Button(
            run_area,
            text="■ 긴급 중지",
            command=self.stop_automation,
            bg="#C62828",
            fg="white",
            activebackground="#8E0000",
            activeforeground="white",
            font=("Malgun Gothic", 10, "bold"),
            state=tk.DISABLED,
        )
        self.stop_button.pack(side=tk.RIGHT, padx=4)

    def _build_action_tab(self) -> None:
        toolbar = ttk.LabelFrame(self.action_tab, text="동작 추가", padding=8)
        toolbar.pack(fill=tk.X)
        buttons: list[tuple[str, Callable[[], None]]] = [
            ("● 실제 동작 녹화", self.start_recording),
            ("■ 녹화 종료(F8)", self.stop_recording),
            ("클릭", lambda: self.capture_point("click")),
            ("더블 클릭", lambda: self.capture_point("double_click")),
            ("오른쪽 클릭", lambda: self.capture_point("right_click")),
            ("가운데 클릭", lambda: self.capture_point("middle_click")),
            ("드래그", self.capture_drag),
            ("스크롤", self.add_scroll),
            ("문자 입력", self.add_text),
            ("키/단축키", self.add_key_or_hotkey),
            ("이미지 찾기", self.add_image_action),
            ("화면 캡처로 이미지 등록", self.add_image_action_from_capture),
            ("조건 분기(이미지)", self.add_image_condition_skip),
            ("창 대기", self.add_window_action),
            ("대기", self.add_wait),
        ]
        self.record_start_button: ttk.Button | None = None
        self.record_stop_button: ttk.Button | None = None
        for index, (text, command) in enumerate(buttons):
            button = ttk.Button(toolbar, text=text, command=command)
            button.grid(row=index // 5, column=index % 5, sticky="ew", padx=3, pady=3)
            toolbar.columnconfigure(index % 5, weight=1)
            if index == 0:
                self.record_start_button = button
            elif index == 1:
                self.record_stop_button = button
                button.configure(state=tk.DISABLED)
        checkbox_row = (len(buttons) + 4) // 5
        ttk.Checkbutton(
            toolbar,
            text="녹화 중 마우스 이동도 기록(단계 수 증가)",
            variable=self.record_moves_var,
        ).grid(row=checkbox_row, column=0, columnspan=5, sticky=tk.W, pady=(4, 0))

        list_frame = ttk.LabelFrame(self.action_tab, text="실행 순서", padding=8)
        list_frame.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)
        self.action_tree = ttk.Treeview(
            list_frame,
            columns=("number", "enabled", "description"),
            show="headings",
            selectmode="browse",
        )
        self.action_tree.heading("number", text="순서")
        self.action_tree.heading("enabled", text="사용")
        self.action_tree.heading("description", text="동작")
        self.action_tree.column("number", width=55, anchor=tk.CENTER, stretch=False)
        self.action_tree.column("enabled", width=55, anchor=tk.CENTER, stretch=False)
        self.action_tree.column("description", width=760)
        self.action_tree.grid(row=0, column=0, sticky="nsew")
        self.action_tree.bind("<Double-1>", lambda _event: self.edit_selected())
        scrollbar = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.action_tree.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.action_tree.configure(yscrollcommand=scrollbar.set)

        edit_buttons = ttk.Frame(list_frame)
        edit_buttons.grid(row=0, column=2, sticky="ns", padx=(8, 0))
        for text, command in [
            ("편집", self.edit_selected),
            ("복제", self.duplicate_selected),
            ("사용/제외", self.toggle_selected),
            ("위로", self.move_up),
            ("아래로", self.move_down),
            ("선택 삭제", self.delete_selected),
            ("전체 삭제", self.clear_all),
            ("▶ 선택 단계부터 실행", self.test_selected_action),
        ]:
            ttk.Button(edit_buttons, text=text, command=command).pack(fill=tk.X, pady=2)

    def _build_settings_tab(self) -> None:
        execution = ttk.LabelFrame(self.settings_tab, text="반복 · 속도 · 오류 복구", padding=12)
        execution.pack(fill=tk.X)
        fields = [
            ("전체 반복 횟수", self.repeat_var, 1, 10000, 1),
            ("단계 사이 대기(초)", self.step_delay_var, 0, 30, 0.05),
            ("실패 재시도 횟수", self.retry_count_var, 0, 20, 1),
            ("재시도 대기(초)", self.retry_delay_var, 0, 300, 0.5),
        ]
        for row, (label, variable, low, high, increment) in enumerate(fields):
            ttk.Label(execution, text=label).grid(row=row, column=0, sticky=tk.W, padx=4, pady=6)
            ttk.Spinbox(
                execution,
                from_=low,
                to=high,
                increment=increment,
                width=12,
                textvariable=variable,
            ).grid(row=row, column=1, sticky=tk.W, padx=4, pady=6)

        files = ttk.LabelFrame(self.settings_tab, text="설정 파일 · 실행 기록", padding=12)
        files.pack(fill=tk.X, pady=(12, 0))
        ttk.Button(files, text="저장(Ctrl+S)", command=self.save_profile).grid(
            row=0, column=0, padx=4
        )
        ttk.Button(files, text="다른 이름으로 저장...(Ctrl+Shift+S)", command=self.save_profile_as).grid(
            row=0, column=1, padx=4
        )
        ttk.Button(files, text="불러오기...(Ctrl+O)", command=self.load_profile).grid(
            row=0, column=2, padx=4
        )
        ttk.Button(files, text="로그 폴더 열기", command=self.open_log_directory).grid(
            row=0, column=3, padx=4
        )
        ttk.Label(files, text="로그 폴더").grid(row=1, column=0, sticky=tk.W, padx=4, pady=(12, 4))
        ttk.Entry(files, textvariable=self.log_dir_var).grid(
            row=1, column=1, columnspan=3, sticky="ew", padx=4, pady=(12, 4)
        )
        ttk.Button(files, text="찾기", command=self.choose_log_directory).grid(
            row=1, column=4, padx=4, pady=(12, 4)
        )
        files.columnconfigure(3, weight=1)
        ttk.Label(files, text="로그 보관 일수(0=무제한)").grid(
            row=2, column=0, sticky=tk.W, padx=4, pady=(8, 4)
        )
        ttk.Spinbox(
            files, from_=0, to=3650, width=8, textvariable=self.log_retention_var
        ).grid(row=2, column=1, sticky=tk.W, padx=4, pady=(8, 4))
        ttk.Label(
            files,
            text=(
                "로그에는 단계 번호·성공/실패만 기록합니다. 실패 시 화면 스크린샷이 저장될 수 있으므로 "
                "회사 보안 정책에 따라 로그 폴더를 관리하세요. 실행 시작 시 보관 기간이 지난 로그·스크린샷은 "
                "자동으로 정리됩니다."
            ),
            foreground="#8A3B12",
            wraplength=850,
        ).grid(row=3, column=0, columnspan=5, sticky=tk.W, padx=4, pady=(8, 0))

    def _build_data_tab(self) -> None:
        data = ttk.LabelFrame(self.data_tab, text="CSV/Excel 행별 반복", padding=12)
        data.pack(fill=tk.X)
        ttk.Checkbutton(
            data, text="데이터 파일의 각 행마다 전체 동작 실행", variable=self.data_enabled_var
        ).grid(row=0, column=0, columnspan=4, sticky=tk.W, pady=4)
        ttk.Label(data, text="데이터 파일").grid(row=1, column=0, sticky=tk.W, padx=4, pady=4)
        ttk.Entry(data, textvariable=self.data_path_var).grid(
            row=1, column=1, columnspan=2, sticky="ew", padx=4, pady=4
        )
        ttk.Button(data, text="찾기", command=self.choose_data_file).grid(row=1, column=3, padx=4)
        ttk.Label(data, text="시트명(비우면 첫 시트)").grid(row=2, column=0, sticky=tk.W, padx=4, pady=4)
        ttk.Entry(data, textvariable=self.sheet_name_var, width=20).grid(
            row=2, column=1, sticky=tk.W, padx=4, pady=4
        )
        ttk.Label(data, text="헤더 행").grid(row=2, column=2, sticky=tk.E, padx=4)
        ttk.Spinbox(data, from_=1, to=10000, width=8, textvariable=self.header_row_var).grid(
            row=2, column=3, sticky=tk.W, padx=4
        )
        ttk.Label(data, text="시작 행").grid(row=3, column=0, sticky=tk.W, padx=4, pady=4)
        ttk.Spinbox(data, from_=1, to=1000000, width=10, textvariable=self.start_row_var).grid(
            row=3, column=1, sticky=tk.W, padx=4
        )
        ttk.Label(data, text="종료 행(0=끝까지)").grid(row=3, column=2, sticky=tk.E, padx=4)
        ttk.Spinbox(data, from_=0, to=1000000, width=10, textvariable=self.end_row_var).grid(
            row=3, column=3, sticky=tk.W, padx=4
        )
        ttk.Button(data, text="데이터 미리 확인", command=self.preview_data).grid(
            row=4, column=0, padx=4, pady=(10, 4), sticky=tk.W
        )
        ttk.Label(
            data,
            text=(
                "문자 입력 단계에서 {열이름}, {A}, {B}, {row} 형식을 사용하세요. "
                "예: {설비명} / Lot {A} / 원본 행 {row}"
            ),
            foreground="#174A7E",
            wraplength=850,
        ).grid(row=4, column=1, columnspan=3, sticky=tk.W, padx=4, pady=(10, 4))
        data.columnconfigure(1, weight=1)

        schedule = ttk.LabelFrame(self.data_tab, text="예약 실행", padding=12)
        schedule.pack(fill=tk.X, pady=(12, 0))
        ttk.Checkbutton(
            schedule,
            text="프로그램이 실행 중일 때 예약 시간에 자동 실행",
            variable=self.schedule_enabled_var,
        ).grid(row=0, column=0, columnspan=2, sticky=tk.W, pady=4)
        ttk.Label(schedule, text="실행 시각(HH:MM)").grid(row=1, column=0, sticky=tk.W, padx=4)
        ttk.Entry(schedule, textvariable=self.schedule_time_var, width=10).grid(
            row=1, column=1, sticky=tk.W, padx=4
        )
        ttk.Checkbutton(schedule, text="매일 반복", variable=self.schedule_daily_var).grid(
            row=1, column=2, padx=10
        )
        ttk.Checkbutton(
            schedule,
            text="Windows 작업 스케줄러 실행 시 창 최소화로 시작",
            variable=self.schedule_minimized_var,
        ).grid(row=2, column=0, columnspan=2, sticky=tk.W, pady=4)
        ttk.Checkbutton(
            schedule,
            text="실행 완료 후 프로그램 자동 종료",
            variable=self.schedule_exit_after_var,
        ).grid(row=2, column=2, columnspan=2, sticky=tk.W, pady=4)
        ttk.Button(
            schedule, text="Windows 작업 스케줄러 등록", command=self.register_windows_task
        ).grid(row=3, column=0, padx=4, pady=(12, 4), sticky=tk.W)
        ttk.Button(
            schedule, text="Windows 예약 삭제", command=self.delete_windows_task
        ).grid(row=3, column=1, padx=4, pady=(12, 4), sticky=tk.W)
        ttk.Label(
            schedule,
            text=(
                "내부 예약은 프로그램을 켜 둬야 합니다. Windows 작업 스케줄러 등록은 저장된 설정을 "
                "사용해 프로그램이 닫혀 있어도 지정 시각에 실행합니다. 무인 실행을 위해 창 최소화와 "
                "완료 후 자동 종료를 함께 켜 두는 것을 권장합니다."
            ),
            wraplength=850,
            foreground="#555555",
        ).grid(row=4, column=0, columnspan=4, sticky=tk.W, padx=4, pady=(4, 0))

    def _show_dependency_error(self) -> None:
        messagebox.showerror(
            "필수 모듈 없음",
            "pyautogui가 필요합니다.\n\n명령 프롬프트:\npy -m pip install pyautogui pynput pillow",
        )

    def _set_status(self, text: str) -> None:
        self.root.after(0, self.status_var.set, text)

    def _selected_action_index(self) -> int | None:
        selected = self.action_tree.selection()
        if not selected:
            return None
        action_id = selected[0]
        return next((index for index, action in enumerate(self.actions) if action.action_id == action_id), None)

    def _refresh_list(self, selected_index: int | None = None) -> None:
        previous = self.action_tree.selection()
        previous_id = previous[0] if previous else None
        for item in self.action_tree.get_children():
            self.action_tree.delete(item)
        for index, action in enumerate(self.actions, start=1):
            self.action_tree.insert(
                "",
                tk.END,
                iid=action.action_id,
                values=(index, "✓" if action.enabled else "—", action.label()),
            )
        target_id: str | None = None
        if selected_index is not None and self.actions:
            selected_index = max(0, min(selected_index, len(self.actions) - 1))
            target_id = self.actions[selected_index].action_id
        elif previous_id and any(action.action_id == previous_id for action in self.actions):
            target_id = previous_id
        if target_id:
            self.action_tree.selection_set(target_id)
            self.action_tree.see(target_id)

    def _insert_action(self, action: Action) -> None:
        """새 동작을 추가한다. 목록에서 선택된 단계가 있으면 그 바로 아래에 끼워 넣고,
        선택된 단계가 없으면 맨 뒤에 추가한다."""
        if len(self.actions) >= MAX_ACTIONS:
            messagebox.showwarning("단계 제한", f"실행 단계는 최대 {MAX_ACTIONS:,}개입니다.")
            return
        index = self._selected_action_index()
        insert_at = len(self.actions) if index is None else index + 1
        self.actions.insert(insert_at, action)
        self._refresh_list(insert_at)

    def _capture_position_after_countdown(self, purpose: str) -> tuple[int, int] | None:
        if pyautogui is None:
            self._show_dependency_error()
            return None
        messagebox.showinfo("좌표 등록", f"확인을 누른 뒤 3초 안에\n{purpose} 위치로 이동하세요.")
        self.root.iconify()
        for remaining in (3, 2, 1):
            self._set_status(f"{remaining}초 후 마우스 위치를 등록합니다...")
            deadline = time.monotonic() + 1.0
            while time.monotonic() < deadline:
                self.root.update()
                time.sleep(0.05)
        point = pyautogui.position()
        self.root.deiconify()
        self.root.lift()
        self._set_status(f"좌표 등록: ({point.x}, {point.y})")
        return int(point.x), int(point.y)

    def capture_point(self, kind: str) -> None:
        point = self._capture_position_after_countdown("클릭할")
        if point is None:
            return
        self._insert_action(Action(kind, {"x": point[0], "y": point[1]}))

    def capture_drag(self) -> None:
        start = self._capture_position_after_countdown("드래그를 시작할")
        if start is None:
            return
        end = self._capture_position_after_countdown("드래그를 끝낼")
        if end is None:
            return
        self._insert_action(
            Action(
                "drag",
                {"x1": start[0], "y1": start[1], "x2": end[0], "y2": end[1], "duration": 0.6},
            )
        )

    def add_scroll(self) -> None:
        clicks = simpledialog.askinteger(
            "스크롤 추가",
            "스크롤 양을 입력하세요.\n양수=위, 음수=아래",
            initialvalue=-5,
            minvalue=-10000,
            maxvalue=10000,
            parent=self.root,
        )
        if clicks is None:
            return
        self._insert_action(Action("scroll", {"clicks": clicks}))

    def add_text(self) -> None:
        dialog = MultilineDialog(
            self.root,
            "문자 입력",
            "입력할 문자를 작성하세요. Excel/CSV 사용 시 {열이름}, {A}, {row}를 사용할 수 있습니다.",
        )
        if dialog.result is None:
            return
        interval = simpledialog.askfloat(
            "입력 속도",
            "문자 사이 대기시간(초)",
            initialvalue=0.01,
            minvalue=0,
            maxvalue=5,
            parent=self.root,
        )
        if interval is None:
            # 입력 속도 창을 취소해도 이미 작성한 문자열은 버리지 않고 기본 속도로 등록한다.
            interval = 0.01
        use_clipboard = messagebox.askyesno(
            "입력 방식 선택",
            "일부 프로그램은 직접 키 입력 방식으로 한글 등이 제대로 들어가지 않을 수 있습니다.\n"
            "클립보드에 붙여넣기(Ctrl+V)로 입력할까요?\n"
            "(붙여넣기 직후 원래 클립보드 내용을 복원합니다)\n\n"
            "잘 모르겠으면 '아니요'를 선택하세요.",
        )
        self._insert_action(
            Action(
                "text",
                {"template": dialog.result, "interval": interval, "use_clipboard": use_clipboard},
            )
        )

    def add_key_or_hotkey(self) -> None:
        value = simpledialog.askstring(
            "키/단축키 추가",
            "키 또는 단축키를 입력하세요.\n예: enter, tab, f5, ctrl+shift+s, win+d",
            parent=self.root,
        )
        if not value:
            return
        keys = [normalize_key_name(item) for item in value.split("+") if item.strip()]
        if not keys:
            return
        action = Action("key", {"key": keys[0]}) if len(keys) == 1 else Action("hotkey", {"keys": keys})
        self._insert_action(action)

    def add_wait(self) -> None:
        seconds = simpledialog.askfloat(
            "대기시간 추가",
            "대기할 시간을 초 단위로 입력하세요.",
            initialvalue=2.0,
            minvalue=0,
            maxvalue=86400,
            parent=self.root,
        )
        if seconds is None:
            return
        self._insert_action(Action("wait", {"seconds": seconds}))

    def _prompt_image_confidence(self, initial: dict[str, Any]) -> float:
        confidence = simpledialog.askfloat(
            "이미지 정확도",
            "정확도(0.10~1.00)\n0.80~0.90 권장",
            initialvalue=float(initial.get("confidence", 0.85)),
            minvalue=0.1,
            maxvalue=1.0,
            parent=self.root,
        )
        return float(initial.get("confidence", 0.85)) if confidence is None else confidence

    def _prompt_image_settings(self, initial: dict[str, Any]) -> dict[str, Any]:
        """이미지 파일이 정해진 뒤 정확도/대기시간/클릭횟수를 물어본다.
        중간에 취소해도 값을 버리지 않고 이전 값(또는 기본값)으로 채운다."""
        confidence = self._prompt_image_confidence(initial)
        timeout = simpledialog.askfloat(
            "최대 대기시간",
            "이미지를 찾을 최대 시간(초)",
            initialvalue=float(initial.get("timeout", 15)),
            minvalue=0,
            maxvalue=86400,
            parent=self.root,
        )
        if timeout is None:
            timeout = float(initial.get("timeout", 15))
        clicks = simpledialog.askinteger(
            "클릭 횟수",
            "찾은 이미지 중심을 몇 번 클릭할까요? (0=찾기만)",
            initialvalue=int(initial.get("clicks", 1)),
            minvalue=0,
            maxvalue=3,
            parent=self.root,
        )
        if clicks is None:
            clicks = int(initial.get("clicks", 1))
        return {"confidence": confidence, "timeout": timeout, "clicks": clicks}

    def _prompt_image_params(self, initial: dict[str, Any] | None = None) -> dict[str, Any] | None:
        initial = initial or {}
        image_path = filedialog.askopenfilename(
            title="찾을 이미지 선택",
            initialfile=Path(str(initial.get("image", ""))).name,
            filetypes=[("이미지", "*.png *.jpg *.jpeg *.bmp"), ("모든 파일", "*.*")],
        )
        if not image_path:
            return None
        settings = self._prompt_image_settings(initial)
        return {"image": image_path, **settings}

    def add_image_action(self) -> None:
        params = self._prompt_image_params()
        if params is None:
            return
        self._insert_action(Action("image_click", params))

    def _capture_screen_region_to_file(self) -> Path | None:
        """화면 위에서 사각형을 드래그로 지정해 그 영역만 이미지 파일로 저장한다."""
        if pyautogui is None:
            self._show_dependency_error()
            return None
        if not messagebox.askyesno(
            "화면 캡처로 이미지 등록",
            "확인을 누르면 창이 최소화되고 화면 전체가 반투명하게 덮입니다.\n"
            "찾을 대상 위에서 마우스를 드래그해 사각형으로 선택한 뒤 손을 떼세요.\n"
            "Esc를 누르면 취소합니다.",
        ):
            return None
        self.root.iconify()
        self.root.update()
        time.sleep(0.3)

        overlay = tk.Toplevel(self.root)
        overlay.attributes("-fullscreen", True)
        try:
            overlay.attributes("-alpha", 0.25)
        except tk.TclError:
            pass
        overlay.attributes("-topmost", True)
        overlay.configure(bg="black")
        canvas = tk.Canvas(overlay, cursor="crosshair", bg="black", highlightthickness=0)
        canvas.pack(fill=tk.BOTH, expand=True)

        start: dict[str, int] = {}
        state: dict[str, Any] = {"rect": None, "region": None, "cancelled": False}

        def on_press(event: tk.Event) -> None:
            start["x"], start["y"] = event.x, event.y
            state["rect"] = canvas.create_rectangle(
                event.x, event.y, event.x, event.y, outline="#FF3B30", width=2
            )

        def on_drag(event: tk.Event) -> None:
            if state["rect"] is not None:
                canvas.coords(state["rect"], start["x"], start["y"], event.x, event.y)

        def on_release(event: tk.Event) -> None:
            if "x" in start:
                state["region"] = (start["x"], start["y"], event.x, event.y)
            overlay.destroy()

        def on_escape(_event: tk.Event) -> None:
            state["cancelled"] = True
            overlay.destroy()

        canvas.bind("<ButtonPress-1>", on_press)
        canvas.bind("<B1-Motion>", on_drag)
        canvas.bind("<ButtonRelease-1>", on_release)
        overlay.bind("<Escape>", on_escape)
        overlay.focus_force()
        overlay.grab_set()
        overlay.wait_window()

        self.root.deiconify()
        self.root.lift()

        if state["cancelled"] or state["region"] is None:
            return None
        x1, y1, x2, y2 = state["region"]
        left, top = min(x1, x2), min(y1, y2)
        width, height = abs(x2 - x1), abs(y2 - y1)
        if width < 5 or height < 5:
            messagebox.showwarning("캡처 실패", "선택 영역이 너무 작습니다.")
            return None

        save_path = filedialog.asksaveasfilename(
            title="캡처한 이미지 저장",
            defaultextension=".png",
            filetypes=[("PNG 이미지", "*.png")],
            initialfile=f"capture_{datetime.now():%Y%m%d_%H%M%S}.png",
        )
        if not save_path:
            return None
        try:
            image = pyautogui.screenshot(region=(left, top, width, height))
            image.save(save_path)
        except Exception as error:
            messagebox.showerror("캡처 실패", str(error))
            return None
        return Path(save_path)

    def add_image_action_from_capture(self) -> None:
        image_path = self._capture_screen_region_to_file()
        if image_path is None:
            return
        settings = self._prompt_image_settings({})
        self._insert_action(Action("image_click", {"image": str(image_path), **settings}))

    def _prompt_image_condition_params(
        self, initial: dict[str, Any] | None = None
    ) -> dict[str, Any] | None:
        """이미지가 있는지/없는지에 따라 다음 몇 단계를 건너뛰는 조건 분기 정보를 입력받는다."""
        initial = initial or {}
        image_path = filedialog.askopenfilename(
            title="조건으로 사용할 이미지 선택",
            initialfile=Path(str(initial.get("image", ""))).name,
            filetypes=[("이미지", "*.png *.jpg *.jpeg *.bmp"), ("모든 파일", "*.*")],
        )
        if not image_path:
            return None
        confidence = self._prompt_image_confidence(initial)
        condition_found = messagebox.askyesno(
            "조건 종류",
            "화면에 이 이미지가 '있을 때' 건너뛸까요?\n"
            "'아니요'를 선택하면 이미지가 '없을 때' 건너뜁니다.",
        )
        skip_count = simpledialog.askinteger(
            "건너뛸 단계 수",
            "조건이 맞을 때 건너뛸 다음 단계 수를 입력하세요.",
            initialvalue=int(initial.get("skip_count", 1)),
            minvalue=1,
            maxvalue=MAX_ACTIONS,
            parent=self.root,
        )
        if skip_count is None:
            return None
        return {
            "image": image_path,
            "confidence": confidence,
            "condition": "found" if condition_found else "not_found",
            "skip_count": skip_count,
        }

    def add_image_condition_skip(self) -> None:
        params = self._prompt_image_condition_params()
        if params is None:
            return
        self._insert_action(Action("image_condition_skip", params))

    def _prompt_window_params(self, initial: dict[str, Any] | None = None) -> dict[str, Any] | None:
        initial = initial or {}
        title = simpledialog.askstring(
            "창 대기",
            "기다릴 창 제목의 일부를 입력하세요.",
            initialvalue=str(initial.get("title", "")),
            parent=self.root,
        )
        if not title:
            return None
        timeout = simpledialog.askfloat(
            "창 대기",
            "최대 대기시간(초)",
            initialvalue=float(initial.get("timeout", 30)),
            minvalue=0,
            maxvalue=86400,
            parent=self.root,
        )
        if timeout is None:
            # 대기시간 창을 취소해도 이미 입력한 제목은 버리지 않고 기본값으로 이어간다.
            timeout = float(initial.get("timeout", 30))
        regex = messagebox.askyesno(
            "창 제목 방식",
            "창 제목을 정규식으로 검색할까요?\n'아니요'를 선택하면 일부 문자열로 검색합니다.",
            parent=self.root,
        )
        activate = messagebox.askyesno(
            "창 활성화",
            "창을 찾은 뒤 앞으로 가져올까요?",
            parent=self.root,
        )
        return {"title": title, "timeout": timeout, "regex": regex, "activate": activate}

    def add_window_action(self) -> None:
        if not WINDOWS:
            messagebox.showwarning("Windows 전용", "창 제목 대기는 Windows에서 지원됩니다.")
            return
        params = self._prompt_window_params()
        if params is None:
            return
        self._insert_action(Action("wait_window", params))

    def start_recording(self) -> None:
        if self.running or self.recording:
            return
        if pynput_keyboard is None or pynput_mouse is None:
            messagebox.showerror(
                "녹화 모듈 없음",
                "실제 동작 녹화에는 pynput이 필요합니다.\n\npy -m pip install pynput",
            )
            return
        if not messagebox.askyesno(
            "실제 동작 녹화",
            "확인을 누르면 프로그램 창이 최소화되고 2초 후 녹화합니다.\n"
            "마우스 클릭·드래그·스크롤과 키보드 입력을 기록합니다.\n\n"
            "녹화를 끝내려면 F8을 누르세요.",
        ):
            return
        with self.record_lock:
            self.recorded_actions = []
            self.record_last_time = time.monotonic() + 2.0
            self.record_mouse_moves = bool(self.record_moves_var.get())
            self.last_move_time = 0.0
            self.last_move_point = None
            self.mouse_press.clear()
            self.keys_down.clear()
        self.recording = True
        if self.record_start_button:
            self.record_start_button.configure(state=tk.DISABLED)
        if self.record_stop_button:
            self.record_stop_button.configure(state=tk.NORMAL)
        self.root.iconify()
        self._set_status("2초 후 실제 동작 녹화를 시작합니다. 종료는 F8입니다.")

        self.mouse_listener = pynput_mouse.Listener(
            on_move=self._record_on_move,
            on_click=self._record_on_click,
            on_scroll=self._record_on_scroll,
        )
        self.keyboard_listener = pynput_keyboard.Listener(
            on_press=self._record_on_key_press,
            on_release=self._record_on_key_release,
        )
        self.mouse_listener.start()
        self.keyboard_listener.start()

    def _append_recorded(self, action: Action, now: float | None = None) -> None:
        if not self.recording:
            return
        now = now or time.monotonic()
        with self.record_lock:
            if now < self.record_last_time:
                return
            delay = now - self.record_last_time
            if delay >= 0.12:
                self.recorded_actions.append(Action("wait", {"seconds": round(min(delay, 3600), 3)}))
            self.recorded_actions.append(action)
            self.record_last_time = now

    def _record_on_move(self, x: int, y: int) -> None:
        if not self.recording or not self.record_mouse_moves:
            return
        now = time.monotonic()
        point = (int(x), int(y))
        if now - self.last_move_time < 0.12:
            return
        if self.last_move_point:
            distance = abs(point[0] - self.last_move_point[0]) + abs(point[1] - self.last_move_point[1])
            if distance < 20:
                return
        self.last_move_time = now
        self.last_move_point = point
        self._append_recorded(Action("move", {"x": point[0], "y": point[1], "duration": 0.1}), now)

    def _record_on_click(self, x: int, y: int, button: Any, pressed: bool) -> None:
        if not self.recording:
            return
        now = time.monotonic()
        name = str(button).replace("Button.", "")
        if pressed:
            self.mouse_press[name] = (int(x), int(y), now)
            return
        start = self.mouse_press.pop(name, (int(x), int(y), now))
        distance = abs(int(x) - start[0]) + abs(int(y) - start[1])
        if name == "left" and distance >= 8:
            self._append_recorded(
                Action(
                    "drag",
                    {
                        "x1": start[0],
                        "y1": start[1],
                        "x2": int(x),
                        "y2": int(y),
                        "duration": max(0.1, round(now - start[2], 3)),
                    },
                ),
                now,
            )
        else:
            kind = {"right": "right_click", "middle": "middle_click"}.get(name, "click")
            self._append_recorded(Action(kind, {"x": int(x), "y": int(y)}), now)

    def _record_on_scroll(self, _x: int, _y: int, _dx: int, dy: int) -> None:
        if self.recording:
            self._append_recorded(Action("scroll", {"clicks": int(dy)}))

    def _record_on_key_press(self, key: Any) -> bool | None:
        if not self.recording:
            return False
        name = pynput_key_name(key)
        if name == STOP_RECORD_KEY:
            self.root.after(0, self.stop_recording)
            return False
        if not name or name in self.keys_down:
            return None
        self.keys_down.add(name)
        self._append_recorded(Action("key_down", {"key": name}))
        return None

    def _record_on_key_release(self, key: Any) -> bool | None:
        if not self.recording:
            return False
        name = pynput_key_name(key)
        if name == STOP_RECORD_KEY:
            return False
        if name:
            self.keys_down.discard(name)
            self._append_recorded(Action("key_up", {"key": name}))
        return None

    def stop_recording(self) -> None:
        if not self.recording:
            return
        self.recording = False
        for listener in (self.mouse_listener, self.keyboard_listener):
            if listener is not None:
                try:
                    listener.stop()
                except Exception:
                    pass
        self.mouse_listener = None
        self.keyboard_listener = None
        self.root.deiconify()
        self.root.lift()
        if self.record_start_button:
            self.record_start_button.configure(state=tk.NORMAL)
        if self.record_stop_button:
            self.record_stop_button.configure(state=tk.DISABLED)
        with self.record_lock:
            recorded = self._clean_recorded_actions(self.recorded_actions)
        remaining = MAX_ACTIONS - len(self.actions)
        if len(recorded) > remaining:
            recorded = recorded[:remaining]
            messagebox.showwarning("단계 제한", f"최대 {MAX_ACTIONS:,}단계까지만 추가했습니다.")
        # 수동 동작 추가와 마찬가지로, 선택된 단계가 있으면 그 바로 아래에 녹화 결과를 끼워 넣는다.
        index = self._selected_action_index()
        insert_at = len(self.actions) if index is None else index + 1
        self.actions[insert_at:insert_at] = recorded
        if recorded:
            self._refresh_list(insert_at + len(recorded) - 1)
        else:
            self._refresh_list()
        self.status_var.set(f"녹화 완료: {len(recorded)}개 단계가 추가되었습니다.")

    @staticmethod
    def _clean_recorded_actions(actions: list[Action]) -> list[Action]:
        result: list[Action] = []
        for action in actions:
            if action.kind == "wait" and float(action.params.get("seconds", 0)) < 0.12:
                continue
            if (
                action.kind == "click"
                and result
                and result[-1].kind == "click"
                and abs(int(action.params["x"]) - int(result[-1].params["x"])) <= 4
                and abs(int(action.params["y"]) - int(result[-1].params["y"])) <= 4
            ):
                result[-1] = Action("double_click", dict(action.params))
                continue
            result.append(action.clone())
        while result and result[0].kind == "wait":
            result.pop(0)
        return result

    def edit_selected(self) -> None:
        try:
            self._edit_selected()
        except (TypeError, ValueError, KeyError) as error:
            messagebox.showerror("단계 편집 오류", str(error))

    def _edit_selected(self) -> None:
        index = self._selected_action_index()
        if index is None:
            return
        action = self.actions[index]
        p = action.params
        updated: dict[str, Any] | None = None
        if action.kind in {"click", "double_click", "right_click", "middle_click", "move"}:
            value = simpledialog.askstring(
                "좌표 편집",
                "x,y 형식으로 입력하세요.",
                initialvalue=f"{p.get('x', 0)},{p.get('y', 0)}",
                parent=self.root,
            )
            if value:
                x_text, y_text = value.split(",", 1)
                updated = dict(p, x=int(x_text.strip()), y=int(y_text.strip()))
        elif action.kind == "drag":
            value = simpledialog.askstring(
                "드래그 편집",
                "x1,y1,x2,y2,시간 형식으로 입력하세요.",
                initialvalue=(
                    f"{p.get('x1', 0)},{p.get('y1', 0)},"
                    f"{p.get('x2', 0)},{p.get('y2', 0)},{p.get('duration', 0.6)}"
                ),
                parent=self.root,
            )
            if value:
                values = [part.strip() for part in value.split(",")]
                if len(values) != 5:
                    raise ValueError("5개 값을 입력해야 합니다.")
                updated = {
                    "x1": int(values[0]),
                    "y1": int(values[1]),
                    "x2": int(values[2]),
                    "y2": int(values[3]),
                    "duration": float(values[4]),
                }
        elif action.kind == "wait":
            value = simpledialog.askfloat(
                "대기 편집",
                "대기시간(초)",
                initialvalue=float(p.get("seconds", 1)),
                minvalue=0,
                maxvalue=86400,
                parent=self.root,
            )
            if value is not None:
                updated = {"seconds": value}
        elif action.kind == "scroll":
            value = simpledialog.askinteger(
                "스크롤 편집",
                "스크롤 양",
                initialvalue=int(p.get("clicks", 0)),
                minvalue=-10000,
                maxvalue=10000,
                parent=self.root,
            )
            if value is not None:
                updated = {"clicks": value}
        elif action.kind in {"key", "key_down", "key_up"}:
            value = simpledialog.askstring(
                "키 편집",
                "키 이름",
                initialvalue=str(p.get("key", "")),
                parent=self.root,
            )
            if value:
                updated = {"key": normalize_key_name(value)}
        elif action.kind == "hotkey":
            value = simpledialog.askstring(
                "단축키 편집",
                "예: ctrl+shift+s",
                initialvalue="+".join(str(item) for item in p.get("keys", [])),
                parent=self.root,
            )
            if value:
                updated = {"keys": [normalize_key_name(item) for item in value.split("+") if item.strip()]}
        elif action.kind == "text":
            dialog = MultilineDialog(
                self.root,
                "문자 입력 편집",
                "입력 문자열 또는 데이터 템플릿을 수정하세요.",
                str(p.get("template", "")),
            )
            if dialog.result is not None:
                use_clipboard = messagebox.askyesno(
                    "입력 방식 선택",
                    "클립보드에 붙여넣기(Ctrl+V)로 입력할까요?\n"
                    "한글 등에서 직접 키 입력이 제대로 안 될 때 선택하세요.",
                )
                updated = dict(p, template=dialog.result, use_clipboard=use_clipboard)
        elif action.kind == "image_click":
            updated = self._prompt_image_params(p)
        elif action.kind == "wait_window":
            updated = self._prompt_window_params(p)
        elif action.kind == "image_condition_skip":
            updated = self._prompt_image_condition_params(p)
        if updated is not None:
            action.params = updated
            self._refresh_list(index)

    def duplicate_selected(self) -> None:
        index = self._selected_action_index()
        if index is None:
            return
        self.actions.insert(index + 1, self.actions[index].clone())
        self._refresh_list(index + 1)

    def toggle_selected(self) -> None:
        index = self._selected_action_index()
        if index is None:
            return
        self.actions[index].enabled = not self.actions[index].enabled
        self._refresh_list(index)

    def move_up(self) -> None:
        index = self._selected_action_index()
        if index is None or index == 0:
            return
        self.actions[index - 1], self.actions[index] = self.actions[index], self.actions[index - 1]
        self._refresh_list(index - 1)

    def move_down(self) -> None:
        index = self._selected_action_index()
        if index is None or index >= len(self.actions) - 1:
            return
        self.actions[index + 1], self.actions[index] = self.actions[index], self.actions[index + 1]
        self._refresh_list(index + 1)

    def delete_selected(self) -> None:
        index = self._selected_action_index()
        if index is None:
            return
        del self.actions[index]
        self._refresh_list(index)

    def clear_all(self) -> None:
        if self.actions and messagebox.askyesno("전체 삭제", "등록된 실행 단계를 모두 삭제할까요?"):
            self.actions.clear()
            self._refresh_list()

    def test_selected_action(self) -> None:
        """선택한 단계부터 끝까지 순서대로 실행한다.
        각 단계가 성공하면 자동으로 다음 단계로 이어지고, 실패하면 그 자리에서 멈춘다
        (일시정지/긴급 중지, 재시도, 실패 스크린샷 등 실제 실행과 동일한 엔진을 사용한다)."""
        if pyautogui is None:
            self._show_dependency_error()
            return
        if self.running or self.recording:
            messagebox.showwarning("실행 중", "자동화 실행/녹화 중에는 시작할 수 없습니다.")
            return
        index = self._selected_action_index()
        if index is None:
            messagebox.showwarning("선택 없음", "시작할 단계를 목록에서 선택하세요.")
            return
        remaining = [action.clone() for action in self.actions[index:] if action.enabled]
        if not remaining:
            messagebox.showwarning("실행할 단계 없음", "선택한 위치부터 사용 상태인 단계가 없습니다.")
            return
        try:
            self._validate_actions(remaining)
        except Exception as error:
            messagebox.showerror("단계 오류", str(error))
            return
        try:
            context = self._get_data_rows()[0]
        except Exception:
            context = {}
        if not messagebox.askyesno(
            "선택 단계부터 실행",
            f"'{self.actions[index].label()}' 단계부터 끝까지 {len(remaining)}개 단계를 실행합니다.\n"
            "각 단계가 성공하면 자동으로 다음 단계로 이어지고, 실패하면 멈춥니다.\n\n"
            "확인을 누르면 3초 후 실행합니다.",
        ):
            return

        mutex_handle = acquire_run_mutex()
        if mutex_handle is None:
            messagebox.showerror(
                "중복 실행 방지",
                "다른 프로세스(예: Windows 작업 스케줄러로 실행된 인스턴스)에서 이미 "
                "자동화를 실행 중인 것으로 보입니다.\n해당 실행이 끝난 뒤 다시 시도하세요.",
            )
            return
        self.run_mutex_handle = mutex_handle

        retry_count = int(self.retry_count_var.get())
        retry_delay = float(self.retry_delay_var.get())
        step_delay = float(self.step_delay_var.get())

        self.stop_event.clear()
        self.pause_event.clear()
        self.running = True
        self.run_button.configure(state=tk.DISABLED)
        self.pause_button.configure(state=tk.NORMAL, text="Ⅱ 일시정지")
        self.stop_button.configure(state=tk.NORMAL)
        self._start_run_log()
        self._write_log("INFO", f"선택 단계부터 실행 시작 (총 {len(remaining)}단계)")
        thread = threading.Thread(
            target=self._automation_worker,
            args=(remaining, [context], 1, step_delay, retry_count, retry_delay),
            daemon=True,
        )
        thread.start()

    def _profile_data(self) -> dict[str, Any]:
        return {
            "version": PROFILE_VERSION,
            "app": APP_TITLE,
            "settings": {
                "repeat_count": max(1, int(self.repeat_var.get())),
                "step_delay": max(0.0, float(self.step_delay_var.get())),
                "retry_count": max(0, int(self.retry_count_var.get())),
                "retry_delay": max(0.0, float(self.retry_delay_var.get())),
                "log_dir": self.log_dir_var.get(),
                "log_retention_days": max(0, int(self.log_retention_var.get())),
            },
            "data_source": {
                "enabled": bool(self.data_enabled_var.get()),
                "path": self.data_path_var.get(),
                "sheet_name": self.sheet_name_var.get(),
                "header_row": max(1, int(self.header_row_var.get())),
                "start_row": max(1, int(self.start_row_var.get())),
                "end_row": max(0, int(self.end_row_var.get())),
            },
            "schedule": {
                "enabled": bool(self.schedule_enabled_var.get()),
                "time": self.schedule_time_var.get(),
                "daily": bool(self.schedule_daily_var.get()),
                "start_minimized": bool(self.schedule_minimized_var.get()),
                "exit_after_run": bool(self.schedule_exit_after_var.get()),
            },
            "actions": [asdict(action) for action in self.actions],
        }

    def _write_profile_to(self, path: Path) -> bool:
        try:
            Path(path).write_text(
                json.dumps(self._profile_data(), ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        except (OSError, ValueError) as error:
            messagebox.showerror("저장 실패", str(error))
            return False
        self.current_profile_path = Path(path).resolve()
        self.status_var.set(f"설정 저장 완료: {self.current_profile_path.name}")
        return True

    def save_profile(self) -> bool:
        """빠른 저장: 이미 열려 있던(불러왔거나 이전에 저장한) 파일이 있으면 그대로
        덮어쓰고, 처음 저장하는 경우에는 '다른 이름으로 저장'과 동일하게 동작한다."""
        if not self.actions:
            messagebox.showwarning("저장할 단계 없음", "먼저 실행 단계를 등록하세요.")
            return False
        if self.current_profile_path is not None:
            return self._write_profile_to(self.current_profile_path)
        return self.save_profile_as()

    def save_profile_as(self) -> bool:
        """다른 이름으로 저장: 항상 파일 이름/위치를 새로 물어본다."""
        if not self.actions:
            messagebox.showwarning("저장할 단계 없음", "먼저 실행 단계를 등록하세요.")
            return False
        initial = self.current_profile_path.name if self.current_profile_path else "RPA_자동화_설정.json"
        path = filedialog.asksaveasfilename(
            title="자동화 설정 다른 이름으로 저장",
            defaultextension=".json",
            filetypes=[("자동화 설정", "*.json"), ("모든 파일", "*.*")],
            initialfile=initial,
        )
        if not path:
            return False
        return self._write_profile_to(Path(path))

    def new_profile(self) -> None:
        """새로 만들기: 현재 등록된 단계를 지우고 저장 파일 연결을 해제한다."""
        if self.running or self.recording:
            return
        if self.actions and not messagebox.askyesno(
            "새로 만들기", "저장하지 않은 변경 사항이 있을 수 있습니다.\n현재 단계를 모두 지우고 새로 시작할까요?"
        ):
            return
        self.actions.clear()
        self.current_profile_path = None
        self._refresh_list()
        self.status_var.set("새 자동화를 시작합니다.")

    def load_profile(self) -> None:
        path = filedialog.askopenfilename(
            title="자동화 설정 불러오기",
            filetypes=[("자동화 설정", "*.json"), ("모든 파일", "*.*")],
        )
        if not path:
            return
        try:
            self._load_profile_path(Path(path))
        except Exception as error:
            messagebox.showerror("불러오기 실패", f"설정 파일을 확인하세요.\n\n{error}")

    def _load_profile_path(self, path: Path) -> None:
        data = json.loads(path.read_text(encoding="utf-8"))
        version = int(data.get("version", 1))
        if version not in {1, 2, 3, 4}:
            raise ValueError(f"지원하지 않는 설정 파일 버전입니다: {version}")
        loaded = [Action.from_dict(item) for item in data.get("actions", [])]
        self._validate_actions(loaded)
        if version == 1:
            settings = {
                "repeat_count": data.get("repeat_count", 1),
                "step_delay": data.get("step_delay", 0.4),
            }
            data_source: dict[str, Any] = {}
            schedule: dict[str, Any] = {}
        else:
            settings = dict(data.get("settings", {}))
            data_source = dict(data.get("data_source", {}))
            schedule = dict(data.get("schedule", {}))
        self.actions = loaded
        self.repeat_var.set(max(1, int(settings.get("repeat_count", 1))))
        self.step_delay_var.set(max(0.0, float(settings.get("step_delay", 0.25))))
        self.retry_count_var.set(max(0, int(settings.get("retry_count", 2))))
        self.retry_delay_var.set(max(0.0, float(settings.get("retry_delay", 1.0))))
        self.log_dir_var.set(str(settings.get("log_dir") or default_log_directory()))
        self.log_retention_var.set(max(0, int(settings.get("log_retention_days", 30))))
        self.data_enabled_var.set(bool(data_source.get("enabled", False)))
        self.data_path_var.set(str(data_source.get("path", "")))
        self.sheet_name_var.set(str(data_source.get("sheet_name", "")))
        self.header_row_var.set(max(1, int(data_source.get("header_row", 1))))
        self.start_row_var.set(max(1, int(data_source.get("start_row", 2))))
        self.end_row_var.set(max(0, int(data_source.get("end_row", 0))))
        self.schedule_enabled_var.set(bool(schedule.get("enabled", False)))
        self.schedule_time_var.set(str(schedule.get("time", "09:00")))
        self.schedule_daily_var.set(bool(schedule.get("daily", True)))
        self.schedule_minimized_var.set(bool(schedule.get("start_minimized", True)))
        self.schedule_exit_after_var.set(bool(schedule.get("exit_after_run", False)))
        self.current_profile_path = path.resolve()
        self._refresh_list()
        self.status_var.set(f"설정 불러오기 완료: {path.name} (v{version} 호환)")

    @staticmethod
    def _validate_actions(actions: list[Action]) -> None:
        if len(actions) > MAX_ACTIONS:
            raise ValueError(f"실행 단계는 최대 {MAX_ACTIONS:,}개입니다.")
        allowed = {
            "click",
            "double_click",
            "right_click",
            "middle_click",
            "move",
            "drag",
            "scroll",
            "wait",
            "hotkey",
            "key",
            "key_down",
            "key_up",
            "text",
            "image_click",
            "wait_window",
            "image_condition_skip",
        }
        for action in actions:
            if action.kind not in allowed:
                raise ValueError(f"허용되지 않은 실행 단계: {action.kind}")
            p = action.params
            if action.kind in {"click", "double_click", "right_click", "middle_click", "move"}:
                int(p["x"])
                int(p["y"])
            elif action.kind == "drag":
                for key in ("x1", "y1", "x2", "y2"):
                    int(p[key])
            elif action.kind == "scroll":
                int(p["clicks"])
            elif action.kind == "wait":
                seconds = float(p["seconds"])
                if not 0 <= seconds <= 86400:
                    raise ValueError("대기시간 범위는 0~86,400초입니다.")
            elif action.kind == "hotkey":
                keys = p["keys"]
                if not isinstance(keys, list) or not keys:
                    raise ValueError("단축키가 비어 있습니다.")
            elif action.kind in {"key", "key_down", "key_up"}:
                if not str(p["key"]).strip():
                    raise ValueError("키 이름이 비어 있습니다.")
            elif action.kind == "text":
                str(p.get("template", ""))
            elif action.kind == "image_click":
                if not str(p.get("image", "")).strip():
                    raise ValueError("이미지 경로가 비어 있습니다.")
                confidence = float(p.get("confidence", 0.85))
                if not 0.1 <= confidence <= 1.0:
                    raise ValueError("이미지 정확도 범위는 0.1~1.0입니다.")
            elif action.kind == "wait_window":
                if not str(p.get("title", "")).strip():
                    raise ValueError("창 제목이 비어 있습니다.")
            elif action.kind == "image_condition_skip":
                if not str(p.get("image", "")).strip():
                    raise ValueError("이미지 경로가 비어 있습니다.")
                confidence = float(p.get("confidence", 0.85))
                if not 0.1 <= confidence <= 1.0:
                    raise ValueError("이미지 정확도 범위는 0.1~1.0입니다.")
                if p.get("condition") not in {"found", "not_found"}:
                    raise ValueError("조건 분기 종류가 올바르지 않습니다.")
                if int(p.get("skip_count", 0)) < 1:
                    raise ValueError("건너뛸 단계 수는 1 이상이어야 합니다.")

    def choose_data_file(self) -> None:
        path = filedialog.askopenfilename(
            title="CSV/Excel 데이터 선택",
            filetypes=[
                ("지원 데이터", "*.csv *.tsv *.xlsx *.xlsm"),
                ("Excel", "*.xlsx *.xlsm"),
                ("CSV", "*.csv *.tsv"),
            ],
        )
        if path:
            self.data_path_var.set(path)

    def _get_data_rows(self) -> list[dict[str, Any]]:
        if not self.data_enabled_var.get():
            return [{}]
        path = Path(self.data_path_var.get()).expanduser()
        end = int(self.end_row_var.get())
        rows = load_tabular_rows(
            path,
            self.sheet_name_var.get().strip(),
            int(self.header_row_var.get()),
            int(self.start_row_var.get()),
            end if end > 0 else None,
        )
        if not rows:
            raise ValueError("실행할 데이터 행이 없습니다.")
        return rows

    def preview_data(self) -> None:
        try:
            rows = self._get_data_rows()
        except Exception as error:
            messagebox.showerror("데이터 확인 실패", str(error))
            return
        sample = rows[:3]
        keys = list(sample[0].keys())[:10] if sample else []
        lines = [f"총 데이터 행: {len(rows):,}개", f"사용 가능 변수: {', '.join(keys)}", ""]
        for index, row in enumerate(sample, start=1):
            safe_values = [f"{key}={str(row.get(key, ''))[:30]}" for key in keys]
            lines.append(f"미리보기 {index}: " + ", ".join(safe_values))
        messagebox.showinfo("데이터 미리 확인", "\n".join(lines))

    def choose_log_directory(self) -> None:
        path = filedialog.askdirectory(title="로그 폴더 선택")
        if path:
            self.log_dir_var.set(path)

    def open_log_directory(self) -> None:
        path = Path(self.log_dir_var.get()).expanduser()
        path.mkdir(parents=True, exist_ok=True)
        if WINDOWS:
            os.startfile(path)  # type: ignore[attr-defined]
        else:
            messagebox.showinfo("로그 폴더", str(path))

    def _start_run_log(self) -> None:
        directory = Path(self.log_dir_var.get()).expanduser()
        directory.mkdir(parents=True, exist_ok=True)
        self.log_directory = directory
        self.log_file = directory / f"RPA_{datetime.now():%Y%m%d_%H%M%S}.log"
        self._write_log("INFO", "실행 시작")
        retention_days = max(0, int(self.log_retention_var.get()))
        if retention_days > 0:
            self._cleanup_old_logs(directory, retention_days)

    def _cleanup_old_logs(self, directory: Path, retention_days: int) -> None:
        cutoff = time.time() - retention_days * 86400
        try:
            entries = list(directory.iterdir())
        except OSError:
            return
        for entry in entries:
            if entry == self.log_file:
                continue
            try:
                if entry.stat().st_mtime < cutoff:
                    if entry.is_dir():
                        shutil.rmtree(entry, ignore_errors=True)
                    else:
                        entry.unlink()
            except OSError:
                continue

    def _write_log(self, level: str, message: str) -> None:
        if not self.log_file:
            return
        safe_message = message.replace("\r", " ").replace("\n", " ")
        line = f"{datetime.now():%Y-%m-%d %H:%M:%S.%f}"[:-3] + f" [{level}] {safe_message}\n"
        try:
            with self.log_lock:
                with self.log_file.open("a", encoding="utf-8") as stream:
                    stream.write(line)
        except OSError:
            pass

    def _capture_failure_screenshot(self, step_number: int) -> Path | None:
        # 실행 중인 워커 스레드에서 호출되므로 Tk 변수(self.log_dir_var)를 직접 읽지 않고
        # 실행 시작 시 메인 스레드에서 캡처해 둔 self.log_directory만 사용한다.
        if pyautogui is None or self.log_directory is None:
            return None
        try:
            directory = self.log_directory / f"{datetime.now():%Y%m%d}"
            directory.mkdir(parents=True, exist_ok=True)
            path = directory / f"failure_{datetime.now():%H%M%S}_{step_number:04d}.png"
            pyautogui.screenshot(str(path))
            return path
        except Exception:
            return None

    def start_automation(self, confirm: bool = True) -> None:
        if pyautogui is None:
            self._show_dependency_error()
            return
        if self.running or self.recording:
            return
        enabled_actions = [action.clone() for action in self.actions if action.enabled]
        if not enabled_actions:
            messagebox.showwarning("실행 단계 없음", "사용 상태인 실행 단계를 등록하세요.")
            return
        try:
            self._validate_actions(enabled_actions)
            repeat_count = int(self.repeat_var.get())
            step_delay = float(self.step_delay_var.get())
            retry_count = int(self.retry_count_var.get())
            retry_delay = float(self.retry_delay_var.get())
            data_rows = self._get_data_rows()
        except Exception as error:
            messagebox.showerror("설정 오류", str(error))
            return
        if not 1 <= repeat_count <= 10000:
            messagebox.showerror("설정 오류", "반복 횟수는 1~10,000 사이여야 합니다.")
            return
        if not 0 <= step_delay <= 30 or not 0 <= retry_count <= 20 or not 0 <= retry_delay <= 300:
            messagebox.showerror("설정 오류", "대기시간 또는 재시도 설정 범위를 확인하세요.")
            return
        if confirm and not messagebox.askyesno(
            "자동화 실행 확인",
            f"사용 단계: {len(enabled_actions)}개\n"
            f"데이터 행: {len(data_rows)}개\n"
            f"전체 반복: {repeat_count}회\n"
            f"최대 실행 단계: {len(enabled_actions) * len(data_rows) * repeat_count:,}회\n\n"
            "3초 후 실행합니다. 대상 프로그램을 준비했나요?",
        ):
            return

        mutex_handle = acquire_run_mutex()
        if mutex_handle is None:
            messagebox.showerror(
                "중복 실행 방지",
                "다른 프로세스(예: Windows 작업 스케줄러로 실행된 인스턴스)에서 이미 "
                "자동화를 실행 중인 것으로 보입니다.\n해당 실행이 끝난 뒤 다시 시도하세요.",
            )
            return
        self.run_mutex_handle = mutex_handle

        self.stop_event.clear()
        self.pause_event.clear()
        self.running = True
        self.run_button.configure(state=tk.DISABLED)
        self.pause_button.configure(state=tk.NORMAL, text="Ⅱ 일시정지")
        self.stop_button.configure(state=tk.NORMAL)
        self._start_run_log()
        thread = threading.Thread(
            target=self._automation_worker,
            args=(enabled_actions, data_rows, repeat_count, step_delay, retry_count, retry_delay),
            daemon=True,
        )
        thread.start()

    def _check_control(self) -> None:
        if self.stop_event.is_set():
            raise AutomationStopped()
        while self.pause_event.is_set():
            if self.stop_event.wait(0.1):
                raise AutomationStopped()

    def _interruptible_wait(self, seconds: float) -> None:
        remaining = max(0.0, seconds)
        last = time.monotonic()
        while remaining > 0:
            self._check_control()
            wait_for = min(0.05, remaining)
            if self.stop_event.wait(wait_for):
                raise AutomationStopped()
            now = time.monotonic()
            if not self.pause_event.is_set():
                remaining -= now - last
            last = now

    def _automation_worker(
        self,
        actions: list[Action],
        data_rows: list[dict[str, Any]],
        repeat_count: int,
        step_delay: float,
        retry_count: int,
        retry_delay: float,
    ) -> None:
        error_message: str | None = None
        stopped = False
        start_time = time.monotonic()
        succeeded_steps = 0
        skipped_steps = 0
        retry_events = 0
        try:
            # 스튜디오 창이 계속 포커스를 갖고 있으면 클릭 없이 바로 단축키를 실행하는
            # 단계에서 키 입력이 대상 프로그램이 아니라 이 창으로 들어가 버린다.
            # 녹화 시작과 마찬가지로 최소화해 대상 프로그램에 포커스를 넘겨준다.
            self.root.after(0, self.root.iconify)
            for countdown in (3, 2, 1):
                self._set_status(f"{countdown}초 후 자동화를 시작합니다...")
                self._interruptible_wait(1)

            total = repeat_count * len(data_rows) * len(actions)
            completed = 0
            for repeat_index in range(repeat_count):
                for data_index, context in enumerate(data_rows):
                    action_index = 0
                    while action_index < len(actions):
                        action = actions[action_index]
                        self._check_control()
                        step_number = completed + 1
                        self._set_status(
                            f"실행 중 {step_number:,}/{total:,} · 반복 {repeat_index + 1}/{repeat_count} · "
                            f"데이터 {data_index + 1}/{len(data_rows)} · {action.label()}"
                        )
                        skip_amount = 0
                        for attempt in range(retry_count + 1):
                            try:
                                skip_amount = self._execute_action(action, context)
                                self._write_log(
                                    "INFO",
                                    f"단계 성공 step={step_number} kind={action.kind} "
                                    f"data_row={context.get('row', '-')}"
                                    + (f" skip={skip_amount}" if skip_amount else ""),
                                )
                                succeeded_steps += 1
                                if skip_amount:
                                    skipped_steps += skip_amount
                                break
                            except AutomationStopped:
                                raise
                            except Exception as error:
                                retry_events += 1
                                self._write_log(
                                    "WARN",
                                    f"단계 실패 step={step_number} kind={action.kind} "
                                    f"attempt={attempt + 1}/{retry_count + 1} "
                                    f"error={type(error).__name__}: {error}",
                                )
                                if attempt >= retry_count:
                                    screenshot = self._capture_failure_screenshot(step_number)
                                    if screenshot:
                                        self._write_log("ERROR", f"실패 화면 저장: {screenshot.name}")
                                    raise
                                self._set_status(
                                    f"단계 {step_number} 실패 · {retry_delay:.1f}초 후 "
                                    f"재시도 {attempt + 2}/{retry_count + 1}"
                                )
                                self._interruptible_wait(retry_delay)
                        completed += 1
                        if completed < total:
                            self._interruptible_wait(step_delay)
                        action_index += 1 + skip_amount
        except AutomationStopped:
            stopped = True
            self._write_log("INFO", "사용자 중지")
        except Exception as error:
            self.stop_event.set()
            error_message = f"{type(error).__name__}: {error}"
            self._write_log("ERROR", error_message)
            self._write_log("DEBUG", traceback.format_exc())
        finally:
            elapsed = time.monotonic() - start_time
            summary = {
                "succeeded_steps": succeeded_steps,
                "skipped_steps": skipped_steps,
                "retry_events": retry_events,
                "elapsed": elapsed,
            }
            self._write_log(
                "INFO",
                f"요약: 성공 {succeeded_steps}단계, 조건 분기로 건너뜀 {skipped_steps}단계, "
                f"재시도 발생 {retry_events}회, 소요 시간 {elapsed:.1f}초",
            )
            self.root.after(0, self._finish_automation, stopped, error_message, summary)

    def _execute_action(self, action: Action, context: dict[str, Any]) -> int:
        """단계를 실행한다. 조건 분기가 발동하면 건너뛸 다음 단계 수를 반환하고,
        그 외에는 0을 반환한다."""
        self._check_control()
        p = action.params
        if action.kind == "click":
            pyautogui.click(int(p["x"]), int(p["y"]), duration=0.15)
        elif action.kind == "double_click":
            pyautogui.doubleClick(int(p["x"]), int(p["y"]), interval=0.12, duration=0.15)
        elif action.kind == "right_click":
            pyautogui.rightClick(int(p["x"]), int(p["y"]), duration=0.15)
        elif action.kind == "middle_click":
            pyautogui.middleClick(int(p["x"]), int(p["y"]), duration=0.15)
        elif action.kind == "move":
            pyautogui.moveTo(int(p["x"]), int(p["y"]), duration=float(p.get("duration", 0.2)))
        elif action.kind == "drag":
            pyautogui.moveTo(int(p["x1"]), int(p["y1"]), duration=0.15)
            pyautogui.dragTo(
                int(p["x2"]),
                int(p["y2"]),
                duration=float(p.get("duration", 0.6)),
                button="left",
            )
        elif action.kind == "scroll":
            pyautogui.scroll(int(p["clicks"]))
        elif action.kind == "wait":
            self._interruptible_wait(float(p["seconds"]))
        elif action.kind == "hotkey":
            pyautogui.hotkey(*[str(key) for key in p["keys"]])
        elif action.kind == "key":
            pyautogui.press(str(p["key"]))
        elif action.kind == "key_down":
            pyautogui.keyDown(str(p["key"]))
        elif action.kind == "key_up":
            pyautogui.keyUp(str(p["key"]))
        elif action.kind == "text":
            value = render_template(str(p.get("template", "")), context)
            if p.get("use_clipboard"):
                self._type_via_clipboard(value)
            else:
                send_unicode_text_windows(value, float(p.get("interval", 0.01)))
        elif action.kind == "image_click":
            point = self._wait_for_image(
                Path(str(p["image"])),
                float(p.get("confidence", 0.85)),
                float(p.get("timeout", 15)),
            )
            clicks = int(p.get("clicks", 1))
            if clicks == 1:
                pyautogui.click(point.x, point.y)
            elif clicks == 2:
                pyautogui.doubleClick(point.x, point.y, interval=0.12)
            elif clicks == 3:
                pyautogui.click(point.x, point.y, clicks=3, interval=0.12)
        elif action.kind == "wait_window":
            self._wait_for_window(
                str(p["title"]),
                bool(p.get("regex", False)),
                float(p.get("timeout", 30)),
                bool(p.get("activate", True)),
            )
        elif action.kind == "image_condition_skip":
            present = self._check_image_present(
                Path(str(p["image"])), float(p.get("confidence", 0.85))
            )
            condition_met = present if p.get("condition") == "found" else not present
            if condition_met:
                return int(p.get("skip_count", 1))
        else:
            raise ValueError(f"지원하지 않는 실행 단계: {action.kind}")
        return 0

    @staticmethod
    def _type_via_clipboard(value: str) -> None:
        """일부 프로그램은 SendInput 유니코드 직접 입력 방식으로 한글 등이 제대로
        들어가지 않을 수 있다. 이런 경우의 대안으로 클립보드에 붙여넣고(Ctrl+V) 실행
        직후 원래 클립보드 내용을 복원한다. Tk 객체를 쓰지 않는 Win32 API만 사용하므로
        워커 스레드에서 안전하게 호출할 수 있다."""
        if not WINDOWS:
            raise RuntimeError("클립보드 붙여넣기 입력은 Windows에서 지원됩니다.")
        previous = win32_get_clipboard_text()
        try:
            win32_set_clipboard_text(value)
            time.sleep(0.05)
            pyautogui.hotkey("ctrl", "v")
            time.sleep(0.05)
        finally:
            win32_set_clipboard_text(previous if previous is not None else "")

    @staticmethod
    def _locate_image_once(image: Path, confidence: float) -> Any:
        try:
            return pyautogui.locateCenterOnScreen(str(image), confidence=confidence)
        except TypeError as error:
            raise RuntimeError(
                "이미지 정확도 기능에는 OpenCV가 필요합니다: py -m pip install opencv-python"
            ) from error
        except Exception as error:
            image_not_found = getattr(pyautogui, "ImageNotFoundException", None)
            if image_not_found is not None and isinstance(error, image_not_found):
                return None
            raise

    def _wait_for_image(self, image: Path, confidence: float, timeout: float) -> Any:
        if not image.exists():
            raise FileNotFoundError(f"이미지 파일을 찾을 수 없습니다: {image}")
        deadline = time.monotonic() + timeout
        while True:
            self._check_control()
            point = self._locate_image_once(image, confidence)
            if point is not None:
                return point
            if time.monotonic() >= deadline:
                raise TimeoutError(f"화면에서 이미지를 찾지 못했습니다: {image.name}")
            self._interruptible_wait(0.4)

    def _check_image_present(self, image: Path, confidence: float) -> bool:
        """조건 분기용: 대기 없이 현재 화면에 이미지가 있는지 한 번만 확인한다."""
        if not image.exists():
            raise FileNotFoundError(f"이미지 파일을 찾을 수 없습니다: {image}")
        self._check_control()
        return self._locate_image_once(image, confidence) is not None

    def _wait_for_window(self, title: str, regex: bool, timeout: float, activate: bool) -> None:
        if not WINDOWS:
            raise RuntimeError("창 제목 대기는 Windows에서만 지원됩니다.")
        deadline = time.monotonic() + timeout
        while True:
            self._check_control()
            found = find_window(title, regex)
            if found:
                if activate:
                    activate_window(found[0])
                return
            if time.monotonic() >= deadline:
                raise TimeoutError(f"창을 찾지 못했습니다: {title}")
            self._interruptible_wait(0.3)

    def toggle_pause(self) -> None:
        if not self.running:
            return
        if self.pause_event.is_set():
            self.pause_event.clear()
            self.pause_button.configure(text="Ⅱ 일시정지")
            self.status_var.set("자동화 실행을 계속합니다.")
            self._write_log("INFO", "실행 계속")
        else:
            self.pause_event.set()
            self.pause_button.configure(text="▶ 계속")
            self.status_var.set("자동화가 일시정지되었습니다.")
            self._write_log("INFO", "일시정지")

    def stop_automation(self) -> None:
        if self.running:
            self.stop_event.set()
            self.pause_event.clear()
            self.status_var.set("중지 요청을 처리하고 있습니다...")

    @staticmethod
    def _format_elapsed(seconds: float) -> str:
        minutes, secs = divmod(int(seconds), 60)
        hours, minutes = divmod(minutes, 60)
        if hours:
            return f"{hours}시간 {minutes}분 {secs}초"
        if minutes:
            return f"{minutes}분 {secs}초"
        return f"{secs}초"

    def _finish_automation(
        self,
        stopped: bool,
        error_message: str | None,
        summary: dict[str, Any] | None = None,
    ) -> None:
        release_run_mutex(self.run_mutex_handle)
        self.run_mutex_handle = None
        self.running = False
        self.pause_event.clear()
        self.root.deiconify()
        self.run_button.configure(state=tk.NORMAL)
        self.pause_button.configure(state=tk.DISABLED, text="Ⅱ 일시정지")
        self.stop_button.configure(state=tk.DISABLED)
        # 예약 실행이 "완료 후 자동 종료"로 설정된 경우, 지켜보는 사람이 없으므로
        # 응답을 기다리는 모달 대화상자를 띄우지 않는다.
        silent = self.is_autorun_session and self.schedule_exit_after_var.get()
        summary_text = ""
        if summary:
            summary_text = (
                f"성공 {summary['succeeded_steps']:,}단계 · "
                f"조건 분기로 건너뜀 {summary['skipped_steps']:,}단계 · "
                f"재시도 발생 {summary['retry_events']:,}회 · "
                f"소요 시간 {self._format_elapsed(summary['elapsed'])}"
            )
        if error_message:
            self.status_var.set("오류로 자동화가 중지되었습니다. 로그를 확인하세요.")
            if not silent:
                messagebox.showwarning(
                    "자동화 중지",
                    f"{error_message}\n\n로그: {self.log_file or '기록 없음'}",
                )
        elif stopped or self.stop_event.is_set():
            self.status_var.set("사용자 요청으로 자동화를 중지했습니다.")
        else:
            self.status_var.set(
                "자동화 실행이 정상적으로 완료되었습니다."
                + (f" ({summary_text})" if summary_text else "")
            )
            self._write_log("INFO", "정상 완료")
            if summary_text and not silent:
                messagebox.showinfo("실행 완료", summary_text)
        if silent:
            self.root.after(500, self.root.destroy)

    @staticmethod
    def _valid_time(value: str) -> bool:
        try:
            datetime.strptime(value.strip(), "%H:%M")
            return True
        except ValueError:
            return False

    def _scheduler_tick(self) -> None:
        try:
            now = datetime.now()
            stamp = now.strftime("%Y-%m-%d %H:%M")
            if (
                self.schedule_enabled_var.get()
                and not self.running
                and not self.recording
                and self._valid_time(self.schedule_time_var.get())
                and now.strftime("%H:%M") == self.schedule_time_var.get().strip()
                and stamp != self.last_schedule_stamp
            ):
                self.last_schedule_stamp = stamp
                self.start_automation(confirm=False)
                if not self.schedule_daily_var.get():
                    self.schedule_enabled_var.set(False)
        finally:
            if self.root.winfo_exists():
                self.root.after(1000, self._scheduler_tick)

    def _task_command(self) -> str:
        if not self.current_profile_path:
            raise ValueError("Windows 예약 등록 전에 설정 파일을 저장하세요.")
        if getattr(sys, "frozen", False):
            parts = [str(Path(sys.executable).resolve()), "--profile", str(self.current_profile_path), "--run"]
        else:
            parts = [
                str(Path(sys.executable).resolve()),
                str(Path(__file__).resolve()),
                "--profile",
                str(self.current_profile_path),
                "--run",
            ]
        return subprocess.list2cmdline(parts)

    def register_windows_task(self) -> None:
        if not WINDOWS:
            messagebox.showwarning("Windows 전용", "Windows 작업 스케줄러 등록은 Windows에서 지원됩니다.")
            return
        if not self._valid_time(self.schedule_time_var.get()):
            messagebox.showerror("시간 오류", "실행 시각을 HH:MM 형식으로 입력하세요.")
            return
        if not self.save_profile():
            return
        task_name = simpledialog.askstring(
            "Windows 예약 등록",
            "예약 작업 이름",
            initialvalue="RPA_자동화_예약",
            parent=self.root,
        )
        if not task_name:
            return
        if not messagebox.askyesno(
            "Windows 예약 등록",
            f"작업 이름: {task_name}\n실행 시각: {self.schedule_time_var.get()}\n"
            f"반복: {'매일' if self.schedule_daily_var.get() else '1회'}\n\n등록할까요?",
        ):
            return
        command = [
            "schtasks",
            "/Create",
            "/TN",
            task_name,
            "/TR",
            self._task_command(),
            "/ST",
            self.schedule_time_var.get().strip(),
            "/F",
        ]
        if self.schedule_daily_var.get():
            command.extend(["/SC", "DAILY"])
        else:
            now = datetime.now()
            scheduled = datetime.strptime(self.schedule_time_var.get().strip(), "%H:%M").time()
            run_date = now.date()
            if scheduled <= now.time():
                from datetime import timedelta

                run_date += timedelta(days=1)
            command.extend(["/SC", "ONCE", "/SD", run_date.strftime("%m/%d/%Y")])
        result = subprocess.run(command, capture_output=True, text=True, encoding="cp949", errors="replace")
        if result.returncode == 0:
            messagebox.showinfo("예약 등록 완료", f"Windows 예약 작업을 등록했습니다.\n\n{task_name}")
        else:
            messagebox.showerror("예약 등록 실패", result.stderr or result.stdout)

    def delete_windows_task(self) -> None:
        if not WINDOWS:
            messagebox.showwarning("Windows 전용", "Windows 작업 스케줄러 삭제는 Windows에서 지원됩니다.")
            return
        task_name = simpledialog.askstring(
            "Windows 예약 삭제",
            "삭제할 예약 작업 이름",
            initialvalue="RPA_자동화_예약",
            parent=self.root,
        )
        if not task_name:
            return
        if not messagebox.askyesno("Windows 예약 삭제", f"'{task_name}' 예약을 삭제할까요?"):
            return
        result = subprocess.run(
            ["schtasks", "/Delete", "/TN", task_name, "/F"],
            capture_output=True,
            text=True,
            encoding="cp949",
            errors="replace",
        )
        if result.returncode == 0:
            messagebox.showinfo("예약 삭제 완료", task_name)
        else:
            messagebox.showerror("예약 삭제 실패", result.stderr or result.stdout)

    def on_close(self) -> None:
        if self.recording:
            self.stop_recording()
        if self.running:
            if not messagebox.askyesno("실행 중", "자동화를 중지하고 프로그램을 종료할까요?"):
                return
            self.stop_event.set()
        self.root.destroy()


def parse_args(argv: Iterable[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=APP_TITLE)
    parser.add_argument("--profile", type=Path, help="시작 시 불러올 JSON 설정")
    parser.add_argument("--run", action="store_true", help="설정 로드 후 자동 실행")
    return parser.parse_args(argv)


def main(argv: Iterable[str] | None = None) -> None:
    args = parse_args(argv)
    enable_windows_dpi_awareness()
    root = tk.Tk()
    try:
        ttk.Style(root).theme_use("vista")
    except tk.TclError:
        pass
    AutomationApp(root, args.profile, args.run)
    root.mainloop()


if __name__ == "__main__":
    main()
