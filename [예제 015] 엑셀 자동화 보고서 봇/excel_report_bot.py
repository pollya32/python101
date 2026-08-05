"""
엑셀 자동화 보고서 봇
=====================================================
설치: pip install -r requirements.txt
실행: python excel_report_bot.py
      (샘플 데이터가 없다면 먼저 python generate_sample_reports.py 실행)

input/ 폴더 안의 모든 xlsx 파일 - 파일마다 여러 시트(월별 등)로 나뉜 실적 데이터를
전부 읽어 하나로 취합하고, 지점별/월별/카테고리별 요약을 계산해
서식이 적용된 보고서 파일(output/보고서_YYYYMMDD_HHMMSS.xlsx)로 저장합니다.
"""
import argparse
import glob
import os
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_INPUT_DIR = os.path.join(BASE_DIR, 'input')
DEFAULT_OUTPUT_DIR = os.path.join(BASE_DIR, 'output')

REQUIRED_COLUMNS = ['날짜', '지점', '담당자', '카테고리', '매출액', '비용']


# ── 1. 데이터 취합 ───────────────────────────────────────────────────────────

def collect_raw_data(input_dir: str) -> pd.DataFrame:
    """input_dir 안의 모든 xlsx 파일 - 모든 시트를 읽어 하나의 DataFrame으로 취합"""
    files = sorted(glob.glob(os.path.join(glob.escape(input_dir), '*.xlsx')))
    if not files:
        raise FileNotFoundError(
            f"'{input_dir}'에 xlsx 파일이 없습니다. generate_sample_reports.py를 먼저 실행하세요."
        )

    frames = []
    for file_path in files:
        sheets = pd.read_excel(file_path, sheet_name=None, engine='openpyxl')
        for sheet_name, df in sheets.items():
            missing = set(REQUIRED_COLUMNS) - set(df.columns)
            if missing:
                print(f'  건너뜀: {os.path.basename(file_path)} / {sheet_name} (누락 컬럼: {missing})')
                continue
            df = df.copy()
            df['출처파일'] = os.path.basename(file_path)
            df['출처시트'] = sheet_name
            frames.append(df)

    if not frames:
        raise ValueError('취합할 유효한 데이터가 없습니다.')

    combined = pd.concat(frames, ignore_index=True)
    combined['날짜'] = pd.to_datetime(combined['날짜'])
    combined['순이익'] = combined['매출액'] - combined['비용']
    return combined.sort_values('날짜').reset_index(drop=True)


# ── 2. 요약 계산 ─────────────────────────────────────────────────────────────

def summarize(df: pd.DataFrame) -> dict:
    branch_summary = (
        df.groupby('지점', as_index=False)[['매출액', '비용', '순이익']]
        .sum()
        .sort_values('매출액', ascending=False)
    )

    monthly = df.copy()
    monthly['월'] = monthly['날짜'].dt.strftime('%Y-%m')
    monthly_summary = (
        monthly.groupby(['월', '지점'], as_index=False)[['매출액', '비용', '순이익']]
        .sum()
        .sort_values(['월', '지점'])
    )

    category_summary = (
        df.groupby('카테고리', as_index=False)[['매출액', '비용', '순이익']]
        .sum()
        .sort_values('매출액', ascending=False)
    )

    return {
        '지점별요약': branch_summary,
        '월별요약': monthly_summary,
        '카테고리별요약': category_summary,
    }


# ── 3. 서식이 적용된 엑셀 저장 ────────────────────────────────────────────────

HEADER_FILL = PatternFill('solid', fgColor='305496')
HEADER_FONT = Font(color='FFFFFF', bold=True)
THIN_SIDE = Side(style='thin', color='D9D9D9')
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
CURRENCY_COLUMNS = {'매출액', '비용', '순이익'}
DATE_COLUMNS = {'날짜'}


def style_sheet(ws, df: pd.DataFrame):
    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = THIN_BORDER
            if cell.row > 1:
                col_name = df.columns[cell.column - 1]
                if col_name in CURRENCY_COLUMNS:
                    cell.number_format = '#,##0"원"'
                elif col_name in DATE_COLUMNS:
                    cell.number_format = 'yyyy-mm-dd'

    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name]])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    ws.freeze_panes = 'A2'


def write_report(raw: pd.DataFrame, summaries: dict, output_path: str):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        raw.to_excel(writer, sheet_name='원본취합', index=False)
        for sheet_name, df in summaries.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

        style_sheet(writer.sheets['원본취합'], raw)
        for sheet_name, df in summaries.items():
            style_sheet(writer.sheets[sheet_name], df)


# ── 4. 실행 ──────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description='여러 엑셀 보고서를 취합해 요약 보고서를 생성합니다.')
    parser.add_argument('--input', default=DEFAULT_INPUT_DIR, help='원본 xlsx 파일들이 있는 폴더')
    parser.add_argument('--output', default=DEFAULT_OUTPUT_DIR, help='결과 보고서를 저장할 폴더')
    args = parser.parse_args()

    print(f"'{args.input}'에서 데이터를 취합하는 중...")
    raw = collect_raw_data(args.input)
    print(f'  총 {len(raw)}건 취합 완료')

    summaries = summarize(raw)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_path = os.path.join(args.output, f'보고서_{timestamp}.xlsx')
    write_report(raw, summaries, output_path)
    print(f'보고서 생성 완료: {output_path}')


if __name__ == '__main__':
    main()
